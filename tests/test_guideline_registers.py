"""Prompt-level regression checks for MF/REF accounting and rendering."""

import sys
import tempfile
import unittest
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import build_guideline_registers as register


class GuidelineRegisterTests(unittest.TestCase):
    def setUp(self):
        self.registers = register.Registers(register.read_headers())
        for index in (register.PRICE_MEDIAN, register.CLASS_COSTS, register.CLASS_MEDIAN,
                      register.FACILITY_MONTHS, register.LG_MONTHS, register.ALL_MONTHS):
            index.clear()
        self.source = {
            "Equipment/Item": "Desk", "Local Government": "Hoima",
            "Facility": "Example Seed Secondary School", "Facility type": "School",
            "Cost": 120000, "Date Placed In Service": date(2026, 7, 1),
            "Equipment status": "Functional", "_row": 2,
        }

    def row(self, source=None, costs=None):
        values = register.build_values(source or self.source, self.registers, borrow=True,
                                       costs=costs or {}, lives={}, dates={}, stats=Counter())
        return {header: value[0] if isinstance(value, tuple) else value
                for header, value in zip(self.registers.headers, values)}

    def test_immaterial_cost_has_no_depreciation(self):
        row = self.row({**self.source, "Cost": 5000})
        self.assertEqual(row["FIXED_ASSETS_COST"], 0)
        self.assertEqual(row["DEPRECIATE_FLAG"], "NO")
        self.assertIsNone(row["ASSET_TYPE"])

    def test_recorded_zero_cost_is_not_an_unpriced_blank(self):
        source = {**self.source, "Equipment/Item": "VLS", "Cost": 0}
        row = self.row(source)
        self.assertEqual(row["FIXED_ASSETS_COST"], 0)
        self.assertEqual(row[register.ATTRIBUTE[10]], 0)
        self.assertEqual(row["DEPRECIATE_FLAG"], "NO")
        self.assertIsNone(row["ASSET_TYPE"])
        self.assertEqual(row["DEPRN_RESERVE"], 0)
        self.assertEqual(row["YTD_DEPRN"], 0)
        for amount in (0, 12345):
            with self.subTest(recorded_amount=amount):
                row = self.row({**source, "Acc Dep Cost": amount, "Ytd Deprn": amount, "Net Book Value": amount})
                self.assertEqual(row["DEPRN_RESERVE"], amount)
                self.assertEqual(row["YTD_DEPRN"], amount)
                self.assertEqual(row[register.ATTRIBUTE[12]], amount)

    def test_explicit_software_uses_annex_software_account(self):
        for item in (
            "Software Licences for the computers: Basic office suite, antivirus/antimalware, utilities",
            "Computer Software",
        ):
            with self.subTest(item=item):
                row = self.row({**self.source, "Equipment/Item": item})
                self.assertEqual(row["ASSET_CATEGORY_MINOR2"], "COMPUTER SOFTWARE")
                self.assertEqual(row["ASSET_EXP_ACCT_ACCOUNT"], "231423")
                self.assertEqual(row["LIFE_IN_MONTHS"], 60)
        hardware = self.row({**self.source, "Equipment/Item": "Desktop computer with bundled software"})
        self.assertEqual(hardware["ASSET_CATEGORY_MINOR2"], "LIGHT ICT HARDWARE")
        for item in ("Annual software licence", "Computer software subscription"):
            with self.subTest(item=item):
                row = self.row({**self.source, "Equipment/Item": item})
                self.assertIsNone(row["ASSET_TYPE"])
                self.assertIsNone(row["ASSET_CATEGORY_MINOR2"])
                self.assertEqual(row["LIFE_IN_MONTHS"], 0)

    def test_generic_land_never_borrows_a_class_price(self):
        register.add_donor(register.CLASS_COSTS, {}, "land", "hoima", (2026,), "", 1000000, "Hoima")
        row = self.row({**self.source, "Equipment/Item": "Land", "Cost": None})
        self.assertIsNone(row["FIXED_ASSETS_COST"])

    def test_missing_cost_keeps_depreciation_amounts_zero(self):
        row = self.row({**self.source, "Equipment/Item": "Special apparatus", "Cost": None})
        self.assertIsNone(row["FIXED_ASSETS_COST"])
        self.assertEqual(row["DEPRN_RESERVE"], 0)
        self.assertEqual(row["YTD_DEPRN"], 0)

    def test_mixed_status_does_not_deny_every_unit(self):
        row = self.row({**self.source, "Equipment status": "18 in use, 10 stolen"})
        self.assertEqual(row[register.ATTRIBUTE[14]], "Functional")
        self.assertEqual(row["ASSET_TYPE"], "CAPITALIZED")
        self.assertGreater(row["DEPRN_RESERVE"], 0)

    def test_work_in_progress_has_no_service_date(self):
        row = self.row({**self.source, "Equipment/Item": "Classroom block",
                        "Remarks": "Under construction"})
        self.assertEqual(row["ASSET_TYPE"], "CIP")
        self.assertIsNone(row["DATE_PLACED_IN_SERVICE"])
        self.assertIsNone(row[register.ATTRIBUTE[8]])

    def test_laboratory_glassware_is_consumable(self):
        for name in ("Volumetric flask", "Conical flask", "Laboratory glassware"):
            with self.subTest(name=name):
                row = self.row({**self.source, "Equipment/Item": name})
                self.assertIsNone(row["ASSET_TYPE"])
                self.assertIsNone(row["ASSET_CATEGORY_MINOR2"])
                self.assertEqual(row["LIFE_IN_MONTHS"], 0)

    def test_outlier_only_local_donor_does_not_override_valid_other_vote(self):
        donor_rows = [self.source,
                      {**self.source, "Local Government": "Masindi", "Cost": 100000},
                      {**self.source, "Local Government": "Masindi", "Cost": 100000},
                      {**self.source, "Local Government": "Hoima", "Cost": 100000000}]
        # Give Hoima only the outlier; three valid namesakes establish the band.
        donor_rows[0] = {**donor_rows[0], "Local Government": "Masindi"}
        costs, _, _ = register.index_sources(iter(donor_rows), self.registers)
        register.index_fallbacks(iter(donor_rows), self.registers)
        register.finalize_cost_donors(costs)
        row = self.row({**self.source, "Cost": None}, costs=costs)
        self.assertEqual(row["FIXED_ASSETS_COST"], 100000)

    def test_depreciation_months_and_recorded_reserve(self):
        row = self.row()
        self.assertEqual(row["DEPRN_RESERVE"], 6000)
        self.assertEqual(row["YTD_DEPRN"], 6000)
        row = self.row({**self.source, "Acc Dep Cost": 119000})
        self.assertEqual(row["DEPRN_RESERVE"], 119000)
        self.assertEqual(row["YTD_DEPRN"], 1000)

    def test_preserves_recorded_figures_even_when_no_charge_applies(self):
        row = self.row({**self.source, "Cost": 5000, "Acc Dep Cost": 6000,
                        "Ytd Deprn": 700, "Net Book Value": 200})
        self.assertEqual(row["DEPRN_RESERVE"], 6000)
        self.assertEqual(row["YTD_DEPRN"], 700)
        self.assertEqual(row[register.ATTRIBUTE[12]], 200)

    def test_whole_group_absence_and_other_item_are_distinct(self):
        absent = self.row({**self.source, "Equipment status": "18 stolen"})
        self.assertIsNone(absent["ASSET_TYPE"])
        self.assertEqual(absent["DEPRN_RESERVE"], 0)
        other = self.row({**self.source, "Remarks": "The microscope was not delivered"})
        self.assertEqual(other["ASSET_TYPE"], "CAPITALIZED")

    def test_specific_asset_without_priced_namesake_can_borrow_class(self):
        minor = register.norm_name(self.registers.classify("Desk")[2])
        register.add_donor(register.CLASS_COSTS, {}, minor, "hoima", (2026,), "", 100000, "Hoima")
        self.assertEqual(self.row({**self.source, "Cost": None})["FIXED_ASSETS_COST"], 100000)

    def test_land_life_is_annex_life_without_depreciation(self):
        row = self.row({**self.source, "Equipment/Item": "Land", "Cost": 1000000})
        self.assertEqual(row["LIFE_IN_MONTHS"], 600)
        self.assertEqual(row["DEPRECIATE_FLAG"], "NO")
        self.assertEqual(row["ASSET_TYPE"], "CAPITALIZED")

    def test_preserves_duplicate_unit_rows_and_readable_headers(self):
        values = register.build_values(self.source, self.registers, borrow=True,
                                       costs={}, lives={}, dates={}, stats=Counter())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "register.xlsx"
            count = register.write_book(path, self.registers.headers, [values] * 120,
                                        lambda count: ["Read Me", "Rules " * 100])
            self.assertEqual(count, 120)
            book = load_workbook(path)
            sheet = book["Asset Register"]
            self.assertEqual(sheet.max_row, 121)
            self.assertEqual(sheet["A2"].value, sheet["A121"].value)
            self.assertTrue(sheet["A1"].alignment.wrap_text)
            self.assertGreaterEqual(sheet.row_dimensions[1].height, 60)
            self.assertGreaterEqual(book["Read Me"].column_dimensions["A"].width, 100)
            self.assertTrue(book["Read Me"]["A2"].alignment.wrap_text)
            book.close()


if __name__ == "__main__":
    unittest.main()
