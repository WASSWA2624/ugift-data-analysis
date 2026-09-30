from docx import Document

d = Document(r"D:\coding\ugift-data-analysis\outputs\narrative-report\UgIFT Asset Verification Report.docx")
print("--- contents around 4.2 ---")
for i, p in enumerate(d.paragraphs):
    if p.text.startswith("4.2") or p.text.startswith("Figure 46"):
        print(i, p.style.name if p.style else "", p.text[:110])
print("--- mda headings ---")
for i, p in enumerate(d.paragraphs):
    if p.style and p.style.name == "Heading 3" and p.text.startswith("4.2."):
        print(i, p.text)
print("--- sample narrative ---")
for i, p in enumerate(d.paragraphs):
    if p.text.startswith("MoFPED held"):
        print(p.text[:240])
        break
print("--- 4.3 follows ---")
for i, p in enumerate(d.paragraphs):
    if p.style and p.style.name == "Heading 2" and p.text.startswith("4.3"):
        print("4.3 at", i, "prev", d.paragraphs[i - 1].text[:80])
        break
