import re
from docx import Document

d = Document(r"D:\coding\ugift-data-analysis\outputs\narrative-report\UgIFT Asset Verification Report.docx")
xml = d.paragraphs[488]._p.xml
for m in re.finditer(r'cx="(\d+)" cy="(\d+)"', xml):
    print(m.group(1), m.group(2), int(m.group(1)) / 914400, int(m.group(2)) / 914400)
print("para align", d.paragraphs[488].alignment)
print("caption", d.paragraphs[489].alignment)
