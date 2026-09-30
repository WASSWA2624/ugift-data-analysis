"""Insert per-MDA narratives into section 4.2 of the UgIFT report."""
from docx import Document

PATH = r"D:\coding\ugift-data-analysis\outputs\narrative-report\UgIFT Asset Verification Report.docx"

SECTIONS = [
    (
        "4.2.5 Ministry of Finance, Planning and Economic Development",
        [
            "MoFPED held 458 assets: 217 ICT items, 122 other assets, 99 furniture and fittings and 20 transport assets. The main items were office chairs and desks, desktop computers, tablets and 20 Toyota Hilux double cabin pickups. All 458 assets (100.0%) were in good condition and in use, and none was broken. Of the 358 movable assets, 272 (76.0%) were engraved and 86 (24.0%) were not engraved; 252 (70.4%) were properly engraved with the UgIFT marking, and 20 (5.6%), the pickups, were not properly engraved, carrying their registration numbers only. The recorded value is UGX 6,152,854,000, accumulated depreciation to 30 September 2026 is UGX 4,272,041,257, and the net book value is UGX 1,880,812,743.",
            "The ministry held 113 computers (73 desktops and 40 laptops), all in good condition and in use; 103 (91.2%) were engraved and 10 (8.8%) were not engraved. Its 17 tablets and 14 office desk phones were all in good condition and engraved. The ministry keeps a UgIFT asset register and the consolidated register for all MDAs that received UgIFT assets.",
            "Expected: Staff computers should support their workload, and portable equipment should remain identifiable and accountable. Found: The computers were working, but users reported that some keyboards and power backup units had failed and that some older computers were freezing. Laptop losses were supported by police letters. The team also reported difficulty obtaining some laptops for engraving.",
            "Gap: Equipment reliability, follow-up of lost assets and completion of identification each required a specific management response. Recommended action: The ministry should assess the failed accessories and freezing computers, follow up the reported laptop losses, and arrange a supervised exercise to engrave all remaining portable equipment.",
            "Expected: Furniture should remain assigned to a location and custodian when offices move. Found: When programme staff moved to premises already furnished, their earlier furniture was left in the old building. Some remained in store, while staff said other items had gone to different offices.",
            "Gap: The move had split the furniture between storage and other offices, requiring confirmation of its present custody and use. Recommended action: The ministry should inspect the old store and recipient offices, confirm the condition and custodian of each item, and approve reuse or other appropriate treatment of furniture no longer required.",
        ],
    ),
    (
        "4.2.6 Ministry of Works and Transport",
        [
            "MoWT held 88 assets: 63 ICT items, 24 other assets and 1 transport asset, a Toyota Hilux pickup. The ICT items were mainly desktop computers, monitors and uninterruptible power supply (UPS) units. Of the 88 assets, 72 (81.8%) were in good condition and in use and 16 (18.2%) were not in use for another reason. None was broken. Of the 72 movable assets, 63 (87.5%) were engraved, 62 (86.1%) of them with the UgIFT marking, and 9 (12.5%) were not engraved. The recorded value is UGX 455,932,376, accumulated depreciation to 30 September 2026 is UGX 326,484,882, and the net book value is UGX 129,447,494.",
            "At the time of the visit, the ministry was tracing the location of the 16 assets not in use, 12 Dell OptiPlex desktop computers and 4 UPS units. None of them was reported damaged. Of the ministry's 33 computers (28 desktops and 5 laptops), 21 (63.6%) were in good condition and in use and the 12 desktops being traced made up the other 36.4%; 28 (84.8%) were engraved and 5 (15.2%) were not engraved.",
            "Repairs and servicing of the Toyota Hilux pickup were handled by MoFPED. The custody and maintenance arrangements should remain clear during handover.",
            "Expected: Installed office and conferencing equipment should support the secretariat and carry a traceable asset identity. Found: The secretariat reported that the installed photocopiers supported its daily work. The colour photocopier, the video conferencing system and five Lenovo desktop computers were working but not engraved.",
            "Gap: Working equipment, including an operational conferencing system, still required clear identification and custody, and the location of 16 assets was being traced. Recommended action: The ministry should engrave the colour photocopier and desktop computers, identify the conferencing system and its components, assign a custodian, and apply a suitable durable identifier without damaging the equipment. It should also complete the tracing of the 12 desktop computers and 4 UPS units and confirm their custodians.",
        ],
    ),
    (
        "4.2.7 Ministry of Education and Sports",
        [
            "MoES held 10,783 assets, all ICT equipment: 9,324 Teacher Effectiveness and Learners Assessment (TELA) handsets, 1,001 tablets (976 of them inspection tablets), 456 laptops, a photocopier and a printer. TELA is the system, financed by UgIFT, that the ministry uses to monitor attendance and performance in schools; the handsets and inspection tablets are deployed in local governments and remain MoES assets. The recorded value is UGX 14,925,128,359, the highest of any MDA. Accumulated depreciation to 30 September 2026 is UGX 10,607,476,704, and the net book value is UGX 4,317,651,655.",
            "All 10,783 assets (100.0%) were in good condition and in use, and none was broken. Of the movable assets, 458 (4.2%) were engraved, all with the UgIFT marking, and 10,325 (95.8%) were not engraved: none of the handsets or tablets carries an engraving. All 456 laptops were in good condition, in use and properly engraved.",
            "Recommended action: The ministry should link each handset and tablet to its serial number and to the school or office that holds it, and apply a durable identifier where the device allows.",
        ],
    ),
    (
        "4.2.8 Ministry of Agriculture, Animal Industry and Fisheries",
        [
            "MAAIF held 2,045 assets: 1,206 ICT items, 837 other assets and 2 transport assets. The ICT items were mainly 819 tablets, 211 computers (162 desktops and 49 laptops), 149 monitors and 15 printers. All 2,045 assets (100.0%) were in good condition and in use, and none was broken. Of the 1,227 movable assets, 1,204 (98.1%) were engraved, 1,200 (97.8%) with the UgIFT marking, and 23 (1.9%) were not engraved, among them two drones, a handheld global positioning system (GPS) unit and a projector. The recorded value is UGX 2,332,681,629, accumulated depreciation to 30 September 2026 is UGX 984,457,129, and the net book value is UGX 1,348,224,500.",
            "All 211 computers and 819 tablets were in good condition, in use and properly engraved. The ministry had working laptops and printers carrying the UgIFT engraving. Keeping those engravings linked to the equipment and its custodian will support continued accountability after programme closure. Routine condition checks should identify new faults and keep usable equipment in service.",
        ],
    ),
    (
        "4.2.9 Ministry of Health",
        [
            "MoH held 264 assets: 200 other assets, 39 ICT items and 25 transport assets. Of these, 263 (99.6%) were in good condition and in use, and 1 (0.4%) was broken. The broken asset was a double cabin pickup that had been boarded off and was not in use. Of the 73 movable assets, 41 (56.2%) were engraved and 32 (43.8%) were not engraved; 14 (19.2%) were properly engraved with the UgIFT marking and 27 (37.0%) were not properly engraved, carrying another marking only: a vehicle registration number (16) or an office asset code (11). The recorded value is UGX 1,216,859,097, accumulated depreciation to 30 September 2026 is UGX 463,977,008, and the net book value is UGX 752,882,089.",
            "The ministry's 26 computers (2 desktops and 24 laptops) were all in good condition and in use; 17 (65.4%) were engraved and 9 (34.6%) were not engraved.",
            "Expected: Office equipment in use should be identifiable by a durable engraving. Found: Two heavy duty printers, one at the Industrial Area engineering office and one at headquarters, were in use and in good condition but not engraved. Nine laptops were also not engraved.",
            "Gap: Working equipment at separate offices still lacked the engraving needed for straightforward physical identification. Recommended action: The ministry should engrave these items, link the engravings to their serial numbers and office locations, and obtain custody confirmation from the receiving units.",
        ],
    ),
    (
        "4.2.10 Office of the Prime Minister",
        [
            "OPM held 10 ICT items, all HP Envy laptops. Seven (70.0%) were in good condition and in use, and three (30.0%) were broken. The three broken assets were damaged laptops and were not in use. All 10 laptops (100.0%) were properly engraved with the UgIFT marking. The recorded value is UGX 53,535,300, accumulated depreciation to 30 September 2026 is UGX 48,181,770, and the net book value is UGX 5,353,530.",
            "The office should assess repair needs for the three damaged laptops and return serviceable equipment to use.",
            "Expected: Computers should have enough capacity for the work assigned to their users. Found: Users reported that the capacity of the HP Envy i3 laptops was below the volume of work handled, such as the performance assessment of local governments, although the machines were described as being in fair condition.",
            "Gap: A usable laptop could still be poorly matched to the workload of its user. Recommended action: The office should assess the requirements of the affected users and decide whether upgrading, reallocating or replacing the laptops would provide suitable capacity.",
        ],
    ),
    (
        "4.2.11 Ministry of Water and Environment",
        [
            "MoWE held 2,939 assets: 2,767 ICT items and 172 other assets. The main items were 2,368 tablets, 2,344 of them Euron tablets, 188 computers (157 Dell OptiPlex all in one desktops and 31 laptops) with their keyboards and UPS units, and 16 real time kinematic GPS machines. All 2,939 assets (100.0%) were in good condition and in use, and none was broken. Of the 2,782 movable assets, 2,740 (98.5%) were engraved, 2,707 (97.3%) with the UgIFT marking, and 42 (1.5%) were not engraved. The recorded value is UGX 5,000,313,297, accumulated depreciation to 30 September 2026 is UGX 2,089,906,834, and the net book value is UGX 2,910,406,463.",
            "All 188 computers were in good condition and in use; 182 (96.8%) were engraved and 6 (3.2%) were not engraved.",
            "Expected: Portable tablets should carry an asset identifier and have an assigned custodian. Found: Ten Apple tablets were not engraved.",
            "Gap: Portable equipment required an identification method that would allow each unit to be traced to its user and location. Recommended action: The ministry should apply suitable durable identifiers, link them to serial numbers, and confirm custody before the next physical check.",
        ],
    ),
    (
        "4.2.12 Ministry of Gender, Labour and Social Development",
        [
            "MoGLSD held 10 ICT items: 8 laptops and 2 tablets. All 10 (100.0%) were in good condition and in use, and none was broken. Eight (80.0%) were properly engraved with the UgIFT marking and 2 (20.0%) were not engraved. The recorded value is UGX 46,180,000, accumulated depreciation to 30 September 2026 is UGX 13,660,000, and the net book value is UGX 32,520,000.",
            "Laptops serving Finance and Administration were working and engraved. Two tablets, one for Finance and Administration and the other for Youth and Children Affairs, were working but not engraved. The ministry should engrave the tablets and link their identifiers to the units and officers responsible for them.",
        ],
    ),
    (
        "4.2.13 National Environment Management Authority",
        [
            "NEMA held 6 ICT items: 4 laptops, a printer and a heavy duty photocopier. All 6 (100.0%) were in good condition and in use, and none was broken. One (16.7%) was engraved and 5 (83.3%) were not engraved, including all 4 laptops. The recorded value is UGX 46,824,409, accumulated depreciation to 30 September 2026 is UGX 32,841,273, and the net book value is UGX 13,983,136.",
            "Expected: Computers and shared office equipment should remain identifiable wherever they are used. Found: Four Lenovo laptops serving the executive office and environmental audit, together with a heavy duty printer, were working but not engraved. The authority was awaiting a MoFPED team to engrave them.",
            "Gap: The equipment could support work, but lacked a durable identifying mark. Recommended action: The authority should engrave the five items, record their serial numbers and custodians, and include them in regular physical checks.",
        ],
    ),
    (
        "4.2.14 Public Procurement and Disposal of Public Assets Authority",
        [
            "PPDA held 5 ICT items: 3 Lenovo ThinkPad laptops, a Samsung Galaxy tablet and an HP printer. All 5 (100.0%) were in good condition, in use and properly engraved with the UgIFT marking. None was broken. The recorded value is UGX 30,153,500, accumulated depreciation to 30 September 2026 is UGX 13,066,516, and the net book value is UGX 17,086,984.",
            "The laptops, tablet and printer were assigned to the legal unit and the strategy and planning unit. The authority should keep the engravings linked to current custodians and update custody whenever equipment moves between units.",
        ],
    ),
    (
        "4.2.15 Office of the Auditor General",
        [
            "OAG held 20 assets: 18 ICT items (16 laptops and 2 printers) and 2 pickup vehicles. All 20 (100.0%) were in good condition, in use and engraved, and none was broken. Six (30.0%) were properly engraved with the UgIFT marking and 14 (70.0%) were not properly engraved, carrying only the office's own asset code (12) or a vehicle registration number (2). Of the 16 laptops, 6 (37.5%) carried the UgIFT marking and 10 (62.5%) the office's asset code. The recorded value is UGX 388,773,680, accumulated depreciation to 30 September 2026 is UGX 111,591,545, and the net book value is UGX 277,182,135.",
            "The two pickups had changed registration numbers. The office should retain the link between the earlier and current registrations, the chassis numbers and the assigned custodians so that each vehicle remains traceable.",
        ],
    ),
    (
        "4.2.16 Ministry of Local Government",
        [
            "MoLG held 13 assets: 12 ICT items and 1 transport asset, a pickup. The ICT items were 6 computers (2 desktops and 4 laptops), 4 computer peripherals, a copier and a heavy duty photocopier. All 13 (100.0%) were in good condition and in use, and none was broken. Eight (61.5%) were engraved, 7 (53.8%) with the UgIFT marking, and 5 (38.5%) were not engraved; of the 6 computers, 3 (50.0%) were engraved and 3 (50.0%) were not. The ministry keeps a well organised UgIFT asset register. The recorded value is UGX 343,320,450, accumulated depreciation to 30 September 2026 is UGX 284,804,160, and the net book value is UGX 58,516,290.",
            "The District Administration unit had a working pickup, photocopier and computer equipment. Two laptops and a desktop set, including its keyboard and mouse, were not engraved. The ministry should engrave this equipment and confirm its location and custodian while keeping the working assets in service.",
        ],
    ),
    (
        "4.2.17 Ministry of Lands, Housing and Urban Development",
        [
            "MoLHUD held 27 assets: 25 Honda motorcycles and 2 other assets. All 27 (100.0%) were in good condition and in use, and none was broken. Of the 26 movable assets, 13 (50.0%) were not properly engraved, carrying only their registration numbers, and 13 (50.0%) were not engraved; none carried the UgIFT marking. The recorded value is UGX 829,749,997, accumulated depreciation to 30 September 2026 is UGX 124,320,000, and the net book value is UGX 705,429,997.",
            "The ministry received survey equipment and motorcycles that it gave out to local governments, and it provided a list of the beneficiaries. The ministry should keep the beneficiary list current and apply the UgIFT engraving so that each motorcycle can be traced to the receiving local government.",
        ],
    ),
]


def main():
    document = Document(PATH)
    anchor = None
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name == "Heading 2" and paragraph.text.startswith("4.3 "):
            anchor = paragraph
            break
    if anchor is None:
        raise SystemExit("anchor 4.3 not found")

    for heading, paragraphs in SECTIONS:
        anchor.insert_paragraph_before(heading, style="Heading 3")
        for text in paragraphs:
            anchor.insert_paragraph_before(text, style="Normal")

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("HC\t"):
            paragraph.insert_paragraph_before("GPS\tGlobal positioning system", style=paragraph.style)
            break

    document.save(PATH)
    print("saved", sum(1 + len(p) for _, p in SECTIONS), "paragraphs")


if __name__ == "__main__":
    main()
