from docx import Document
from docx.oxml.ns import qn

d = Document(r"D:\coding\ugift-data-analysis\outputs\narrative-report\UgIFT Asset Verification Report.docx")
# list entry whose visible text starts with Figure 46 and includes a page field
for i, p in enumerate(d.paragraphs):
    if p.text.startswith("Figure 46:") and "figure_46" in p._p.xml:
        open(r"D:\coding\ugift-data-analysis\tmp_lof_fig46.xml", "w", encoding="utf-8").write(p._p.xml)
        print("list", i, "len", len(p._p.xml))
        break
p = d.paragraphs[508]
open(r"D:\coding\ugift-data-analysis\tmp_cap_fig46.xml", "w", encoding="utf-8").write(p._p.xml)
print("cap", len(p._p.xml))
# contents entry 4.2
for i, p in enumerate(d.paragraphs):
    if p.text.startswith("4.2 Ministries"):
        open(r"D:\coding\ugift-data-analysis\tmp_toc_42.xml", "w", encoding="utf-8").write(p._p.xml)
        print("toc", i)
        break
