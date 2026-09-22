# -*- coding: utf-8 -*-
"""Inspect Tororo LG xlsx and Kyamwinula xls."""
from pathlib import Path

src = Path(r"d:\coding\apps\ugift\tmp\team13-wa-later")

try:
    import openpyxl
    wb = openpyxl.load_workbook(src / "TORORO DLG ASSET REGISTER FOR SEED SCHOOLS 25.26.xlsx", data_only=True)
    print("=== TORORO DLG XLSX sheets:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print("\n-- sheet", name, "dims", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
        for i, row in enumerate(ws.iter_rows(max_row=min(25, ws.max_row or 1), values_only=True), 1):
            vals = [str(c)[:40] if c is not None else "" for c in row[:12]]
            if any(vals):
                print(i, " | ".join(vals))
except Exception as e:
    print("xlsx error", type(e), e)

try:
    import xlrd
    book = xlrd.open_workbook(str(src / "KYAMWINULA inventory.xls"))
    print("\n=== KYAMWINULA XLS sheets:", book.sheet_names())
    for name in book.sheet_names():
        sh = book.sheet_by_name(name)
        print("\n--", name, "rows", sh.nrows, "cols", sh.ncols)
        for r in range(min(20, sh.nrows)):
            vals = [str(sh.cell_value(r, c))[:40] for c in range(min(10, sh.ncols))]
            if any(v.strip() for v in vals):
                print(r+1, " | ".join(vals))
except Exception as e:
    print("xls error", type(e), e)
    # try xlrd via pandas
    try:
        import pandas as pd
        xl = pd.ExcelFile(src / "KYAMWINULA inventory.xls")
        print("pandas sheets", xl.sheet_names)
        for name in xl.sheet_names:
            df = pd.read_excel(xl, name, header=None)
            print(name, df.shape)
            print(df.head(15).to_string())
    except Exception as e2:
        print("pandas also failed", e2)
