"""Build LG contacts PDF: CAO/Town Clerk, contact, and facilities under each LG."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    CondPageBreak,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import allocation  # noqa: E402
import facilities as fac_mod  # noqa: E402

OUT_PDF = ROOT / "final-UGiFT-report-karamojja-6-teams" / "LG-and-facility-leaders-contacts.pdf"
CAO_LIST = (
    ROOT
    / "facility-registers"
    / "team-13"
    / "_team-documents"
    / "CAOs-and-TCs-list-1st-July-2026.docx"
)
REGISTERS = ROOT / "facility-registers"

TEAM_LABELS = {
    10: "Team 10  -  Karamoja (south / central)",
    11: "Team 11  -  Karamoja (north)",
    12: "Team 12  -  Bukedi / Mbale",
    13: "Team 13  -  Bukedi / Bugisu border",
    14: "Team 14  -  Bugisu / Mt Elgon",
    15: "Team 15  -  Sebei",
}

HEADER_BG = colors.HexColor("#1F3A5F")
HEADER_MID = colors.HexColor("#2A4A73")
LEADER_BG = colors.HexColor("#E8EEF5")
FAC_HEAD_BG = colors.HexColor("#F0F3F7")
ROW_ALT = colors.HexColor("#F7F8FA")
GRID = colors.HexColor("#C5CDD8")
MUTED = colors.HexColor("#555555")
SCHOOL_TAG = colors.HexColor("#1F3A5F")
HEALTH_TAG = colors.HexColor("#1B5E4A")

DASHES = {
    "\u2010": "-",
    "\u2011": "-",
    "\u2012": "-",
    "\u2013": "-",
    "\u2014": "-",
    "\u2212": "-",
    "\u00a0": " ",
    "\ufffd": "-",
}


def ascii_hyphen(text: str) -> str:
    out = text or ""
    for src, dst in DASHES.items():
        out = out.replace(src, dst)
    return re.sub(r"[ \t]+", " ", out).strip()


def esc(text: str) -> str:
    return (
        ascii_hyphen(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def clean_phone(raw: str) -> str:
    raw = ascii_hyphen(raw)
    parts = [p for p in re.split(r"[\s/]+", raw) if p]
    return " / ".join(parts) if parts else "-"


def clean_email(raw: str) -> str:
    raw = ascii_hyphen(raw)
    parts = [p for p in re.split(r"\s+", raw) if p and "@" in p]
    return " / ".join(parts) if parts else "-"


def split_leader_name(raw: str, is_mc: bool) -> tuple[str, str]:
    name = ascii_hyphen(raw)
    if is_mc:
        name = re.sub(r"\s+Ag\.?\s*T\.?C\.?\s*$", "", name, flags=re.I).strip()
        return name, "Town Clerk"
    acting = bool(re.search(r"\bAg\.?\s*CAO\b|\bAg\.?\s*CAO\b|Ag\.CAO", name, re.I))
    name = re.sub(r"\s*[-]?\s*Ag\.?\s*CAO\.?\s*$", "", name, flags=re.I)
    name = re.sub(r"\s+Ag\.CAO\s*$", "", name, flags=re.I)
    name = name.strip(" -")
    return name, "CAO (Ag.)" if acting else "CAO"


def lg_display(lg: str) -> str:
    if lg.upper().endswith(" MC"):
        return f"{lg[:-3].strip()} Municipal Council"
    return f"{lg} District"


def folder_for(team: int, lg: str) -> Path | None:
    tdir = REGISTERS / f"team-{team}"
    if not tdir.is_dir():
        return None
    # Never strip " MC" down to the district name: that would attach district
    # facilities to the municipal council.
    candidates = [lg, lg.replace(" MC", " Mc"), lg.replace("MC", "Mc")]
    wanted = {c.lower() for c in candidates}
    for p in tdir.iterdir():
        if p.is_dir() and not p.name.startswith("_") and p.name.lower() in wanted:
            return p
    return None


def is_health(folder_name: str) -> bool:
    low = folder_name.lower()
    return bool(re.search(r"(?:^|-)hc(?:-|$)|health|hc-?i{1,3}$", low))


def facility_label(folder_name: str) -> str:
    name = folder_name.replace("-", " ")
    name = name.replace("St Johns", "St John's")
    name = name.replace("St Marys", "St Mary's")
    name = re.sub(r"\s+", " ", name).strip()
    return name


def facilities_from_folder(lg_dir: Path) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for p in sorted(lg_dir.iterdir(), key=lambda x: x.name.lower()):
        if not p.is_dir() or p.name.startswith("_"):
            continue
        kind = "Health centre" if is_health(p.name) else "School"
        rows.append((kind, facility_label(p.name)))
    rows.sort(key=lambda r: (0 if r[0] == "School" else 1, r[1].lower()))
    return rows


def facilities_from_programme(lg: str) -> list[tuple[str, str]]:
    key = lg.replace(" MC", " Mc") if lg.upper().endswith(" MC") else lg
    # facilities.py uses 'Tororo Mc' style
    alt = lg[:-3] + " Mc" if lg.upper().endswith(" MC") else lg
    schools = fac_mod.SCHOOLS.get(lg) or fac_mod.SCHOOLS.get(alt) or fac_mod.SCHOOLS.get(key) or []
    health = fac_mod.HEALTH.get(lg) or fac_mod.HEALTH.get(alt) or fac_mod.HEALTH.get(key) or []
    rows: list[tuple[str, str]] = []
    for item in schools:
        rows.append(("School", item[0]))
    for item in health:
        rows.append(("Health centre", item[0]))
    return rows


def load_molg() -> tuple[dict[str, tuple[str, str, str]], dict[str, tuple[str, str, str]]]:
    doc = Document(str(CAO_LIST))

    def rows(table):
        out = []
        for r in table.rows[1:]:
            cells = [ascii_hyphen(c.text) for c in r.cells]
            if not cells or len(set(cells)) == 1:
                continue
            if cells[1].lower() in ("cities", "municipalities", "district"):
                continue
            out.append(cells)
        return out

    cao = {}
    for r in rows(doc.tables[0]):
        if r[1]:
            cao[r[1]] = (r[2], clean_phone(r[3]), clean_email(r[4]))
    tc = {}
    for r in rows(doc.tables[1]):
        if r[1]:
            tc[r[1]] = (r[2], clean_phone(r[3]), clean_email(r[4]))
    return cao, tc


def make_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=3,
            textColor=HEADER_BG,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            alignment=TA_CENTER,
            spaceAfter=8,
            textColor=colors.HexColor("#333333"),
        ),
        "team": ParagraphStyle(
            "team",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            spaceBefore=8,
            spaceAfter=5,
            textColor=HEADER_BG,
        ),
        "lg": ParagraphStyle(
            "lg",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=11,
            textColor=colors.white,
        ),
        "lg_meta": ParagraphStyle(
            "lg_meta",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=10,
            alignment=TA_RIGHT,
            textColor=colors.white,
        ),
        "label": ParagraphStyle(
            "label",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=6.5,
            leading=8,
            textColor=MUTED,
        ),
        "value": ParagraphStyle(
            "value",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=HEADER_BG,
        ),
        "role": ParagraphStyle(
            "role",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=HEADER_MID,
        ),
        "fac_head": ParagraphStyle(
            "fac_head",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=10,
            textColor=HEADER_BG,
        ),
        "fac": ParagraphStyle(
            "fac",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10.5,
            textColor=colors.HexColor("#222222"),
        ),
        "kind_s": ParagraphStyle(
            "kind_s",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7,
            leading=9,
            textColor=SCHOOL_TAG,
        ),
        "kind_h": ParagraphStyle(
            "kind_h",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7,
            leading=9,
            textColor=HEALTH_TAG,
        ),
        "note": ParagraphStyle(
            "note",
            parent=base["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=7.5,
            leading=10,
            textColor=MUTED,
        ),
        "footer": ParagraphStyle(
            "footer",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            textColor=MUTED,
            spaceBefore=6,
        ),
    }


def build_card(styles, number: int, team: int, lg: str, leader, facilities, note: str, card_w: float):
    name, role, phone, email = leader
    n_fac = len(facilities)
    fac_count = f"{n_fac} facilit{'y' if n_fac == 1 else 'ies'}" if n_fac else "No facilities on file"

    header = Table(
        [[
            Paragraph(esc(f"{number}.  {lg_display(lg)}"), styles["lg"]),
            Paragraph(esc(f"Team {team}  |  {fac_count}"), styles["lg_meta"]),
        ]],
        colWidths=[card_w * 0.62, card_w * 0.38],
    )
    header.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), HEADER_BG),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )

    leader_tbl = Table(
        [[
            Paragraph("LEADER", styles["label"]),
            Paragraph("TELEPHONE", styles["label"]),
            Paragraph("EMAIL", styles["label"]),
        ], [
            Paragraph(f"{esc(name)}<br/><font size='7.5'>{esc(role)}</font>", styles["value"]),
            Paragraph(esc(phone), styles["role"]),
            Paragraph(esc(email), styles["role"]),
        ]],
        colWidths=[card_w * 0.40, card_w * 0.28, card_w * 0.32],
    )
    leader_tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), LEADER_BG),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), ( -1, 0), 4),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 0),
                ("TOPPADDING", (0, 1), (-1, 1), 1),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 5),
            ]
        )
    )

    fac_head = Table(
        [[Paragraph("Facilities under this local government", styles["fac_head"])]],
        colWidths=[card_w],
    )
    fac_head.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), FAC_HEAD_BG),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )

    if facilities:
        fac_data = []
        for i, (kind, label) in enumerate(facilities):
            kstyle = styles["kind_s"] if kind == "School" else styles["kind_h"]
            fac_data.append([
                Paragraph(esc(kind), kstyle),
                Paragraph(esc(label), styles["fac"]),
            ])
        fac_tbl = Table(fac_data, colWidths=[card_w * 0.22, card_w * 0.78])
        cmds = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (0, -1), 6),
            ("LEFTPADDING", (1, 0), (1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ("LINEBELOW", (0, 0), (-1, -2), 0.25, GRID),
        ]
        for i in range(len(fac_data)):
            if i % 2 == 1:
                cmds.append(("BACKGROUND", (0, i), (-1, i), ROW_ALT))
        fac_tbl.setStyle(TableStyle(cmds))
        body = fac_tbl
    else:
        body = Table(
            [[Paragraph(esc(note or "None recorded."), styles["note"])]],
            colWidths=[card_w],
        )
        body.setStyle(
            TableStyle(
                [
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )

    note_tbl = None
    if note and facilities:
        note_tbl = Table(
            [[Paragraph(esc(note), styles["note"])]],
            colWidths=[card_w],
        )
        note_tbl.setStyle(
            TableStyle(
                [
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FBF6EE")),
                ]
            )
        )

    inner = [header, leader_tbl, fac_head]
    if note_tbl is not None:
        inner.append(note_tbl)
    inner.append(body)
    card = Table([[inner_el] for inner_el in inner], colWidths=[card_w])
    card.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.6, HEADER_BG),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return KeepTogether([card, Spacer(1, 3.2 * mm)])


def collect_entries(cao_map, tc_map):
    entries = []
    for team in sorted(allocation.TEAMS):
        _members, _pay, _sup, rows = allocation.TEAMS[team]
        for lg, n_sch, n_hc in rows:
            is_mc = lg.upper().endswith(" MC")
            base = lg[:-3].strip() if is_mc else lg
            if is_mc:
                raw_name, phone, email = tc_map.get(base, ("", "-", "-"))
            else:
                raw_name, phone, email = cao_map.get(base, ("", "-", "-"))
            name, role = split_leader_name(raw_name or "-", is_mc)
            if not raw_name:
                name = "-"

            lg_dir = folder_for(team, lg)
            note = ""
            if lg_dir is not None:
                facs = facilities_from_folder(lg_dir)
            else:
                facs = facilities_from_programme(lg)
                if n_sch + n_hc == 0:
                    facs = []
                    note = "No UgIFT facilities were allocated to this local government."
                elif facs:
                    note = ""
                else:
                    note = "No facility return on file yet."

            entries.append(
                {
                    "team": team,
                    "lg": lg,
                    "leader": (name, role, phone, email),
                    "facilities": facs,
                    "note": note,
                }
            )
    return entries


def add_page_number(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawCentredString(
        A4[0] / 2,
        8 * mm,
        f"UgIFT teams 10-15  |  CAO and facilities by local government  |  page {canvas.getPageNumber()}",
    )
    canvas.restoreState()


def main():
    styles = make_styles()
    cao_map, tc_map = load_molg()
    entries = collect_entries(cao_map, tc_map)

    page_w, _page_h = A4
    margin = 12 * mm
    usable = page_w - 2 * margin

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title="UgIFT teams 10-15 - CAO and facilities by local government",
        author="UgIFT field verification",
    )

    story = []
    story.append(
        Paragraph(
            "UgIFT teams 10-15 - CAO and facilities by local government",
            styles["title"],
        )
    )
    story.append(
        Paragraph(
            "Each local government shows the Chief Administrative Officer or, for a municipal "
            "council, the Town Clerk, with telephone and email from the Ministry of Local "
            "Government deployment list (1 July 2026). Facilities listed are those filed under "
            "that local government in the field registers. Supervisor: Kassim Luminsa, "
            "+256 702 806116.",
            styles["subtitle"],
        )
    )

    current_team = None
    for i, e in enumerate(entries, start=1):
        card = build_card(
            styles,
            i,
            e["team"],
            e["lg"],
            e["leader"],
            e["facilities"],
            e["note"],
            usable,
        )
        if e["team"] != current_team:
            current_team = e["team"]
            if current_team != min(TEAM_LABELS):
                # Keep a team heading from sitting alone at the foot of a page.
                story.append(CondPageBreak(48 * mm))
            story.append(Paragraph(TEAM_LABELS[current_team], styles["team"]))
        story.append(card)

    story.append(
        Paragraph(
            "Sources: MoLG CAOs and TCs list 1 July 2026; facility-registers folders for teams "
            "10-15; programme schools and health-centre list where a local government has no "
            "facility return on file.",
            styles["footer"],
        )
    )

    doc.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)
    print(f"Wrote {OUT_PDF}")
    print(f"Local governments: {len(entries)}")
    missing = [e["lg"] for e in entries if e["leader"][0] == "-"]
    if missing:
        print("Missing leaders:", missing)


if __name__ == "__main__":
    main()
