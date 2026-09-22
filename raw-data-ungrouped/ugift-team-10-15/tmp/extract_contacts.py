"""Extract LG and facility leader contacts for teams 10-15."""
from __future__ import annotations

import re
from pathlib import Path
from collections import defaultdict

from docx import Document

ROOT = Path("facility-registers")
PHONE_RE = re.compile(
    r"(?:\+?256[\s\-]*)?0?7\d[\d\s\-/]{6,12}\d",
    re.I,
)
ROLE_HINT = re.compile(
    r"head\s*teacher|deputy\s*head|in[-\s]?charge|person\s+in\s+charge|"
    r"\bDOS\b|Director of Studies|informant|CAO|Town Clerk|DEO|DHO|CFO|"
    r"District Education|District Health|Chief Administrative|"
    r"welcomed by|met with|attended by|received by|spoke with",
    re.I,
)


def doc_text(path: Path) -> str:
    try:
        doc = Document(path)
    except Exception as e:
        return f""
    parts: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)
    return "\n".join(parts)


def clean_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    if digits.startswith("256") and len(digits) >= 12:
        digits = "0" + digits[3:]
    if len(digits) == 9 and digits.startswith("7"):
        digits = "0" + digits
    if len(digits) >= 10:
        return digits[:10]
    return raw.strip()


def interesting_lines(text: str) -> list[str]:
    out = []
    for line in text.splitlines():
        line = " ".join(line.split())
        if not line:
            continue
        if PHONE_RE.search(line) or ROLE_HINT.search(line):
            if len(line) > 400:
                line = line[:400] + "…"
            out.append(line)
    return out


# --- Sample patterns from a few docs ---
print("=== SAMPLE HITS ===")
samples = []
for team in range(10, 16):
    samples += list((ROOT / f"team-{team}").rglob("*LG-courtesy-call.docx"))[:1]
    samples += list((ROOT / f"team-{team}").rglob("*-school-*.docx"))[:1]
    samples += list((ROOT / f"team-{team}").rglob("*-health-centre-*.docx"))[:1]
    samples += list((ROOT / f"team-{team}").rglob("*team-field-note.docx"))[:1]

for s in samples[:18]:
    text = doc_text(s)
    hits = interesting_lines(text)
    if not hits:
        continue
    print(f"\n## {s.as_posix()}")
    for h in hits[:12]:
        print(f"  {h}")

# --- Full pass: collect facility report hits with phones ---
print("\n\n=== FACILITY DOCS WITH PHONES ===")
for team in range(10, 16):
    team_dir = ROOT / f"team-{team}"
    print(f"\n# Team {team}")
    for path in sorted(team_dir.rglob("*.docx")):
        if path.name.startswith("~"):
            continue
        # skip process reports / team docs that aren't facility or LG
        rel = path.relative_to(team_dir).as_posix()
        if "_team-documents" in rel and "field-note" not in path.name.lower():
            continue
        text = doc_text(path)
        phones = [clean_phone(m.group(0)) for m in PHONE_RE.finditer(text)]
        # skip verifier/supervisor phones noise somewhat
        role_lines = [
            ln
            for ln in interesting_lines(text)
            if PHONE_RE.search(ln)
            or re.search(
                r"head\s*teacher|in[-\s]?charge|informant|deputy|DOS|"
                r"CAO|Town Clerk|DEO|DHO|CFO|welcomed|met with",
                ln,
                re.I,
            )
        ]
        if not role_lines:
            continue
        # only print if looks like facility/LG leader (has phone near role OR identification)
        keep = []
        for ln in role_lines:
            if re.search(
                r"head\s*teacher|in[-\s]?charge|informant|deputy|DOS|"
                r"CAO|Town Clerk|DEO|DHO|CFO|welcomed|met with|received by|"
                r"person in charge|Incharge",
                ln,
                re.I,
            ):
                keep.append(ln)
            elif PHONE_RE.search(ln) and re.search(
                r"Mr\.|Ms\.|Mrs\.|Dr\.|Miss ", ln
            ):
                keep.append(ln)
        if keep:
            print(f"\n{rel}")
            for ln in keep[:15]:
                print(f"  {ln}")
