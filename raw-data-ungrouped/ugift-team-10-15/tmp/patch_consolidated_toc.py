# -*- coding: utf-8 -*-
"""Number team/LG headings, replace TOC, add page numbers; then Word-export PDF."""
from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
import build_consolidated_final_report as B
import paths

DOCX = os.path.join(paths.ROOT, "final-UGiFT-report-karamojja-6-teams",
                    "UGiFT-consolidated-field-report.docx")
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
SIZES = {1: 16, 2: 14, 3: 12, 4: 11}


def set_heading(para, text, level):
    para.style = para.part.document.styles["Heading %d" % level]
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
        run = para.runs[0]
    else:
        run = para.add_run(text)
    B.set_run_font(run, size=SIZES[level], bold=True, color=NAVY)
    para.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    para.paragraph_format.space_after = Pt(8)


def insert_before(paragraph, text="", *, size=11, bold=False, color=None,
                  page_break=False, toc=False):
    new_p = OxmlElement("w:p")
    paragraph._element.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if page_break:
        new_para.add_run().add_break(WD_BREAK.PAGE)
        return new_para
    run = new_para.add_run(text) if text else new_para.add_run()
    if text:
        B.set_run_font(run, size=size, bold=bold, color=color)
        new_para.paragraph_format.space_after = Pt(12 if bold else 6)
    if toc:
        B._add_field(run, r' TOC \o "1-3" \h \z \u ')
    return new_para


def main():
    doc = Document(DOCX)
    team_labels, lg_labels = B.field_heading_labels()
    plain_team = {}
    for team, region, _lgs in B.SPINE:
        plain_team["Team %d - %s" % (team, region)] = team_labels[team]
        # already numbered headings stay idempotent
        plain_team[team_labels[team]] = team_labels[team]
    plain_lg = {}
    for (team, lg), label in lg_labels.items():
        plain_lg[B.lg_title(lg)] = label
        plain_lg[label] = label

    n_team = n_lg = n_fac = 0
    for para in doc.paragraphs:
        style = para.style.name if para.style else ""
        text = para.text.strip()
        if style == "Heading 1" and text in plain_team:
            set_heading(para, plain_team[text], 2)
            n_team += 1
        elif style == "Heading 2" and text in plain_lg:
            set_heading(para, plain_lg[text], 3)
            n_lg += 1
        elif style == "Heading 3" and " (" in text:
            set_heading(para, text, 4)
            n_fac += 1
        elif style == "Heading 2" and text.startswith("4.") and "Team " in text:
            set_heading(para, text, 2)
            n_team += 1
        elif style == "Heading 3" and text[:2] == "4.":
            set_heading(para, text, 3)
            n_lg += 1

    print("relabelled teams", n_team, "lgs", n_lg, "facilities", n_fac)

    intro = None
    for para in doc.paragraphs:
        if para.text.strip() == "1. Introduction" and (para.style.name or "").startswith("Heading"):
            intro = para
            break
    if intro is None:
        raise SystemExit("could not find Introduction heading")

    started = False
    removed = 0
    for para in list(doc.paragraphs):
        text = para.text.strip()
        if text == "Table of contents":
            started = True
        if not started:
            continue
        if para._element is intro._element:
            break
        para._element.getparent().remove(para._element)
        removed += 1
    print("removed old TOC paragraphs", removed)

    insert_before(intro, "Table of contents", size=16, bold=True, color=NAVY)
    insert_before(intro, toc=True)
    insert_before(intro, page_break=True)

    B.add_footer_page_numbers(doc)
    doc.save(DOCX)
    print("saved", DOCX)

    pdf = B.export_pdf(DOCX)
    print("pdf", pdf)


if __name__ == "__main__":
    main()
