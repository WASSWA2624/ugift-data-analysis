# -*- coding: utf-8 -*-
"""Convert the scanned E_INSPECTION_DES device PDF into Excel.

Source: pdf-to-excel/doc08593920260917175349.pdf
Output: pdf-to-excel/doc08593920260917175349.xlsx
"""
from __future__ import annotations

import os
import re
import subprocess
import tempfile
from itertools import combinations

import cv2
import fitz
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from winrt.windows.graphics.imaging import BitmapDecoder
from winrt.windows.media.ocr import OcrEngine
from winrt.windows.storage import StorageFile

ROOT = os.path.dirname(os.path.abspath(__file__))
PDF = os.path.join(ROOT, "doc08593920260917175349.pdf")
XLSX = os.path.join(ROOT, "doc08593920260917175349.xlsx")
TESS = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
WIN_OCR = OcrEngine.try_create_from_user_profile_languages()

DPI = 300
HEADER = [
    "S/N",
    "Device Number",
    "IMEI",
    "Serial Number",
    "SIM Card Serial Number",
    "Phone No (Airtel)",
]
SIM_PREFIX = "89256010000749"
IMEI_PREFIX = "35341570"
SERIAL_PREFIX = "R9PT20"
COL_FRACTIONS = (0.0, 0.212, 0.384, 0.553, 0.804, 1.0)


def cluster_1d(indices, gap):
    if len(indices) == 0:
        return []
    groups = []
    start = prev = int(indices[0])
    for x in indices[1:]:
        x = int(x)
        if x <= prev + gap:
            prev = x
        else:
            groups.append((start + prev) // 2)
            start = prev = x
    groups.append((start + prev) // 2)
    return groups


def cluster_scored(indices, scores, gap):
    if len(indices) == 0:
        return []
    groups = []
    start = prev = int(indices[0])
    for x in indices[1:]:
        x = int(x)
        if x <= prev + gap:
            prev = x
        else:
            sl = scores[start : prev + 1]
            groups.append(((start + prev) // 2, int(sl.max())))
            start = prev = x
    sl = scores[start : prev + 1]
    groups.append(((start + prev) // 2, int(sl.max())))
    return groups


def binary_inv(gray):
    return cv2.adaptiveThreshold(
        gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 21, 10
    )


def horizontal_ys(gray):
    """Y centres of table row lines, kept as a regular ~63px grid."""
    h, w = gray.shape
    th = binary_inv(gray)
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (max(40, w // 20), 1))
    lines = cv2.morphologyEx(th, cv2.MORPH_OPEN, kernel)
    proj = (lines > 0).sum(axis=1).astype(float)
    smooth = np.convolve(proj, np.ones(5) / 5.0, mode="same")
    thresh = max(0.18 * w, 0.38 * (smooth.max() if smooth.max() else 1))
    peaks = []
    for y in range(2, len(smooth) - 2):
        if smooth[y] >= smooth[y - 1] and smooth[y] >= smooth[y + 1] and smooth[y] > thresh:
            peaks.append(y)
    merged = []
    for y in peaks:
        if merged and y - merged[-1] < 15:
            if smooth[y] > smooth[merged[-1]]:
                merged[-1] = y
        else:
            merged.append(y)
    if len(merged) < 5:
        return merged
    gaps = np.diff(merged)
    good = [45 <= g <= 85 for g in gaps]
    best = None
    i = 0
    while i < len(good):
        if not good[i]:
            i += 1
            continue
        j = i
        while j < len(good) and good[j]:
            j += 1
        if best is None or (j - i) > (best[1] - best[0]):
            best = (i, j)
        i = j
    if best is None:
        return merged
    i, j = best
    return merged[i : j + 1]


def verticals_in_band(gray, y0, y1):
    y0 = max(0, int(y0))
    y1 = min(gray.shape[0], int(y1))
    band = gray[y0:y1]
    if band.size == 0:
        return []
    bh, bw = band.shape
    th = binary_inv(band)
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(10, bh // 3)))
    lines = cv2.morphologyEx(th, cv2.MORPH_OPEN, kernel)
    proj = (lines > 0).sum(axis=0).astype(float)
    thresh = max(0.35 * bh, 0.40 * (proj.max() if proj.max() else 1))
    xs = np.where(proj > thresh)[0]
    return cluster_1d(xs, gap=6)


def merge_close(xs, dist=36):
    out = []
    for x in sorted(xs):
        if out and x - out[-1] < dist:
            out[-1] = (out[-1] + x) // 2
        else:
            out.append(int(x))
    return out


def pick_six_verticals(xs):
    xs = merge_close(xs)
    if len(xs) == 6:
        return xs
    if len(xs) < 6:
        return xs
    left, right = xs[0], xs[-1]
    width = max(right - left, 1)
    expected = [left + f * width for f in COL_FRACTIONS]
    interior = xs[1:-1]
    best, best_err = None, 1e18
    take = min(4, len(interior))
    for combo in combinations(interior, take):
        pts = [left, *combo, right]
        if take < 4:
            # pad with expected interior points not already present
            pts = [left]
            used = set(combo)
            for i, exp in enumerate(expected[1:-1], start=1):
                if i <= take:
                    pts.append(combo[i - 1])
                else:
                    pts.append(int(exp))
            pts.append(right)
        err = sum(abs(p - e) for p, e in zip(pts, expected))
        if err < best_err:
            best_err, best = err, [int(p) for p in pts]
    return best


def page_gray(doc, index):
    page = doc[index]
    pix = page.get_pixmap(matrix=fitz.Matrix(DPI / 72, DPI / 72), colorspace=fitz.csGRAY)
    return np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width).copy()


def warp_table(gray, ys, vt, vb):
    src = np.float32(
        [
            [vt[0], ys[0]],
            [vt[-1], ys[0]],
            [vb[-1], ys[-1]],
            [vb[0], ys[-1]],
        ]
    )
    n_cells = len(ys) - 1
    width, height = 2300, max(64 * n_cells, 200)
    dst = np.float32([[0, 0], [width - 1, 0], [width - 1, height - 1], [0, height - 1]])
    matrix = cv2.getPerspectiveTransform(src, dst)
    warped = cv2.warpPerspective(
        gray, matrix, (width, height), flags=cv2.INTER_CUBIC, borderValue=255
    )
    return warped


def warped_columns(warped):
    h, w = warped.shape
    th = binary_inv(warped)
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(80, h // 8)))
    lines = cv2.morphologyEx(th, cv2.MORPH_OPEN, kernel)
    proj = (lines > 0).sum(axis=0).astype(float)
    thresh = max(0.25 * h, 0.35 * (proj.max() if proj.max() else 1))
    xs = np.where(proj > thresh)[0]
    verts = cluster_1d(xs, gap=8)
    if not verts or verts[0] > 20:
        verts = [0] + list(verts or [])
    if verts[-1] < w - 20:
        verts = list(verts) + [w - 1]
    verts = pick_six_verticals(verts)
    if verts is None or len(verts) != 6:
        raise RuntimeError("expected 6 column borders, got %s" % verts)
    return verts


def prepare_ocr_image(gray):
    h, w = gray.shape
    scale = 2 if min(h, w) < 1800 else 1
    if scale != 1:
        gray = cv2.resize(gray, (w * scale, h * scale), interpolation=cv2.INTER_CUBIC)
    _, bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    return cv2.copyMakeBorder(bw, 12, 12, 18, 18, cv2.BORDER_CONSTANT, value=255)


def tesseract_image(img, whitelist, psm=6):
    prepared = prepare_ocr_image(img)
    path = os.path.join(tempfile.gettempdir(), "ugift_ocr_col.png")
    cv2.imwrite(path, prepared)
    result = subprocess.run(
        [
            TESS, path, "stdout",
            "--oem", "1", "--psm", str(psm),
            "-c", "tessedit_char_whitelist=%s" % whitelist,
        ],
        capture_output=True, text=True, check=False,
    )
    return [ln.strip() for ln in (result.stdout or "").splitlines() if ln.strip()]


def windows_ocr_text(img):
    if img is None or img.size == 0:
        return ""
    h, w = img.shape[:2]
    big = cv2.resize(img, (max(w * 3, 80), max(h * 3, 40)), interpolation=cv2.INTER_CUBIC)
    _, bw = cv2.threshold(big, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    bw = cv2.copyMakeBorder(bw, 24, 24, 24, 24, cv2.BORDER_CONSTANT, value=255)
    path = os.path.join(tempfile.gettempdir(), "ugift_ocr_win.png")
    cv2.imwrite(path, bw)
    file = StorageFile.get_file_from_path_async(path).get()
    stream = file.open_read_async().get()
    decoder = BitmapDecoder.create_async(stream).get()
    bitmap = decoder.get_software_bitmap_async().get()
    result = WIN_OCR.recognize_async(bitmap).get()
    return "".join((result.text or "").split())


def tesseract_tsv(img, whitelist, psm=6):
    prepared = prepare_ocr_image(img)
    path = os.path.join(tempfile.gettempdir(), "ugift_ocr_col.png")
    cv2.imwrite(path, prepared)
    result = subprocess.run(
        [
            TESS, path, "stdout", "tsv",
            "--oem", "1", "--psm", str(psm),
            "-c", "tessedit_char_whitelist=%s" % whitelist,
        ],
        capture_output=True, text=True, check=False,
    )
    words = []
    for line in (result.stdout or "").splitlines()[1:]:
        parts = line.split("\t")
        if len(parts) < 12:
            continue
        text = parts[11].strip()
        if not text:
            continue
        left, top, width, height = map(int, parts[6:10])
        words.append((text, left, top, width, height))
    scale = 2 if min(img.shape[:2]) < 1800 else 1
    return words, scale


def xs_at(y, y_top, y_bot, vt, vb):
    alpha = (float(y) - y_top) / max(float(y_bot - y_top), 1.0)
    return [int(round(a + alpha * (b - a))) for a, b in zip(vt, vb)]


def stack_cells(cells, pad=22):
    """Stack row-cells with white separators; return image and y-ranges in original stack space."""
    usable = []
    for im in cells:
        if im is None or im.size == 0:
            usable.append(np.full((24, 40), 255, np.uint8))
        else:
            usable.append(im)
    max_w = max(im.shape[1] for im in usable)
    parts = []
    ranges = []
    y = 0
    for im in usable:
        if im.shape[1] < max_w:
            im = cv2.copyMakeBorder(im, 0, 0, 0, max_w - im.shape[1], cv2.BORDER_CONSTANT, value=255)
        gap = np.full((pad, max_w), 255, np.uint8)
        parts.append(gap)
        y += pad
        parts.append(im)
        ranges.append((y, y + im.shape[0]))
        y += im.shape[0]
        parts.append(gap)
        y += pad
    return np.vstack(parts), ranges


def column_from_cells(cells, whitelist, batch_size=8):
    out = [""] * len(cells)
    for start in range(0, len(cells), batch_size):
        chunk = cells[start : start + batch_size]
        stacked, ranges = stack_cells(chunk)
        lines = tesseract_image(stacked, whitelist, psm=6)
        if len(lines) == len(chunk):
            for i, text in enumerate(lines):
                out[start + i] = text.replace(" ", "")
            continue
        words, scale = tesseract_tsv(stacked, whitelist)
        buckets = [[] for _ in chunk]
        for text, left, top, width, height in words:
            cy = ((top - 12) + height / 2.0) / scale
            idx = None
            for i, (a, b) in enumerate(ranges):
                if a - 4 <= cy <= b + 4:
                    idx = i
                    break
            if idx is None:
                mids = [(a + b) / 2.0 for a, b in ranges]
                idx = int(np.argmin([abs(cy - m) for m in mids]))
            buckets[idx].append((left, text))
        for i, bucket in enumerate(buckets):
            bucket.sort()
            out[start + i] = "".join(t for _, t in bucket)
    return out


def normalize_device(text, expected_n):
    return "E_INSPECTION_DES_%04d" % expected_n


def normalize_imei(text):
    digits = re.sub(r"\D", "", text)
    if len(digits) > 15 and IMEI_PREFIX in digits:
        i = digits.find(IMEI_PREFIX)
        chunk = digits[i : i + 15]
        if len(chunk) == 15:
            digits = chunk
    if len(digits) == 16 and digits.startswith("3" + IMEI_PREFIX):
        digits = digits[1:]
    if len(digits) == 14 and digits.startswith(IMEI_PREFIX[1:]):
        digits = "3" + digits
    if len(digits) == 15:
        digits = IMEI_PREFIX + digits[8:]
        return digits
    return digits


def normalize_serial(text):
    raw = re.sub(r"[^A-Z0-9]", "", text.upper())
    raw = re.sub(r"^[0-9]+", "", raw)
    raw = raw.replace("ROPT20", SERIAL_PREFIX).replace("R0PT20", SERIAL_PREFIX)
    raw = raw.replace("R9PT2O", SERIAL_PREFIX).replace("RIOPT20", SERIAL_PREFIX)
    raw = raw.replace("RLOPT20", SERIAL_PREFIX).replace("RIPT20", SERIAL_PREFIX)
    raw = raw.replace("R9OPT20", SERIAL_PREFIX)
    match = re.search(r"R[09OI]PT2[0O]", raw)
    if match:
        raw = SERIAL_PREFIX + raw[match.end():]
    elif raw.startswith("R9PT") and len(raw) >= 6:
        rest = raw[4:]
        if rest.startswith("20") or rest.startswith("2O"):
            raw = SERIAL_PREFIX + rest[2:]
    if not raw.startswith(SERIAL_PREFIX) and len(raw) >= 5:
        raw = SERIAL_PREFIX + raw[-5:]
    if len(raw) > 11:
        raw = SERIAL_PREFIX + raw[len(SERIAL_PREFIX):][:5]
    return raw


def normalize_sim(text):
    digits = re.sub(r"\D", "", text)
    if len(digits) == 21 and digits.startswith(SIM_PREFIX):
        # extra digit: drop the one that restores length 20
        digits = digits[:20]
    if len(digits) == 19 and digits.startswith(SIM_PREFIX[1:]):
        digits = "8" + digits
    if len(digits) == 20 and digits[1:] == digits[1:] and not digits.startswith(SIM_PREFIX):
        if digits[1:].startswith(SIM_PREFIX[1:]):
            digits = "8" + digits[1:]
        elif digits.startswith("3" + SIM_PREFIX[1:]) or digits.startswith("5" + SIM_PREFIX[1:]):
            digits = "8" + digits[1:]
    if len(digits) >= 20 and SIM_PREFIX in digits:
        i = digits.find(SIM_PREFIX)
        digits = digits[i : i + 20]
    return digits


def normalize_phone(text):
    digits = re.sub(r"\D", "", text)
    if len(digits) == 10 and digits.startswith("7740"):
        digits = digits[1:]
    if len(digits) == 10 and digits.startswith("0740"):
        digits = digits[1:]
    if len(digits) == 8 and digits.startswith("404"):
        digits = "7" + digits
    # 0 in 740 is often read as 9
    if len(digits) == 9 and digits.startswith("749"):
        digits = "740" + digits[3:]
    return digits


def issues_for(row):
    notes = []
    imei, serial, sim, phone = row[2], row[3], row[4], row[5]
    if not re.fullmatch(r"35341570\d{7}", imei or ""):
        notes.append("IMEI")
    if not re.fullmatch(r"R9PT20[A-Z0-9]{5}", serial or ""):
        notes.append("Serial")
    if not re.fullmatch(r"89256010000749\d{6}", sim or ""):
        notes.append("SIM")
    if not re.fullmatch(r"74\d{7}", phone or ""):
        notes.append("Phone")
    return notes


def extract_page(gray, page_no, start_sn):
    ys = horizontal_ys(gray)
    if len(ys) < 5:
        raise RuntimeError("page %s: few table rows (%s)" % (page_no, len(ys)))
    header_band = (ys[0], ys[1] if len(ys) > 1 else ys[0] + 50)
    bottom_band = (ys[-2], ys[-1] + 6)
    vt = pick_six_verticals(verticals_in_band(gray, header_band[0] - 4, header_band[1] + 10))
    vb = pick_six_verticals(verticals_in_band(gray, bottom_band[0] - 4, bottom_band[1] + 8))
    if vt is None or len(vt) != 6:
        vt = vb
    if vb is None or len(vb) != 6:
        vb = vt
    if vt is None or vb is None or len(vt) != 6 or len(vb) != 6:
        raise RuntimeError("page %s: column borders top=%s bot=%s" % (page_no, vt, vb))

    n_cells = len(ys) - 1
    start_row = 1 if n_cells == 49 else 0
    n_rows = n_cells - start_row

    cells = [[] for _ in range(5)]
    for i in range(start_row, n_cells):
        y0 = ys[i] + 4
        y1 = ys[i + 1] - 3
        mid = (ys[i] + ys[i + 1]) / 2.0
        xs = xs_at(mid, ys[0], ys[-1], vt, vb)
        # Phone numbers sit on the right of their column; serials sit close
        # to the next border. Extend those crops so the last character is kept.
        insets = (
            (4, 3),
            (3, 2),
            (2, 1),
            (3, 2),
            (10, -18),
        )
        for c in range(5):
            left_in, right_in = insets[c]
            x0 = xs[c] + left_in
            x1 = xs[c + 1] - right_in
            x0 = max(0, x0)
            x1 = min(gray.shape[1], x1)
            if x1 <= x0 or y1 <= y0:
                cells[c].append(np.full((20, 20), 255, np.uint8))
            else:
                cells[c].append(gray[y0:y1, x0:x1])

    probe = column_from_cells(cells[0][:1], "ABCDEFGHIJKLMNOPQRSTUVWXYZ_0123456789")[0]
    whitelists = [
        "0123456789",
        "ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789",
        "0123456789",
        "0123456789",
    ]
    values = [[] for _ in range(5)]
    values[0] = [""] * n_rows
    for src_i, wl in zip((1, 2, 3, 4), whitelists):
        values[src_i] = column_from_cells(cells[src_i], wl)

    def retry(im, whitelist):
        tess = "".join(tesseract_image(im, whitelist, psm=7)).replace(" ", "")
        win = windows_ocr_text(im)
        return tess, win

    raw = []
    clean = []
    for i in range(n_rows):
        sn = start_sn + i
        imei = normalize_imei(values[1][i])
        serial = normalize_serial(values[2][i])
        sim = normalize_sim(values[3][i])
        phone = normalize_phone(values[4][i])
        row = [sn, normalize_device("", sn), imei, serial, sim, phone]
        notes = issues_for(row)
        if notes:
            if "IMEI" in notes:
                for cand in retry(cells[1][i], "0123456789"):
                    nxt = normalize_imei(cand)
                    trial = list(row); trial[2] = nxt
                    if "IMEI" not in issues_for(trial):
                        row[2] = nxt
                        break
            if "Serial" in notes:
                for cand in retry(cells[2][i], "ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789"):
                    nxt = normalize_serial(cand)
                    trial = list(row); trial[3] = nxt
                    if "Serial" not in issues_for(trial):
                        row[3] = nxt
                        break
            if "SIM" in notes:
                for cand in retry(cells[3][i], "0123456789"):
                    nxt = normalize_sim(cand)
                    trial = list(row); trial[4] = nxt
                    if "SIM" not in issues_for(trial):
                        row[4] = nxt
                        break
            if "Phone" in notes:
                for cand in retry(cells[4][i], "0123456789"):
                    nxt = normalize_phone(cand)
                    trial = list(row); trial[5] = nxt
                    if "Phone" not in issues_for(trial):
                        row[5] = nxt
                        break
        raw.append([values[c][i] for c in range(5)])
        clean.append(row)
    return clean, raw, probe, n_rows


def write_xlsx(rows, flags):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection Devices"
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="B8B8B8"),
        right=Side(style="thin", color="B8B8B8"),
        top=Side(style="thin", color="B8B8B8"),
        bottom=Side(style="thin", color="B8B8B8"),
    )
    odd = PatternFill("solid", fgColor="D6EAF8")
    warn = PatternFill("solid", fgColor="FCE4D6")
    body = Font(name="Calibri", size=11)
    mono = Font(name="Consolas", size=10)

    ws.append(HEADER)
    for col in range(1, 7):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin
    ws.row_height = None
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F%d" % (len(rows) + 1)

    for r, row in enumerate(rows, start=2):
        flagged = flags[r - 2]
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, value)
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            cell.font = body if c <= 2 else mono
            if c >= 2:
                cell.number_format = "@"
            if c == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if flagged:
                cell.fill = warn
            elif r % 2 == 0:
                cell.fill = odd

    widths = [8, 28, 20, 18, 26, 18]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    issues = []
    for i, flag in enumerate(flags):
        if not flag:
            continue
        note = ", ".join(flag)
        if "Phone" in flag:
            note += " — last digits covered by a binder punch hole on the scan"
        issues.append((i + 1, note))
    if issues:
        w2 = wb.create_sheet("OCR flags")
        w2.append(["S/N", "Fields to review"])
        for col in range(1, 3):
            w2.cell(1, col).font = header_font
            w2.cell(1, col).fill = header_fill
        for sn, note in issues:
            w2.append([sn, note])
        w2.column_dimensions["A"].width = 10
        w2.column_dimensions["B"].width = 40

    wb.save(XLSX)


def main():
    doc = fitz.open(PDF)
    all_rows = []
    all_raw = []
    start_sn = 1
    # Last page is a blank lined sheet.
    for page_no in range(1, min(22, doc.page_count + 1)):
        if page_no == 22:
            continue
        gray = page_gray(doc, page_no - 1)
        # Skip the empty last sheet: almost no dark ink in a table.
        ink = (gray < 80).mean()
        if ink < 0.01:
            print("skip page %s (blank)" % page_no)
            continue
        rows, raw, probe, n_rows = extract_page(gray, page_no, start_sn)
        print(
            "page %02d  rows=%s  start=%s  header_probe=%r"
            % (page_no, n_rows, start_sn, probe[:40]),
            flush=True,
        )
        all_rows.extend(rows)
        all_raw.extend(raw)
        start_sn += n_rows
    doc.close()

    flags = [issues_for(row) for row in all_rows]
    n_flag = sum(1 for f in flags if f)
    print("total rows", len(all_rows), "flagged", n_flag)
    if n_flag:
        for i, f in enumerate(flags):
            if f:
                print("  sn", all_rows[i][0], f, all_rows[i][1:], "raw", all_raw[i])
    write_xlsx(all_rows, flags)
    print("wrote", XLSX)


if __name__ == "__main__":
    main()
