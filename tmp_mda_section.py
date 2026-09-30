"""Per-MDA charts and exact draft narratives for section 4.2."""
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FIG_DIR = Path(r"D:\coding\ugift-data-analysis\outputs\narrative-report\figures")
DOC_PATH = Path(r"D:\coding\ugift-data-analysis\outputs\narrative-report\UgIFT Asset Verification Report.docx")
BAR = "#243D4A"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

MDAS = [
    {
        "key": "mofped",
        "fig": "46a",
        "short": "MoFPED",
        "bars": [
            ("ICT items", 217),
            ("Other assets", 122),
            ("Furniture and fittings", 99),
            ("Transport assets", 20),
        ],
        "heading": "4.2.5 Ministry of Finance, Planning and Economic Development",
        "paragraphs": [
            "MoFPED held 458 assets: 217 ICT items, 122 other assets, 99 furniture and fittings and 20 transport assets. The main items were office chairs and desks, desktop computers, tablets and 20 Toyota Hilux double cabin pickups. All 458 assets (100.0%) were working and in use, and none was broken down. Of the 358 verified assets, 272 (76.0%) were engraved and 86 (24.0%) were not engraved; 252 (70.4%) were properly engraved with the UgIFT marking, and 20 (5.6%), the pickups, were not properly engraved, carrying their registration numbers only. The recorded value of the ministry's UgIFT assets is UGX 6,152,854,000.",
            "The ministry held 113 computers (73 desktops and 40 laptops), all working and in use; 103 (91.2%) were engraved and 10 (8.8%) were not engraved. Its 17 tablets and 14 office desk phones were all working and engraved. The ministry keeps a UgIFT asset register and the consolidated register for all MDAs that received UgIFT assets.",
            "Expected: Staff computers should support their workload, and portable equipment should remain identifiable and accountable. Found: The computers were working, but users reported that some keyboards and power backup units had failed and that some older computers were freezing. Laptop losses were supported by police letters. The team also reported difficulty obtaining some laptops for engraving.",
            "Gap: Equipment reliability, follow-up of lost assets and completion of identification each required a specific management response. Recommended action: The ministry should assess the failed accessories and freezing computers, follow up the reported laptop losses, and arrange a supervised exercise to engrave all remaining portable equipment.",
            "Expected: Furniture should remain assigned to a location and custodian when offices move. Found: When programme staff moved to premises already furnished, their earlier furniture was left in the old building. Some remained in store, while staff said other items had gone to different offices.",
            "Gap: The move had split the furniture between storage and other offices, requiring confirmation of its present custody and use. Recommended action: The ministry should inspect the old store and recipient offices, confirm the condition and custodian of each item, and approve reuse or other appropriate treatment of furniture no longer required.",
        ],
    },
    {
        "key": "mowt",
        "fig": "46b",
        "short": "MoWT",
        "bars": [("ICT items", 63), ("Other assets", 24), ("Transport assets", 1)],
        "heading": "4.2.6 Ministry of Works and Transport",
        "paragraphs": [
            "MoWT held 88 assets: 63 ICT items, 24 other assets and 1 transport asset, a Toyota Hilux pickup. The ICT items were mainly desktop computers, monitors and uninterruptible power supply (UPS) units. Of the 88 assets, 72 (81.8%) were working and in use and 16 (18.2%) were not working and not in use; none was broken down. Of the 72 verified assets, 63 (87.5%) were engraved, 62 (86.1%) of them with the UgIFT marking, and 9 (12.5%) were not engraved. The recorded value is UGX 455,932,376.",
            "At the time of the visit, the ministry was tracing the location of the 16 assets not in use, 12 Dell OptiPlex desktop computers and 4 UPS units. None of them was reported damaged. Of the ministry's 33 computers (28 desktops and 5 laptops), 21 (63.6%) were working and in use and the 12 desktops being traced made up the other 36.4%; 28 (84.8%) were engraved and 5 (15.2%) were not engraved.",
            "Repairs and servicing of the Toyota Hilux pickup were handled by MoFPED. The custody and maintenance arrangements should remain clear during handover.",
            "Expected: Installed office and conferencing equipment should support the secretariat and carry a traceable asset identity. Found: The secretariat reported that the installed photocopiers supported its daily work. The colour photocopier, the video conferencing system and five Lenovo desktop computers were working but not engraved.",
            "Gap: Working equipment, including an operational conferencing system, still required clear identification and custody, and the location of 16 assets was being traced. Recommended action: The ministry should engrave the colour photocopier and desktop computers, identify the conferencing system and its components, assign a custodian, and apply a suitable durable identifier without damaging the equipment. It should also complete the tracing of the 12 desktop computers and 4 UPS units and confirm their custodians.",
        ],
    },
    {
        "key": "moes",
        "fig": "46c",
        "short": "MoES",
        "bars": [
            ("TELA handsets", 9324),
            ("Tablets", 1001),
            ("Laptops", 456),
            ("Photocopier", 1),
            ("Printer", 1),
        ],
        "heading": "4.2.7 Ministry of Education and Sports",
        "paragraphs": [
            "MoES held 10,783 assets, all ICT equipment: 9,324 TELA handsets, 1,001 tablets (976 of them inspection tablets), 456 laptops, a photocopier and a printer. TELA is the system, financed by UgIFT, that the ministry uses to monitor attendance and performance in schools; the handsets and inspection tablets are deployed in local governments and remain MoES assets. The recorded value is UGX 14,925,128,359, the highest of any MDA.",
            "All 10,783 assets (100.0%) were working and in use. Of the verified assets, 458 (4.2%) were engraved, all with the UgIFT marking, and 10,325 (95.8%) were not engraved: none of the handsets or tablets carries an engraving. All 456 laptops were working, in use and properly engraved.",
            "Recommended action: The ministry should link each handset and tablet to its serial number and to the school or office that holds it, and apply a durable identifier where the device allows.",
        ],
    },
    {
        "key": "maaif",
        "fig": "46d",
        "short": "MAAIF",
        "bars": [("ICT items", 1206), ("Other assets", 837), ("Transport assets", 2)],
        "heading": "4.2.8 Ministry of Agriculture, Animal Industry and Fisheries",
        "paragraphs": [
            "MAAIF held 2,045 assets: 1,206 ICT items, 837 other assets and 2 transport assets. The ICT items were mainly 819 tablets, 211 computers (162 desktops and 49 laptops), 149 monitors and 15 printers. All 2,045 assets (100.0%) were working and in use. Of the 1,227 verified assets, 1,204 (98.1%) were engraved, 1,200 (97.8%) with the UgIFT marking, and 23 (1.9%) were not engraved, among them two drones, a handheld global positioning system (GPS) unit and a projector. The recorded value is UGX 2,332,681,629.",
            "All 211 computers and 819 tablets were working, in use and properly engraved. The ministry had working laptops and printers carrying the UgIFT engraving. Keeping those engravings linked to the equipment and its custodian will support continued accountability after programme closure. Routine condition checks should identify new faults and keep usable equipment in service.",
        ],
    },
    {
        "key": "moh",
        "fig": "46e",
        "short": "MoH",
        "bars": [("Other assets", 200), ("ICT items", 39), ("Transport assets", 25)],
        "heading": "4.2.9 Ministry of Health",
        "paragraphs": [
            "MoH held 264 assets: 200 other assets, 39 ICT items and 25 transport assets. Of these, 263 (99.6%) were working and in use; one double cabin pickup (0.4%) had been boarded off and was not in use, and none was broken down. Of the 73 verified assets, 41 (56.2%) were engraved and 32 (43.8%) were not engraved; 14 (19.2%) were properly engraved with the UgIFT marking and 27 (37.0%) were not properly engraved, carrying another marking only: a vehicle registration number (16) or an office asset code (11). The recorded value is UGX 1,216,859,097.",
            "The ministry's 26 computers (2 desktops and 24 laptops) were all working and in use; 17 (65.4%) were engraved and 9 (34.6%) were not engraved.",
            "Expected: Office equipment in use should be identifiable by a durable engraving. Found: Two heavy duty printers, one at the Industrial Area engineering office and one at headquarters, were in use and in good condition but not engraved. Nine laptops were also not engraved.",
            "Gap: Working equipment at separate offices still lacked the engraving needed for straightforward physical identification. Recommended action: The ministry should engrave these items, link the engravings to their serial numbers and office locations, and obtain custody confirmation from the receiving units.",
        ],
    },
    {
        "key": "opm",
        "fig": "46f",
        "short": "OPM",
        "bars": [("Working and in use", 7), ("Damaged", 3)],
        "heading": "4.2.10 Office of the Prime Minister",
        "paragraphs": [
            "OPM held 10 ICT items, all HP Envy laptops. Seven (70.0%) were working and in use, and three (30.0%) were damaged, broken down and out of use. All 10 laptops (100.0%) were properly engraved with the UgIFT marking. The recorded value is UGX 53,535,300.",
            "The office should assess repair needs for the three damaged laptops and return serviceable equipment to use.",
            "Expected: Computers should have enough capacity for the work assigned to their users. Found: Users reported that the capacity of the HP Envy i3 laptops was below the volume of work handled, such as the performance assessment of local governments, although the machines were described as being in fair condition.",
            "Gap: A usable laptop could still be poorly matched to the workload of its user. Recommended action: The office should assess the requirements of the affected users and decide whether upgrading, reallocating or replacing the laptops would provide suitable capacity.",
        ],
    },
    {
        "key": "mowe",
        "fig": "46g",
        "short": "MoWE",
        "bars": [("ICT items", 2767), ("Other assets", 172)],
        "heading": "4.2.11 Ministry of Water and Environment",
        "paragraphs": [
            "MoWE held 2,939 assets: 2,767 ICT items and 172 other assets. The main items were 2,368 tablets, 2,344 of them Euron tablets, 188 computers (157 Dell OptiPlex all in one desktops and 31 laptops) with their keyboards and UPS units, and 16 real time kinematic GPS machines. All 2,939 assets (100.0%) were working and in use. Of the 2,782 verified assets, 2,740 (98.5%) were engraved, 2,707 (97.3%) with the UgIFT marking, and 42 (1.5%) were not engraved. The recorded value is UGX 5,000,313,297.",
            "All 188 computers were working and in use; 182 (96.8%) were engraved and 6 (3.2%) were not engraved.",
            "Expected: Portable tablets should carry an asset identifier and have an assigned custodian. Found: Ten Apple tablets were not engraved.",
            "Gap: Portable equipment required an identification method that would allow each unit to be traced to its user and location. Recommended action: The ministry should apply suitable durable identifiers, link them to serial numbers, and confirm custody before the next physical check.",
        ],
    },
    {
        "key": "mogl",
        "fig": "46h",
        "short": "MoGLSD",
        "bars": [("Laptops", 8), ("Tablets", 2)],
        "heading": "4.2.12 Ministry of Gender, Labour and Social Development",
        "paragraphs": [
            "MoGLSD held 10 ICT items: 8 laptops and 2 tablets. All 10 (100.0%) were working and in use. Eight (80.0%) were properly engraved with the UgIFT marking and 2 (20.0%) were not engraved. The recorded value is UGX 46,180,000.",
            "Laptops serving Finance and Administration were working and engraved. Two tablets, one for Finance and Administration and the other for Youth and Children Affairs, were working but not engraved. The ministry should engrave the tablets and link their identifiers to the units and officers responsible for them.",
        ],
    },
    {
        "key": "nema",
        "fig": "46i",
        "short": "NEMA",
        "bars": [("Laptops", 4), ("Printer", 1), ("Photocopier", 1)],
        "heading": "4.2.13 National Environment Management Authority",
        "paragraphs": [
            "NEMA held 6 ICT items: 4 laptops, a printer and a heavy duty photocopier. All 6 (100.0%) were working and in use. One (16.7%) was engraved and 5 (83.3%) were not engraved, including all 4 laptops. The recorded value is UGX 46,824,409.",
            "Expected: Computers and shared office equipment should remain identifiable wherever they are used. Found: Four Lenovo laptops serving the executive office and environmental audit, together with a heavy duty printer, were working but not engraved. The authority was awaiting a MoFPED team to engrave them.",
            "Gap: The equipment could support work, but lacked a durable identifying mark. Recommended action: The authority should engrave the five items, record their serial numbers and custodians, and include them in regular physical checks.",
        ],
    },
    {
        "key": "ppda",
        "fig": "46j",
        "short": "PPDA",
        "bars": [("Laptops", 3), ("Tablet", 1), ("Printer", 1)],
        "heading": "4.2.14 Public Procurement and Disposal of Public Assets Authority",
        "paragraphs": [
            "PPDA held 5 ICT items: 3 Lenovo ThinkPad laptops, a Samsung Galaxy tablet and an HP printer. All 5 (100.0%) were working, in use and properly engraved with the UgIFT marking. The recorded value is UGX 30,153,500.",
            "The laptops, tablet and printer were assigned to the legal unit and the strategy and planning unit. The authority should keep the engravings linked to current custodians and update custody whenever equipment moves between units.",
        ],
    },
    {
        "key": "oag",
        "fig": "46k",
        "short": "OAG",
        "bars": [("Laptops", 16), ("Printers", 2), ("Pickup vehicles", 2)],
        "heading": "4.2.15 Office of the Auditor General",
        "paragraphs": [
            "OAG held 20 assets: 18 ICT items (16 laptops and 2 printers) and 2 pickup vehicles. All 20 (100.0%) were working, in use and engraved: 6 (30.0%) were properly engraved with the UgIFT marking and 14 (70.0%) were not properly engraved, carrying only the office's own asset code (12) or a vehicle registration number (2). Of the 16 laptops, 6 (37.5%) carried the UgIFT marking and 10 (62.5%) the office's asset code. The recorded value is UGX 388,773,680.",
            "The two pickups had changed registration numbers. The office should retain the link between the earlier and current registrations, the chassis numbers and the assigned custodians so that each vehicle remains traceable.",
        ],
    },
    {
        "key": "molg",
        "fig": "46l",
        "short": "MoLG",
        "bars": [
            ("Computers", 6),
            ("Computer peripherals", 4),
            ("Copier", 1),
            ("Photocopier", 1),
            ("Pickup", 1),
        ],
        "heading": "4.2.16 Ministry of Local Government",
        "paragraphs": [
            "MoLG held 13 assets: 12 ICT items and 1 transport asset, a pickup. The ICT items were 6 computers (2 desktops and 4 laptops), 4 computer peripherals, a copier and a heavy duty photocopier. All 13 (100.0%) were working and in use. Eight (61.5%) were engraved, 7 (53.8%) with the UgIFT marking, and 5 (38.5%) were not engraved; of the 6 computers, 3 (50.0%) were engraved and 3 (50.0%) were not. The ministry keeps a well organised UgIFT asset register. The recorded value is UGX 343,320,450.",
            "The District Administration unit had a working pickup, photocopier and computer equipment. Two laptops and a desktop set, including its keyboard and mouse, were not engraved. The ministry should engrave this equipment and confirm its location and custodian while keeping the working assets in service.",
        ],
    },
    {
        "key": "molhud",
        "fig": "46m",
        "short": "MoLHUD",
        "bars": [("Motorcycles", 25), ("Other assets", 2)],
        "heading": "4.2.17 Ministry of Lands, Housing and Urban Development",
        "paragraphs": [
            "MoLHUD held 27 assets: 25 Honda motorcycles and 2 other assets. All 27 (100.0%) were working and in use. Of the 26 verified assets, 13 (50.0%) were not properly engraved, carrying only their registration numbers, and 13 (50.0%) were not engraved; none carried the UgIFT marking.",
            "The ministry received survey equipment and motorcycles that it gave out to local governments, and it provided a list of the beneficiaries. The ministry should keep the beneficiary list current and apply the UgIFT engraving so that each motorcycle can be traced to the receiving local government.",
        ],
    },
]


def chart_path(item):
    return FIG_DIR / f"chart_mda_{item['key']}.png"


EXPECTED = {
    "mofped": 458,
    "mowt": 88,
    "moes": 10783,
    "maaif": 2045,
    "moh": 264,
    "opm": 10,
    "mowe": 2939,
    "mogl": 10,
    "nema": 6,
    "ppda": 5,
    "oag": 20,
    "molg": 13,
    "molhud": 27,
}


def draw_charts():
    plt.rcParams["font.family"] = "Calibri"
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    for item in MDAS:
        total = sum(value for _, value in item["bars"])
        if total != EXPECTED[item["key"]]:
            raise SystemExit(f"{item['key']} chart total {total} != {EXPECTED[item['key']]}")
        labels = [label for label, _ in item["bars"]]
        values = [value for _, value in item["bars"]]
        if sum(values) <= 0:
            raise SystemExit(f"empty chart {item['key']}")
        height = 1.6 + 0.48 * len(labels)
        fig, ax = plt.subplots(figsize=(8.2, height), dpi=200)
        positions = list(range(len(labels)))
        ax.barh(positions, values, color=BAR, height=0.62)
        ax.set_yticks(positions, labels)
        ax.invert_yaxis()
        ax.set_xlabel("Number of assets")
        ax.set_title(f"Assets verified: {item['short']}", loc="left", fontsize=13, fontweight="bold", color="#1A1A1A", pad=10)
        ceiling = max(values) * 1.22
        ax.set_xlim(0, ceiling)
        for pos, value in zip(positions, values):
            ax.text(value + ceiling * 0.015, pos, f"{value:,}", va="center", ha="left", fontsize=9, color="#1A1A1A")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.spines["left"].set_color("#B0B0B0")
        ax.spines["bottom"].set_color("#B0B0B0")
        ax.tick_params(axis="y", length=0, labelsize=10)
        ax.tick_params(axis="x", labelsize=9, colors="#333333")
        ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda n, _p: f"{int(n):,}"))
        fig.text(0.01, 0.01, "Source: UgIFT asset verification records.", fontsize=8, color="#666666")
        fig.tight_layout(rect=(0, 0.04, 1, 1))
        fig.savefig(chart_path(item), dpi=200, facecolor="white")
        plt.close(fig)


def _run(parent, text=None, tab=False, field=None, size=None, no_proof=False):
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), size)
        props.append(sz)
    if no_proof:
        props.append(OxmlElement("w:noProof"))
    if len(props):
        run.append(props)
    if text is not None:
        node = OxmlElement("w:t")
        if text.startswith(" ") or text.endswith(" "):
            node.set(XML_SPACE, "preserve")
        node.text = text
        run.append(node)
    if tab:
        run.append(OxmlElement("w:tab"))
    if field:
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), field)
        run.append(char)
    parent.append(run)
    return run


def _instr(parent, text, size):
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), size)
    props.append(sz)
    run.append(props)
    instr = OxmlElement("w:instrText")
    instr.set(XML_SPACE, "preserve")
    instr.text = text
    run.append(instr)
    parent.append(run)


def list_entry(title, bookmark):
    paragraph = OxmlElement("w:p")
    props = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9965")
    tabs.append(tab)
    props.append(tabs)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "30")
    spacing.set(qn("w:line"), "245")
    spacing.set(qn("w:lineRule"), "auto")
    props.append(spacing)
    align = OxmlElement("w:jc")
    align.set(qn("w:val"), "left")
    props.append(align)
    paragraph.append(props)
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), bookmark)
    link.set(qn("w:history"), "1")
    _run(link, text=title, size="18")
    _run(link, tab=True, size="18")
    _run(link, field="begin", size="18")
    _instr(link, f" PAGEREF {bookmark} ", "18")
    _run(link, field="separate", size="18")
    _run(link, text="0", size="18", no_proof=True)
    _run(link, field="end", size="18")
    paragraph.append(link)
    return paragraph


def bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    props = paragraph._p.find(qn("w:pPr"))
    if props is None:
        paragraph._p.insert(0, start)
    else:
        props.addnext(start)
    paragraph._p.append(end)


def next_bookmark_id(document):
    highest = 0
    for node in document.element.body.iter(qn("w:bookmarkStart")):
        highest = max(highest, int(node.get(qn("w:id")) or 0))
    return highest + 1


def remove_previous_mda_block(document):
    removing = False
    pending = []
    for paragraph in document.paragraphs:
        text = paragraph.text
        if paragraph.style and paragraph.style.name == "Heading 3" and text.startswith("4.2.5 "):
            removing = True
        if paragraph.style and paragraph.style.name == "Heading 2" and text.startswith("4.3 "):
            break
        if removing:
            pending.append(paragraph._element)
    for element in pending:
        element.getparent().remove(element)
    for paragraph in list(document.paragraphs):
        xml = paragraph._p.xml
        if "PAGEREF figure_46" in xml and any(f"figure_46{letter}" in xml for letter in "abcdefghijklm"):
            paragraph._element.getparent().remove(paragraph._element)


def insert_narratives(document):
    anchor = None
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name == "Heading 2" and paragraph.text.startswith("4.3 "):
            anchor = paragraph
            break
    if anchor is None:
        raise SystemExit("section 4.3 not found")
    bookmark_id = next_bookmark_id(document)
    for item in MDAS:
        heading = anchor.insert_paragraph_before(item["heading"], style="Heading 3")
        picture = anchor.insert_paragraph_before("", style="Normal")
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(str(chart_path(item)), width=Inches(6.92))
        caption_text = f"Figure {item['fig']}: Assets verified: {item['short']}"
        caption = anchor.insert_paragraph_before(caption_text, style="Caption1")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        bookmark(caption, f"figure_{item['fig']}", bookmark_id)
        bookmark_id += 1
        for text in item["paragraphs"]:
            anchor.insert_paragraph_before(text, style="Normal")
        _ = heading


def insert_figure_list(document):
    host = None
    for paragraph in document.paragraphs:
        xml = paragraph._p.xml
        if "PAGEREF figure_46 " in xml and "figure_46a" not in xml:
            host = paragraph._element
            break
    if host is None:
        raise SystemExit("list entry for Figure 46 not found")
    entries = []
    for item in MDAS:
        title = f"Figure {item['fig']}: Assets verified: {item['short']}"
        entries.append(list_entry(title, f"figure_{item['fig']}"))
    for entry in reversed(entries):
        host.addnext(entry)


def main():
    draw_charts()
    document = Document(DOC_PATH)
    remove_previous_mda_block(document)
    insert_narratives(document)
    insert_figure_list(document)
    document.save(DOC_PATH)
    print("updated", DOC_PATH)


if __name__ == "__main__":
    main()
