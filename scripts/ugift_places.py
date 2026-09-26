"""Local-government and facility names used by every register stage.

The known local governments come from raw-data-grouped/facility-reconciliation.csv
and the vote names in Location(3)2.xlsx. Known facility names per local government
come from the reconciliation. Stage 1 uses these to decide whether a column or
banner really names a place; Stage 2 uses them to write one spelling per fact,
BOOK_TYPE_CODE, LOCATION_SEGMENT1 and LOCATION_SEGMENT3.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RECONCILIATION = ROOT / "raw-data-grouped" / "facility-reconciliation.csv"
LOCATION_MASTER = ROOT / "Location(3)2.xlsx"

PLACEHOLDER = {
    "", "n a", "na", "nil", "nill", "none", "null", "not applicable", "not indicated",
    "not stated", "unknown", "tbd", "x", "xx", "xxx", "not known", "not given", "not there",
    "not sure", "not certain", "not specified", "not provided", "unspecified",
}
LG_SUFFIXES = (
    # (regex on the normalised text, kind)
    (r"\b(district local gov t|district local govt|district local government|dis local gov t|"
     r"district lg|district|dlg|d l g|dc|local government|local govt|lg)$", "DLG"),
    (r"\b(municipal council|municipality council|municipality|municipal|mc|m c)$", "MC"),
    (r"\b(city council|city|cc)$", "CITY"),
    (r"\b(town council|tc)$", "TC"),
)
HEALTH_TAIL = re.compile(
    r"(?i)\b(?:health\s*(?:cent(?:re|er|ere)|ctr|c)|h\s*[./]?\s*c|hc|ch(?=\s*(?:iii|111)))\s*[.\-]?\s*"
    r"(?:iv|iii|ii|i|1v|111|11|1|2|3|4|three|two|four)?\b\.?"
)
SCHOOL_TAIL = re.compile(
    r"(?i)\b(?:seed\s*)?(?:secondary|sec\.?|s\.?)?\s*(?:school|sch\.?|s\.?s\.?s\.?|s\.?s\.?)\b\.?"
    r"|\bseed\b\.?"
)
ROMAN = {"i", "ii", "iii", "iv", "v", "vi"}
SMALL_WORDS = {"of", "and", "the", "for", "de", "da"}
# Misspellings seen in the returns that are more than one edit from the name.
ALIASES = {
    "amorata": "amolatar",
    "amoratar": "amolatar",
    "fortportal": "fort portal",
    "fort portal city": "fort portal",
    "packwach": "pakwach",
    "bukomansmbi": "bukomansimbi",
    "ssabagabo makindye": "makindye ssabagabo",
    "makindye sabagabo": "makindye ssabagabo",
    "rakia": "rakai",
    "bulisa": "buliisa",
}
# Spelling to show where the sources disagree with each other.
DISPLAY_OVERRIDES = {
    "fort portal": "Fort Portal", "luwero": "Luweero", "kasanda": "Kassanda", "rakia": "Rakai", "buliisa": "Buliisa",
}


def _edit_distance(left: str, right: str) -> int:
    previous = list(range(len(right) + 1))
    for i, a in enumerate(left, 1):
        current = [i]
        for j, b in enumerate(right, 1):
            current.append(min(previous[j] + 1, current[j - 1] + 1, previous[j - 1] + (a != b)))
        previous = current
    return previous[-1]


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", clean_text(value).casefold()).strip()


def is_placeholder(value: object) -> bool:
    text = clean_text(value)
    if not text:
        return True
    folded = norm(text)
    if folded in PLACEHOLDER:
        return True
    # Dotted or dashed template blanks such as "…........DISTRICT LOCAL GOVERNMENT".
    if re.fullmatch(r"[\s.…_\-–—:]*", text):
        return True
    if re.search(r"[.…_]{3,}", text) and len(re.sub(r"[.…_\s\-–—:]", "", text)) <= 30:
        letters = re.sub(r"[.…_\s\-–—:]", "", text).casefold()
        if letters in {"", "districtlocalgovernment", "district", "healthcentreiii", "healthcenteriii",
                       "seedschool", "seedsecondaryschool", "nameoflg", "nameoffacility", "lg"} \
                or re.fullmatch(r"(name\s*of\s*)?(lg|district|facility|school|health\s*(centre|center)\s*(iii)?)", letters):
            return True
    return False


@dataclass(frozen=True)
class LocalGovernment:
    base: str            # display spelling of the name, e.g. "Madi-Okollo", "Fort Portal"
    kind: str            # DLG, MC, CITY or TC

    @property
    def key(self) -> str:
        return f"{norm(self.base)} {self.kind.casefold()}".strip()

    @property
    def display(self) -> str:
        if self.kind == "DLG":
            return self.base
        if self.kind == "MC":
            return f"{self.base} MC"
        if self.kind == "CITY":
            return f"{self.base} City"
        return f"{self.base} TC"

    @property
    def book_type_code(self) -> str:
        name = re.sub(r"[\\\-]+", " ", self.base).upper()
        name = re.sub(r"[^A-Z0-9 ]", "", name)
        name = re.sub(r"\s+", " ", name).strip()
        suffix = {"DLG": "", "MC": " MC", "CITY": " CITY", "TC": " TC"}[self.kind]
        return f"{name}{suffix} BK"

    @property
    def location_segment1(self) -> str:
        name = self.base.upper()
        name = re.sub(r"\s*-\s*", r"\\-", name)
        name = re.sub(r"[^A-Z0-9 \\\-]", "", name)
        name = re.sub(r"\s+", " ", name).strip()
        return f"{name} {self.kind}"


def _title(text: str) -> str:
    text = re.sub(r"(?i)\b(st|mt|dr)\.(?=[A-Za-z])", r"\1. ", text)
    words = []
    for index, word in enumerate(text.split()):
        low = word.casefold()
        if low in ROMAN:
            words.append(low.upper())
        elif low in SMALL_WORDS and index:
            words.append(low)
        else:
            words.append("-".join(part[:1].upper() + part[1:].casefold() for part in word.split("-")))
    return " ".join(words)


def split_lg(text: object) -> tuple[str, str] | None:
    """Split "Busia Municipal Council" into ("busia", "MC"). Returns None for blanks."""
    raw = clean_text(text).replace("\\-", "-").replace("\\", "")
    if is_placeholder(raw):
        return None
    folded = norm(raw)
    folded = re.sub(r"^(name\s+of\s+(the\s+)?(lg|local\s+government|district)\s*:?\s*)", "", folded).strip()
    kind = "DLG"
    for pattern, found in LG_SUFFIXES:
        match = re.search(pattern, folded)
        if match:
            folded = folded[: match.start()].strip()
            kind = found
            break
    folded = re.sub(r"\b(district|dist|dis|the)\b", " ", folded)
    folded = re.sub(r"\s+", " ", folded).strip()
    if not folded:
        return None
    return folded, kind


@lru_cache(maxsize=1)
def known_local_governments() -> dict[str, LocalGovernment]:
    """Known local governments keyed by (normalised base, kind)."""
    found: dict[str, LocalGovernment] = {}
    spellings: dict[str, str] = {}

    def add(text: str) -> None:
        split = split_lg(text)
        if not split:
            return
        base, kind = split
        if base in ALIASES:
            base = ALIASES[base]
        display = spellings.get(base)
        if display is None and base in DISPLAY_OVERRIDES:
            display = DISPLAY_OVERRIDES[base]
            spellings[base] = display
        if display is None:
            display = _title(base.replace(" ", " "))
            # Keep the hyphen the reconciliation uses for names such as Madi-Okollo.
            hyphenated = re.sub(r"\s*-\s*", "-", clean_text(text).replace("\\", ""))
            match = re.match(r"^([A-Za-z]+-[A-Za-z]+)", hyphenated)
            if match and norm(match.group(1)) == base:
                display = _title(match.group(1))
            spellings[base] = display
        lg = LocalGovernment(display, kind)
        found.setdefault(lg.key, lg)

    if RECONCILIATION.exists():
        with RECONCILIATION.open(encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                for column in ("lg", "master_lg"):
                    if row.get(column):
                        add(row[column])
    if LOCATION_MASTER.exists():
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(LOCATION_MASTER, read_only=True, data_only=True)
            for (value,) in workbook.active.iter_rows(min_row=2, max_col=1, values_only=True):
                if not value:
                    continue
                segment = re.split(r"(?<!\\)-", str(value))[0]
                if re.search(r"(?i)\b(dlg|district|mc|municipal|city|cc|council|gov)", segment):
                    add(segment)
            workbook.close()
        except Exception:
            pass
    return found


@lru_cache(maxsize=1)
def _base_index() -> dict[str, list[LocalGovernment]]:
    index: dict[str, list[LocalGovernment]] = {}
    for lg in known_local_governments().values():
        index.setdefault(norm(lg.base), []).append(lg)
    return index


@lru_cache(maxsize=4096)
def resolve_lg(text: object, fuzzy: bool = True) -> LocalGovernment | None:
    """Match free text to a known local government, or None when it names none."""
    split = split_lg(text)
    if not split:
        return None
    base, kind = split
    if len(base) < 3 or re.search(r"\d", base):
        return None
    index = _base_index()
    candidates = index.get(base)
    if not candidates and base in ALIASES:
        candidates = index.get(ALIASES[base])
    if not candidates and fuzzy and len(base) >= 5:
        # One typing slip is tolerated (Kakumuro -> Kakumiro). Two is not: Buikwe and
        # Bukwo, or Amuru and Amuria, are different local governments two edits apart.
        matches = [
            lgs for known_base, lgs in index.items()
            if known_base[0] == base[0] and abs(len(known_base) - len(base)) <= 1
            and _edit_distance(base, known_base) == 1
        ]
        if len(matches) == 1:
            candidates = matches[0]
    if not candidates:
        return None
    for lg in candidates:
        if lg.kind == kind:
            return lg
    if kind == "DLG":
        # "Busia" alone is the district when both Busia and Busia MC exist.
        for lg in candidates:
            if lg.kind == "DLG":
                return lg
        return None
    # The text names a municipality or city the list lacks under that kind; keep the kind.
    return LocalGovernment(candidates[0].base, kind)


def lg_from_text(text: object) -> LocalGovernment | None:
    """Local government named anywhere in a longer text, e.g. a banner."""
    raw = clean_text(text)
    if not raw:
        return None
    direct = resolve_lg(raw)
    if direct:
        return direct
    match = re.search(
        r"(?i)([A-Za-z][A-Za-z' \-]{2,40}?)\s+(district\s+local\s+gov(?:ernment|'?t)|district|"
        r"municipal(?:ity)?\s+council|municipality|city\s+council|city|dlg|d\.?c\.?|mc|lg)\b",
        raw,
    )
    if match:
        candidate = resolve_lg(f"{match.group(1)} {match.group(2)}", fuzzy=False)
        if candidate:
            return candidate
        words = match.group(1).split()
        for start in range(len(words)):
            candidate = resolve_lg(" ".join(words[start:]) + " " + match.group(2), fuzzy=False)
            if candidate:
                return candidate
    return None


# ---------------------------------------------------------------- facilities

def facility_kind(*texts: object) -> str:
    """'School', 'Health centre', or '' when the words do not say."""
    blob = " ".join(clean_text(text) for text in texts)
    if re.search(r"(?i)\bhealth\b|\bh\s*[./]?\s*c\s*(?:i{1,3}|1{1,3}|2|3|4)?\b|\bhc\s*(?:i{1,3}|1{1,3}|2|3|4)?\b|hospital|clinic|dispensary|blood\s*bank", blob):
        return "Health centre"
    if re.search(r"(?i)school|\bsss\b|\bss\b|\bs\.s\.s\b|\bs\.s\b|\bseed\b|secondary|\bsec\b|\bsch\b", blob):
        return "School"
    return ""


def facility_base(name: object) -> str:
    """Facility name without its type words, in normalised form, for matching."""
    text = clean_text(name)
    text = re.sub(r"(?i)\b(upgraded|proposed|new)\b", " ", text)
    text = HEALTH_TAIL.sub(" ", text)
    text = SCHOOL_TAIL.sub(" ", text)
    text = re.sub(r"(?i)\b(iii|ii|iv|111|11|3|2)\b", " ", text)
    text = re.sub(r"(?i)\b(st|saint)\.?\b", "st", text)
    return norm(text)


def facility_key(name: object, kind: str = "") -> str:
    base = facility_base(name)
    kind = kind or facility_kind(name)
    tag = {"School": "s", "Health centre": "h"}.get(kind, "")
    return f"{base}|{tag}" if base else ""


@lru_cache(maxsize=1)
def known_facilities() -> dict[str, dict[str, tuple[str, str]]]:
    """{lg key: {facility key: (display name, kind)}} from the reconciliation."""
    index: dict[str, dict[str, tuple[str, str]]] = {}
    if not RECONCILIATION.exists():
        return index
    with RECONCILIATION.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            kind = {"School": "School", "Health centre": "Health centre"}.get(row.get("type", ""), "")
            lgs = {resolve_lg(row.get(column, "")) for column in ("lg", "master_lg")}
            names = [row.get("field_name", ""), row.get("name", "")]
            display = ""
            for name in names:
                if name and not is_placeholder(name):
                    display = facility_display(name, kind, official=True)
                    break
            if not display:
                continue
            for lg in lgs:
                if lg is None:
                    continue
                bucket = index.setdefault(lg.key, {})
                for name in names:
                    key = facility_key(name, kind)
                    if key and key not in bucket:
                        bucket[key] = (display, kind)
    return index


def facility_display(name: object, kind: str = "", *, official: bool = False) -> str:
    """One spelling of a facility name, ending in Seed Secondary School or Health Centre III."""
    text = clean_text(name)
    if is_placeholder(text):
        return ""
    text = re.sub(r"(?i)^(name\s+of\s+(the\s+)?(health\s+)?(facility|school|health\s+cent(re|er))\s*:?\s*)", "", text).strip(" :-")
    text = re.sub(r"(?i)\s+\.s\.?(?=\s|$)", " ", text)
    text = re.sub(r"[…]+", " ", text)
    text = re.sub(r"(?i)\bsch\b\.?", "School", text)
    # Stray tokens after the name: "Lll" for III, "UgIFT", "DC" for the district.
    text = re.sub(r"(?i)\s+\b(?:lll|ugift|ug\s*ift|dc|dlg)\b\.?", " ", text)
    text = re.sub(r"(?i)\bhigh\s+(?:s\.?s\.?s?\.?|sec(?:ondary)?\.?(?:\s*(?:school|sch))?)\b\.?", "High School", text)
    named_kind = facility_kind(text)
    kind = named_kind or kind
    if kind == "Health centre" and re.search(r"(?i)hospital|blood bank|clinic", text):
        # A hospital or blood bank keeps its own name; the suffix rule is for health centres.
        text = re.sub(r"\s+", " ", text).strip(" ,.-")
    elif kind == "Health centre":
        text = HEALTH_TAIL.sub(" Health Centre III ", text)
        text = re.sub(r"(?i)\bhealth centre iii\b(?:\s+health centre iii\b)+", "Health Centre III", text)
        text = re.sub(r"\s+", " ", text).strip(" ,.-")
        if not re.search(r"(?i)\bhealth centre iii\b", text):
            text = f"{text} Health Centre III"
        # Words after the suffix ("Kyeihara Health Centre III Mitooma") are not the name.
        cut = re.match(r"(?i)^(.+?\bhealth centre iii)\b.*$", text)
        if cut and cut.group(1).strip().casefold() != "health centre iii":
            text = cut.group(1)
    elif kind == "School" and re.search(
        r"(?i)\b(primary|p\.?\s?s\.?|nursery|high school|college|university|institute|technical|vocational institute|ptc|polytechnic|community school)\b", text
    ) and not re.search(r"(?i)\bseed\b", text):
        # A primary school, high school or college keeps its own name; the seed
        # secondary school suffix is for the UgIFT seed schools.
        text = re.sub(r"\s+", " ", text).strip(" ,.-")
    elif kind == "School":
        # Normalise every spelling of the suffix, then make sure the name carries it once.
        text = re.sub(
            r"(?i)\b(?:seed\s*)?(?:secondary|sec\.?|s\.?)\s*(?:school|sch\.?|s\.?)\b\.?|"
            r"\b(?:seed\s*)?s\.?s\.?s\.?\b|\bseed\s+school\b|\bseed\s+sch\b\.?",
            " Seed Secondary School ",
            text,
        )
        text = re.sub(r"(?i)\bseed\s+seed\b", "Seed", text)
        text = re.sub(r"(?i)\bseed secondary school\b(?:\s+seed secondary school\b)+", "Seed Secondary School", text)
        text = re.sub(r"(?i)\b(?:seed\s+secondary\s+)+(?=seed secondary school\b)", "", text)
        text = re.sub(r"(?i)\bschool\s+seed secondary school\b", "Seed Secondary School", text)
        text = re.sub(r"\s+", " ", text).strip(" ,.-")
        if not re.search(r"(?i)\bseed secondary school\b", text):
            text = re.sub(r"(?i)\s*\bseed\b\s*$", "", text).strip()
            text = re.sub(r"(?i)\s*\bschool\b\.?\s*$", "", text).strip()
            text = f"{text} Seed Secondary School"
        cut = re.match(r"(?i)^(.+?\bseed secondary school)\b.*$", text)
        if cut and cut.group(1).strip().casefold() != "seed secondary school":
            text = cut.group(1)
    text = re.sub(r"\s+", " ", text).strip(" ,.-")
    titled = _title(text)
    titled = re.sub(r"(?i)\bhealth centre iii\b", "Health Centre III", titled)
    titled = re.sub(r"(?i)\bseed secondary school\b", "Seed Secondary School", titled)
    return titled


def canonical_facility(name: object, lg: LocalGovernment | None, kind: str = "") -> tuple[str, str]:
    """(display name, kind) using the reconciliation spelling when the facility is known."""
    # The name's own words decide its kind: "Maregamo Health Centre III" is a health
    # centre even when the table it came from is a school table.
    named_kind = facility_kind(name)
    if named_kind:
        kind = named_kind
    kind = kind or named_kind
    if lg is not None:
        bucket = known_facilities().get(lg.key, {})
        key = facility_key(name, kind)
        if key in bucket:
            return bucket[key]
        base = facility_base(name)
        # "KWANIA OWINYI HC III": the district typed in front of the facility name.
        lg_words = norm(lg.base)
        if lg_words and base.startswith(lg_words + " ") and len(base) > len(lg_words) + 3:
            base = base[len(lg_words) + 1:]
            for tag in {"School": ("s",), "Health centre": ("h",)}.get(kind, ("s", "h")):
                if f"{base}|{tag}" in bucket:
                    return bucket[f"{base}|{tag}"]
        tags = {"School": ("s",), "Health centre": ("h",)}.get(kind, ("s", "h"))
        for tag in tags:
            probe = f"{base}|{tag}"
            if probe in bucket:
                return bucket[probe]
        # "Aporu Okol" names "Dr Aporu Okol"; "Malera Kabarwa" names "Kabarwa": one
        # known facility of the kind whose name is a whole word of the other.
        if len(base) >= 5:
            contained = {
                bucket[known]
                for known in bucket
                if known.split("|")[1] in tags and len(known.split("|")[0]) >= 5
                and (re.search(rf"\b{re.escape(known.split('|')[0])}\b", base) or re.search(rf"\b{re.escape(base)}\b", known.split("|")[0]))
            }
            if len(contained) == 1:
                return contained.pop()
        # A misspelt known facility (Angatta for Angetta, Adekenin for Adeknino) still
        # names it when only one facility of that kind in the local government is that
        # close: one edit for short names, two for names of six letters or more.
        if len(base) >= 4:
            allowed = 2 if len(base) >= 6 else 1
            for distance in range(1, allowed + 1):
                close = {
                    bucket[known]
                    for known in bucket
                    if known.split("|")[1] in tags
                    # Same first letter, or one name is the other minus its first letter
                    # ("Wemba" for "Iwemba").
                    and (known.split("|")[0][:1] == base[:1] or (distance == 1 and (known.split("|")[0].endswith(base) or base.endswith(known.split("|")[0]))))
                    and _edit_distance(known.split("|")[0], base) == distance
                }
                if len(close) == 1:
                    return close.pop()
                if close:
                    break
    return facility_display(name, kind), kind


def lg_of_known_facility(name: object, kind: str = "") -> LocalGovernment | None:
    """The one local government whose reconciled facility list holds this name,
    or None when no or several local governments do."""
    kind = facility_kind(name) or kind
    key = facility_key(name, kind)
    if not key or key.split("|")[0] in {"", "st"} or len(key.split("|")[0]) < 5:
        return None
    owners = {lg_key for lg_key, bucket in known_facilities().items() if key in bucket}
    if len(owners) != 1:
        return None
    lg_key = owners.pop()
    return known_local_governments().get(lg_key)


def check_examples() -> None:
    assert split_lg("Busia Municipal Council") == ("busia", "MC")
    assert split_lg("PALLISA DISTRICT LOCAL GOVERNMENT") == ("pallisa", "DLG")
    assert split_lg("MADI\\-OKOLLO DLG") == ("madi okollo", "DLG")
    assert split_lg("…...............DISTRICT LOCAL GOVERNMENT") is None
    assert is_placeholder("N/A") and is_placeholder("-") and not is_placeholder("Busia")
    lg = resolve_lg("Madi-Okollo")
    assert lg and lg.book_type_code == "MADI OKOLLO BK" and lg.location_segment1 == "MADI\\-OKOLLO DLG", lg
    assert resolve_lg("Hoima District").book_type_code == "HOIMA BK"
    assert resolve_lg("Kiira Municipal Council").book_type_code == "KIIRA MC BK"
    assert resolve_lg("Busia MC").display == "Busia MC" and resolve_lg("Busia").display == "Busia"
    assert resolve_lg("Tororo District Local Government").location_segment1 == "TORORO DLG"
    assert resolve_lg("KAKUMURO").display == "Kakumiro"
    assert resolve_lg("Muyuge District").display == "Mayuge"
    assert resolve_lg("KYANKWAZI").display == "Kyankwanzi"
    assert resolve_lg("ENT Basic for HCIII") is None
    assert resolve_lg("Book shelves)") is None
    assert resolve_lg("2 at staff house") is None
    assert resolve_lg("Functional") is None
    assert resolve_lg("Education") is None
    assert resolve_lg("Ministry of Water and Environment") is None
    assert lg_from_text("NAME OF LG: DOKOLO DISTRICT LOCAL GOVERNMENT").display == "Dokolo"
    assert lg_from_text("BITSYA HCIII, BUHWEJU DC").display == "Buhweju"
    assert facility_display("OGUR SEED SCHOOL") == "Ogur Seed Secondary School"
    assert facility_display("Iyolwa Seed Sch") == "Iyolwa Seed Secondary School"
    assert facility_display("BUGIRI MUNICIPAL COUNCIL HC III") == "Bugiri Municipal Council Health Centre III"
    assert facility_display("Muggi Health Center III") == "Muggi Health Centre III"
    assert facility_display("Kidubuli HC II") == "Kidubuli Health Centre III"
    assert facility_display("St Mugagga Vocational Seed Secondary School") == "St Mugagga Vocational Seed Secondary School"
    assert facility_display("Luwube Muslim Seed Secondary School") == "Luwube Muslim Seed Secondary School"
    assert facility_display("St.Andrews Ndwaddemutwe SSS") == "St. Andrews Ndwaddemutwe Seed Secondary School"
    assert resolve_lg("AMORATA DISTRICT").display == "Amolatar"
    assert resolve_lg("Amuru").display == "Amuru" and resolve_lg("Buikwe").display == "Buikwe"
    assert facility_display("Arua Regional Blood Bank") == "Arua Regional Blood Bank"
    assert facility_display("AWEI HC 111") == "Awei Health Centre III"
    assert facility_display("Sop Sop") == "Sop Sop"
    assert facility_kind("Sop Sop", "Health centre") == "Health centre"
    assert facility_key("OGUR SEED SCHOOL") == facility_key("Ogur Seed Secondary School")
    assert facility_key("Kidubuli HC II") == facility_key("Kidubuli Health Centre III")
    lg = resolve_lg("Tororo")
    assert canonical_facility("Kamuli HC II", lg)[0] == "Kamuli Health Centre III"
    assert facility_kind("Ministry of Water and Environment") == ""


if __name__ == "__main__":
    check_examples()
    print(f"{len(known_local_governments())} known local governments")
    for text in ("Madi-Okollo", "Busia MC", "Fort Portal City", "Kiira Municipal Council", "Mbarara City", "Lira City"):
        lg = resolve_lg(text)
        print(f"{text!r:32} -> {lg.display!r:22} {lg.book_type_code!r:24} {lg.location_segment1!r}")
    print("ok")
