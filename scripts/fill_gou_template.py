"""Fill the GOU asset-register template from the shared UgIFT register.

Columns A-AW keep the sample headers. A shared-register column is written
there only when the 2023 asset-accounting guidelines and the sample row give
that column a meaning the source can fill. Every other source column is kept
once, in order, under ATTRIBUTE1, ATTRIBUTE2, and so on.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "asset-register-2026-09-23" / "UgIFT Shared Asset Register.xlsx"
TEMPLATE = ROOT / "Sample Header of Asset Register..xlsx"
OUTPUT = ROOT / "outputs" / "asset-register-2026-09-23" / "UgIFT Asset Register GOU Template.xlsx"

# Source columns that already have a home in A-AW.
MAPPED_TO_TEMPLATE = {
    "Equipment/Item",
    "Department",
    "Asset Number",
    "Life in Months",
    "Tag Number (engrave no.)",
    "Date Placed In Service",
    "Cost",
    "Acc Dep Cost",
    "Ytd Deprn",
    "Local Government",
    "Facility",
}

NOT_IN_USE = re.compile(
    r"not\s+(?:in\s+use|received|available|functional|functioning)|"
    r"obsolete|unserviceable|disposed|for disposal|\blost\b|\bmissing\b",
    re.I,
)
IN_USE = re.compile(
    r"\bin use\b|functional|functioning|working well|\bavailable\b|good condition",
    re.I,
)


def clean(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).replace("\xa0", " ").replace("\n", " ").replace("\r", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def as_date(value: object) -> object:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return clean(value)


def as_number(value: object) -> object:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    text = clean(value)
    if not text:
        return None
    bare = text.replace(",", "")
    if re.fullmatch(r"-?\d+", bare):
        return int(bare)
    if re.fullmatch(r"-?\d+\.\d+", bare):
        number = float(bare)
        return int(number) if number.is_integer() else number
    return text


def book_code(lg: str | None) -> str | None:
    if not lg:
        return None
    text = lg.upper()
    for old, new in (
        (" DISTRICT LOCAL GOV'T", ""),
        (" DISTRICT LOCAL GOVT", ""),
        (" DISTRICT LOCAL GOVERNMENT", ""),
        (" MUNICIPALITY COUNCIL", " MC"),
        (" MUNICIPAL COUNCIL", " MC"),
        (" MUNICIPALITY", " MC"),
        (" CITY COUNCIL", " CITY"),
        (" DISTRICT LG", ""),
        (" LOCAL GOVERNMENT", ""),
        (" DISTRICT", ""),
        (" DLG", ""),
    ):
        if text.endswith(old):
            text = (text[: -len(old)] + new).strip()
            break
    text = text.replace("\\", "").replace("-", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return f"{text} BK" if text else None


def location_segment1(lg: str | None) -> str | None:
    """Vote name in the sample form, such as MADI\\-OKOLLO DLG."""
    if not lg:
        return None
    text = lg.upper().replace("\\", "")
    kind = "DLG"
    for old, new_kind in (
        (" DISTRICT LOCAL GOV'T", "DLG"),
        (" DISTRICT LOCAL GOVT", "DLG"),
        (" DISTRICT LOCAL GOVERNMENT", "DLG"),
        (" MUNICIPALITY COUNCIL", "MC"),
        (" MUNICIPAL COUNCIL", "MC"),
        (" MUNICIPALITY", "MC"),
        (" CITY COUNCIL", "CITY"),
        (" DISTRICT LG", "DLG"),
        (" LOCAL GOVERNMENT", "DLG"),
        (" DISTRICT", "DLG"),
        (" DLG", "DLG"),
        (" CITY", "CITY"),
        (" MC", "MC"),
    ):
        if text.endswith(old):
            text = text[: -len(old)].strip()
            kind = new_kind
            break
    text = re.sub(r"\s*-\s*", r"\\-", text)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return None
    return f"{text} {kind}"


def in_use_flag(status: str | None) -> str | None:
    if not status:
        return None
    if NOT_IN_USE.search(status):
        return "NO"
    if IN_USE.search(status):
        return "YES"
    return None


def template_headers() -> list[str]:
    workbook = load_workbook(TEMPLATE, read_only=True, data_only=True)
    headers = [clean(value) or "" for value in next(workbook.active.iter_rows(max_row=1, values_only=True))]
    workbook.close()
    return headers


def attribute_headers(source_headers: list[str], template: list[str]) -> list[str]:
    pending = [name for name in source_headers if name not in MAPPED_TO_TEMPLATE]
    headers = list(template)
    for index, name in enumerate(pending, 1):
        headers[48 + index] = f"ATTRIBUTE{index}({name})"
    return headers


def build_row(values: tuple[object, ...], source_headers: list[str]) -> list[object]:
    source = {name: values[index] if index < len(values) else None for index, name in enumerate(source_headers)}
    item = clean(source.get("Equipment/Item"))
    description = clean(source.get("Item Description"))
    department = clean(source.get("Department"))
    facility = clean(source.get("Facility"))
    lg = clean(source.get("Local Government"))
    life = as_number(source.get("Life in Months"))
    status = clean(source.get("Equipment status"))
    row = [None] * 64
    row[0] = book_code(lg)
    row[1] = item or description
    row[7] = 1
    row[8] = location_segment1(lg)
    row[9] = department
    row[10] = facility
    row[12] = as_number(source.get("Cost"))
    row[31] = as_date(source.get("Date Placed In Service"))
    if isinstance(life, (int, float)) and life > 0:
        row[32] = "YES"
        row[33] = "STL"
        row[34] = life
    row[36] = as_number(source.get("Acc Dep Cost"))
    row[37] = as_number(source.get("Ytd Deprn"))
    row[39] = clean(source.get("Asset Number"))
    row[41] = clean(source.get("Tag Number (engrave no.)"))
    row[46] = in_use_flag(status)
    attribute_index = 49
    for name in source_headers:
        if name in MAPPED_TO_TEMPLATE:
            continue
        value = source.get(name)
        row[attribute_index] = as_date(value) if "date" in name.casefold() else (
            as_number(value) if name in {"Recoverable cost", "Net Book Value"} else clean(value) if not isinstance(value, (int, float, datetime, date)) else value
        )
        if isinstance(value, datetime):
            row[attribute_index] = value.date()
        elif isinstance(value, date):
            row[attribute_index] = value
        attribute_index += 1
    return row


def main() -> None:
    template = template_headers()
    source_book = load_workbook(SOURCE, read_only=True, data_only=True)
    source_sheet = source_book["Asset Register"]
    rows = source_sheet.iter_rows(values_only=True)
    source_headers = [clean(value) or "" for value in next(rows)]
    headers = attribute_headers(source_headers, template)
    output = Workbook(write_only=True)
    register = output.create_sheet("Asset Register")
    register.append(headers)
    count = 0
    for values in rows:
        if not any(value not in (None, "") for value in values):
            continue
        register.append(build_row(values, source_headers))
        count += 1
        if count % 50000 == 0:
            print(f"{count:,}", flush=True)
    source_book.close()
    notes = output.create_sheet("Read Me")
    for line in (
        "UgIFT asset register in the GOU template",
        "Filled from UgIFT Shared Asset Register. The shared register was not changed.",
        "Columns A to AW follow Sample Header of Asset Register and the 2023 asset-accounting guidelines.",
        "BOOK_TYPE_CODE is the local government in the sample form, ending in BK. LOCATION_SEGMENT1 is the same government as the vote name.",
        "DESCRIPTION is the equipment item. LOCATION_SEGMENT2 is the department. LOCATION_SEGMENT3 is the facility.",
        "FIXED_ASSETS_UNITS is 1 because each row is one item. Cost is the historical cost. Accumulated depreciation is the depreciation reserve. Year-to-date depreciation is YTD_DEPRN.",
        "Where a life in months is recorded, the row is marked depreciable on the straight-line method used in the guidelines. The life itself is the life recorded in the source.",
        "IN_USE_FLAG is YES or NO only where the equipment status says so. The full status text is kept in an attribute column.",
        "Recoverable cost is not written as salvage value. The guidelines treat salvage as the residual built into a class life, and recoverable amount as an impairment test.",
        "Asset category, account codes, prorate convention and asset type stay blank. The shared register does not state them.",
        "Attribute columns hold the source columns that have no field in A to AW, in source order.",
        "",
        "Attribute headers:",
        *[header for header in headers[49:] if header and header.startswith("ATTRIBUTE") and "(" in header],
        "",
        f"Asset rows: {count:,}",
    ):
        notes.append([line])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    output.save(OUTPUT)
    print(f"wrote {count:,} rows to {OUTPUT}")


if __name__ == "__main__":
    main()
