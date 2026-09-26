"""Merge shared UgIFT asset registers into one workbook.

Sources are the shared workbooks under raw-data-grouped (team, district and
multi-team files), not facility-by-facility copies. The column layout follows
new-templates-to-follow. A source line that states a quantity becomes that
many rows, one physical item each.
"""

from __future__ import annotations

import csv
import hashlib
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, date
from pathlib import Path

import xlrd
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from ugift_places import (
    LocalGovernment,
    _edit_distance,
    canonical_facility,
    facility_base,
    facility_key,
    facility_kind,
    fuzzy_owner,
    is_placeholder,
    known_facilities,
    known_local_governments,
    lg_from_text,
    lg_of_known_facility,
    resolve_lg,
)

ROOT = Path(__file__).resolve().parents[1]
GROUPED = ROOT / "raw-data-grouped"
RECONCILIATION = GROUPED / "facility-reconciliation.csv"
OUTPUT = ROOT / "outputs" / "asset-register-2026-09-23" / "ALL_UGIFT_ASSET_REGISTER_SK_TEMPLATE.xlsx"

HEADERS = [
    "Equipment/Item",
    "Department",
    "Asset Number",
    "Item Description",
    "Life in Months",
    "Tag Number (engrave no.)",
    "Date Of Purchase",
    "Date Placed In Service",
    "Recoverable cost",
    "Cost",
    "Acc Dep Cost",
    "Net Book Value",
    "Ytd Deprn",
    "Equipment status",
    "Remarks",
    "Local Government",
    "Facility",
    "Facility type",
    "Unit",
    "Source file",
    "Source location",
]

MAX_QUANTITY = 500
BULK_ITEM = re.compile(
    r"(?i)\b(desks?|chairs?|stools?|tables?|shel(?:f|ves|ving)|bench(?:es)?|beds?|cupboards?|couch(?:es)?|cylinders?)\b"
)
MODEL_TAIL = re.compile(
    r"(?i)^(laserjet|laptop|printer|elitebook|probook|latitude|inspiron|pavilion|"
    r"thinkpad|monitor|cpu|iphone|samsung|nokia|tecno|inch|gen|core|mhz|gb)$"
)
# A model family anywhere in the name makes a trailing number a model, not a count.
MODEL_FAMILY = re.compile(r"(?i)\b(laserjet|laptop|elitebook|probook|latitude|inspiron|pavilion|thinkpad|thinkbook|ideapad|printer|deskjet|officejet|prolite|optiplex|vostro)\b")
WORD_COUNTS = {
    "a": 1, "an": 1, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
    "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
}
COUNT_WORD = r"(?:\d{1,4}|one|two|three|four|five|six|seven|eight|nine|ten|a|an)"


@dataclass
class Asset:
    item: str = ""
    department: str = ""
    asset_number: str = ""
    description: str = ""
    life: object = None
    tag: str = ""
    purchase: object = None
    service: object = None
    recoverable: object = None
    cost: object = None
    acc_dep: object = None
    nbv: object = None
    ytd: object = None
    status: str = ""
    remarks: str = ""
    lg: str = ""
    facility: str = ""
    facility_type: str = ""
    explicit_qty: int | None = None
    source_file: str = ""
    source_location: str = ""
    extras: dict = field(default_factory=dict)

    def signal(self) -> bool:
        return any([
            self.item, self.description, self.tag, self.status, self.remarks,
            self.cost not in (None, ""), self.asset_number,
        ])


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\xa0", " ").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", clean(value).casefold()).strip()


def squash(value: object) -> str:
    """Header text without spaces, so 'Equipme nt Item' and 'Equipment/Item' compare equal."""
    return norm(value).replace(" ", "")


def as_count(token: str) -> int | None:
    token = token.casefold()
    if token.isdigit():
        return int(token)
    return WORD_COUNTS.get(token)


def excel_date(value: object) -> object:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 60_000:
        try:
            return xlrd.xldate_as_datetime(float(value), 0).date()
        except Exception:
            return value
    return value


def number_or_text(value: object) -> object:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value
    text = clean(value)
    if not text:
        return None
    bare = text.replace(",", "")
    if re.fullmatch(r"-?\d+", bare):
        return int(bare)
    if re.fullmatch(r"-?\d+\.\d+", bare):
        number = float(bare)
        return int(number) if number.is_integer() else number
    return text


def family_key(path: Path) -> str:
    stem = path.stem
    stem = re.sub(r"\s*[\(\[]\d+[\)\]]\s*$", "", stem)
    stem = re.sub(r"\s*-\s*copy\s*$", "", stem, flags=re.I)
    return re.sub(r"\s+", " ", stem).strip().casefold()


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def path_context(path: Path) -> tuple[str, str]:
    """Local government and facility implied by the grouped folder path."""
    parts = path.relative_to(GROUPED).parts
    if "_district-documents" in parts:
        index = parts.index("_district-documents")
        if index:
            return parts[index - 1].replace("-", " "), ""
    if parts and re.fullmatch(r"team-\d+", parts[0]) and len(parts) >= 3:
        if parts[1] not in {"_team-documents", "_district-documents"}:
            facility = ""
            if parts[2] not in {"_district-documents", "_team-documents"}:
                facility = parts[2].replace("-", " ")
            return parts[1].replace("-", " "), facility
    return "", ""


def default_lg(path: Path) -> str:
    return path_context(path)[0]


def excluded_programme_file(path: Path) -> bool:
    """Leave out programme-documents, but keep the data-management chat."""
    parts = path.parts
    if "programme-documents" not in parts:
        return False
    index = parts.index("programme-documents")
    remainder = parts[index + 1:]
    return not (remainder and remainder[0] == "data-management-chat")


# Files under raw-data-grouped that carry an asset table but are not a return of
# the facility they are filed under. Each is left out with its reason on Read Me.
EXCLUDED_SOURCES: dict[str, str] = {
    "team-13/Tororo/Kamuli-HC-III/88_facility-fixed-asset-register_ref-kamuli.xls":
        "IFMS fixed-asset register format example: its rows name Jinja Regional Referral Hospital "
        "wards, vehicles and plots, not Kamuli Health Centre III assets.",
}

# Folders whose workbooks are programme lists or later copies of returns that are
# already read from the team folders. Each is left out with its reason on Read Me.
EXCLUDED_FOLDERS: dict[str, str] = {
    "_multi-team/teams-10-15/final-UGiFT-report-karamojja-6-teams":
        "The MDA status registers there are programme supply and WIP lists (55,468 of 58,818 rows are "
        "'programme supply, not field-verified' transcriptions of programme-documents), and the LG register and "
        "registers-by-facility workbooks are later copies of the team-10 to team-15 Asset-Verification-Toolkit "
        "returns that are read from the team folders. Reading both would count each facility two to five times.",
    "_multi-team/teams-10-15/pdf-to-excel":
        "Conversions of the TELA phone and tablet distribution lists and a scanned supply schedule: distribution "
        "lists record what was supplied, not what a facility holds.",
}

# Every folder a byte-identical copy of a source was filed under. A combined
# health-centre-and-school return is filed under both facilities, so the copy that
# is read keeps both contexts and each table is assigned to the facility of its kind.
CONTEXTS: dict[Path, list[tuple[str, str]]] = {}


def candidate_files() -> list[Path]:
    found: list[Path] = []
    for path in GROUPED.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in {".xls", ".xlsx", ".docx"}:
            continue
        if path.name.startswith("~$"):
            continue
        if excluded_programme_file(path):
            continue
        relative = path.relative_to(GROUPED).as_posix()
        if relative in EXCLUDED_SOURCES or any(relative.startswith(folder + "/") for folder in EXCLUDED_FOLDERS):
            continue
        found.append(path)
    grouped: dict[tuple[str, str], list[Path]] = defaultdict(list)
    for path in found:
        grouped[(str(path.parent).casefold(), path.stem.casefold())].append(path)
    chosen: list[Path] = []
    for paths in grouped.values():
        suffixes = {item.suffix.lower() for item in paths}
        if ".docx" in suffixes and suffixes & {".xls", ".xlsx"}:
            # The same stem as a spreadsheet and a Word file: keep the spreadsheet.
            paths.sort(key=lambda item: (
                {".xlsx": 0, ".xls": 1, ".docx": 2}.get(item.suffix.lower(), 3),
                -item.stat().st_size, len(str(item)), str(item),
            ))
            chosen.append(paths[0])
        else:
            chosen.extend(paths)
    by_digest: dict[str, list[Path]] = defaultdict(list)
    for path in sorted(chosen, key=lambda item: str(item).casefold()):
        by_digest[file_hash(path)].append(path)
    unique: list[Path] = []
    CONTEXTS.clear()
    for copies in by_digest.values():
        # Prefer the copy filed under a facility folder over a team-level copy.
        copies.sort(key=lambda item: (not path_context(item)[1], not path_context(item)[0], str(item).casefold()))
        kept = copies[0]
        contexts: list[tuple[str, str]] = []
        for copy in copies:
            context = path_context(copy)
            if context[0] and context not in contexts:
                contexts.append(context)
        CONTEXTS[kept] = contexts
        unique.append(kept)
    unique.sort(key=lambda item: str(item).casefold())
    return unique


def docx_has_asset_table(path: Path) -> bool:
    """The prompt's test: an Equipment/Item header, or a description column with
    condition, quantity, tag, or cost. Word splits header words across runs, so the
    text is compared with tags and spaces removed."""
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
    except (OSError, zipfile.BadZipFile, KeyError):
        return False
    tables = re.findall(rb"<w:tbl>.*?</w:tbl>", xml, re.S)
    for table in tables:
        text = re.sub(rb"<[^>]+>", b"", table)
        blob = re.sub(rb"[^a-z0-9]+", b"", text.lower())
        if b"equipmentitem" in blob or b"equipmentasset" in blob:
            return True
        if b"description" in blob and any(
            token in blob for token in (b"condition", b"status", b"quantity", b"qty", b"tagnumber", b"engrave", b"cost")
        ):
            return True
        if b"item" in blob and any(token in blob for token in (b"quantity", b"qty", b"supplied", b"found", b"functional")):
            return True
    return False


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEALTH_SECTION = re.compile(r"(?i)^\s*(health\s+cent(?:re|er)s?(\s+checklist)?|health\s+facilit(?:y|ies))\b")
SCHOOL_SECTION = re.compile(
    r"(?i)^\s*(school\s+verification|school\s+interview|school\s+furniture|ict\s+items\s+to\s+verify|"
    r"buildings\s+to\s+verify|seed\s+schools?|schools?\s+checklist)\b"
)


def docx_cell_text(tc) -> str:
    return clean(xml_text(tc))


MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def xml_text(element) -> str:
    """The text of a Word element: line breaks and paragraph ends become spaces, and
    the duplicate copy of a text box kept for old Word versions (mc:Fallback) is
    read once."""
    pieces: list[str] = []

    def walk(node) -> None:
        if node.tag == MC_FALLBACK:
            return
        if node.tag == f"{W_NS}t":
            pieces.append(node.text or "")
        elif node.tag in (f"{W_NS}br", f"{W_NS}tab", f"{W_NS}cr"):
            pieces.append(" ")
        elif node.tag == f"{W_NS}p" and pieces:
            pieces.append(" ")
        for child in node:
            walk(child)

    walk(element)
    return "".join(pieces)


def docx_grid_rows(table_element) -> list[list[object]]:
    """Table rows laid on the column grid, so a horizontally merged cell fills its
    first grid column and leaves the rest blank instead of repeating its text."""
    rows: list[list[object]] = []
    for tr in table_element.iter(f"{W_NS}tr"):
        row: list[object] = []
        for tc in tr.findall(f"{W_NS}tc"):
            span = 1
            properties = tc.find(f"{W_NS}tcPr")
            if properties is not None:
                grid = properties.find(f"{W_NS}gridSpan")
                if grid is not None:
                    try:
                        span = max(1, int(grid.get(f"{W_NS}val", "1")))
                    except ValueError:
                        span = 1
            row.append(docx_cell_text(tc))
            row.extend([""] * (span - 1))
        rows.append(row)
    return rows


def docx_tables(path: Path) -> list[tuple[str, list[list[object]], str]]:
    """Tables in document order, each preceded by the paragraphs written since the
    previous table (the toolkit's "Name of LG ... Name of SCHOOL ..." banners), with
    the checklist section the table sits in: Health centre, School, or blank."""
    document = Document(path)
    tables: list[tuple[str, list[list[object]], str]] = []
    pending: list[list[object]] = []
    section = ""
    index = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.replace(W_NS, "")
        if tag == "p":
            text = clean(xml_text(child))
            if not text:
                continue
            if HEALTH_SECTION.match(text):
                section = "Health centre"
            elif SCHOOL_SECTION.match(text):
                section = "School"
            if len(text) <= 200:
                pending.append([text])
        elif tag == "tbl":
            index += 1
            rows = docx_grid_rows(child)
            labels = label_rows(rows)
            if labels:
                # A two-column details table ("Name of Local Government" | "Amuria")
                # is a banner for the asset tables that follow it.
                pending.extend(labels)
            tables.append((f"Table {index}", list(pending) + rows, section))
            pending = []
    return tables


LABEL_CELL = re.compile(
    r"(?i)^\s*(name\s+of\s+(the\s+)?)?(lg|l\.g\.?|local\s+government|district|health\s+facility|facility|school|"
    r"health\s+cent(re|er)(\s*iii)?|seed\s+school|facility\s+name|school\s+name)\s*:?\s*$"
)


def label_rows(rows: list[list[object]]) -> list[list[object]]:
    """Banner pseudo-rows from a details table of label | value pairs."""
    found: list[list[object]] = []
    if len(rows) > 40:
        return found
    for row in rows:
        cells = [clean(value) for value in row if clean(value)]
        if len(cells) != 2 or not LABEL_CELL.match(cells[0]) or is_placeholder(cells[1]):
            continue
        label = cells[0].rstrip(": ")
        if re.search(r"(?i)lg|local\s+government|district", label):
            found.append([f"NAME OF LG: {cells[1]}"])
        else:
            found.append([f"NAME OF FACILITY: {cells[1]}"])
    return found


def sheet_rows(path: Path) -> list[tuple[str, list[list[object]]]]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        loaded = []
        try:
            for name in book.sheet_names():
                sheet = book.sheet_by_name(name)
                rows = []
                for row_index in range(sheet.nrows):
                    values = []
                    for col in range(sheet.ncols):
                        cell = sheet.cell(row_index, col)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            try:
                                values.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                            except Exception:
                                values.append(cell.value)
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                            values.append(None)
                        elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                            values.append(int(cell.value))
                        else:
                            values.append(cell.value)
                    rows.append(values)
                loaded.append((name, rows))
        finally:
            book.release_resources()
        return loaded
    workbook = load_workbook(path, read_only=True, data_only=True)
    loaded = []
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            loaded.append((name, [list(row) for row in sheet.iter_rows(values_only=True)]))
    finally:
        workbook.close()
    return loaded


def classify_header(row: list[object]) -> dict[str, int] | None:
    """Map a template header row (Equipment/Item, Department, ...) to field names.

    Header words are compared without spaces because Word tables often break a
    word across runs ("Equipme nt Item", "Item descriptio n", "Date of purcha se").
    """
    keys = [squash(value) for value in row]
    if not any(keys):
        return None
    has_equipment = any(key.startswith("equipmentitem") or key == "equipment" for key in keys)
    mapping: dict[str, int] = {}
    fallback_remarks: int | None = None
    for index, key in enumerate(keys):
        if not key or key in INDEX_KEYS:
            continue
        text = clean(row[index])
        if re.search(r"supplied|delivered|calloff|ordered|deficit|expected|requisition", key):
            # What was supplied or ordered is not what was found.
            continue
        field_name = None
        if key.startswith("equipmentitem") or key.startswith("equipmentname") or key in ITEM_KEYS:
            # "Equipment / Item Description", "Equipment Name", "Asset Description":
            # the item column, whatever sits in the line-number column before it.
            field_name = "item"
        elif "status" in key and ("note" in key or "location" in key):
            # "Status / Location / Notes" is free text; a plain status column wins.
            fallback_remarks = index if fallback_remarks is None else fallback_remarks
        elif "status" in key or key == "condition" or "functionality" in key:
            field_name = "status"
        elif "remark" in key:
            field_name = "remarks"
        elif "tag" in key or "engrave" in key:
            field_name = "tag"
        elif key.startswith("life") or "lifeinmonth" in key:
            field_name = "life"
        elif "purcha" in key or "datepurchased" in key or key.startswith("dateofpur"):
            field_name = "purchase"
        elif "placed" in key:
            field_name = "service"
        elif "recover" in key:
            field_name = "recoverable"
        elif "accdep" in key or "accumulated" in key:
            field_name = "acc_dep"
        elif "netbook" in key:
            field_name = "nbv"
        elif "ytd" in key or "deprn" in key or "depreciation" in key:
            field_name = "ytd"
        elif "assetnumber" in key or key == "assetno":
            field_name = "asset_number"
        elif "descriptio" in key:
            field_name = "description"
        elif key.startswith("equipmentitem") or key == "equipment" or key == "equipmentasset":
            field_name = "item"
        elif key == "item":
            field_name = "description" if has_equipment else "item"
        elif "department" in key:
            field_name = "department"
        elif key in {"qty", "quantity"} or key.startswith(("qty", "quantity", "numberverified", "noverified", "physicalcount", "countverified")) or key.endswith("qty"):
            field_name = "explicit_qty"
        elif "serial" in key:
            field_name = "serial"
        elif key.startswith("manufactur") or key == "make":
            field_name = "manufacturer"
        elif key.startswith("model"):
            field_name = "model"
        elif key in {"cost", "initialcost", "costugx"} or key.startswith("cost"):
            field_name = "cost"
        elif "localgov" in key or key in {"district", "localgovernment", "localgovt"}:
            field_name = "lg"
        elif any(token in key for token in ("healthcentre", "healthcenter", "hospital", "school", "location", "facility")):
            field_name = "facility"
        elif key in {"education", "educational", "health", "edn"} and index == mapping.get("item", -2) + 1:
            # The Department header cell overtyped with the department itself.
            field_name = "department"
        elif key.endswith("district") or resolve_lg(text, fuzzy=False) is not None:
            # A consolidation sheet that typed the block's district into the header cell.
            field_name = "lg"
        elif looks_like_facility(text) and facility_kind(text):
            field_name = "facility"
        if field_name and field_name not in mapping:
            mapping[field_name] = index
            if field_name == "facility" and facility_kind(text) and not re.search(r"(?i)location|physical", text):
                # "HEALTH CENTRE" / "SCHOOL" over the column: its bare names are that kind.
                mapping["facility_header_kind"] = facility_kind(text)
    if "remarks" not in mapping and fallback_remarks is not None:
        mapping["remarks"] = fallback_remarks
    if "item" not in mapping and "description" in mapping and (not keys[0] or keys[0] in INDEX_KEYS or 0 in mapping.values()):
        # The description column beside a line-number column names the item.
        mapping["item"] = mapping.pop("description")
    if "item" not in mapping and "department" in mapping and ("asset_number" in mapping or "tag" in mapping) \
            and len(mapping) >= 5 and 0 not in mapping.values() and keys[0] not in INDEX_KEYS:
        # A header whose first cell was overtyped ("MCH" for "Equipment/ Item").
        mapping["item"] = 0
    if "item" in mapping and len([name for name in mapping if name != "facility_header_kind"]) >= 4:
        return mapping
    return None


# Line-number columns: never the item, never a count.
INDEX_KEYS = {
    "vote", "itemno", "sn", "sno", "srno", "slno", "no", "nos", "number", "pgno", "s", "sr", "sl", "count",
    "line", "id", "ref", "sr.no", "itemnumber", "lineno", "num",
}
TOTAL_ITEM = re.compile(r"(?i)^\s*(?:sub|grand)?\s*totals?\b")
ITEM_KEYS = {
    "equipment", "equipmentasset", "equipments", "equipmentsitem", "assetequipment", "assetdescription",
    "assetsdescription", "itemname", "nameofitem", "assetname", "nameofasset", "equipmentdescription",
    "nameofequipment", "descriptionofasset", "descriptionofequipment",
}


def looks_like_facility(text: str) -> bool:
    """A named school or health facility: a type word alone ("Health", "Education",
    "Seed school"), a count ("2 at school") or a department is not a facility name."""
    text = clean(text)
    if len(text) < 4 or len(text) > 80 or is_placeholder(text) or re.match(r"\d", text):
        return False
    if re.search(r"(?i)regional blood bank|\bhospital\b|ministry of", text):
        return True
    if not facility_kind(text):
        return False
    if HEALTH_SECTION.match(text) or SCHOOL_SECTION.match(text) or re.search(
        r"(?i)\b(checklist|interview|discussion|verification|furniture|ict items|buildings to|equipment)\b", text
    ):
        return False
    if re.search(
        r"(?i)\b(fences?|blocks?|tanks?|latrines?|toilets?|desks?|chairs?|tables?|beds?|stools?|kits?|sets?|machines?|"
        r"computers?|printers?|uniforms?|textbooks?|books?|shelves|shelf|thermometers?|instrument|basic|pit|kitchen|bins?)\b",
        text,
    ) and not re.search(r"(?i)name\s+of", text):
        # "School fence", "Instrument set, ENT Basic for HCIII": an item, not a place.
        return False
    base = facility_base(text)
    if re.search(r"\d{3,}", base) or re.match(r"(ugift|team|sheet|table|copy|final|residential|non residential|buildings?)\b", base):
        return False
    if len(base.split()) > 6:
        return False
    return len(base) >= 3 and base not in {
        "education", "educational", "health", "edn", "primary", "secondary", "seed", "general", "the", "at",
        "furniture", "items", "buildings", "name", "facility", "ugift", "ugift health", "ugift health 2",
    }


def strip_lg_prefix(text: str) -> tuple[str, str]:
    """'LIRA DISTRICT LOCAL GOVERNMENT ANYOMOREM HEALTH CENTRE III' -> (lg, facility)."""
    match = re.match(
        r"(?i)^\s*(?:name\s+of\s+lg\s*:?\s*)?([A-Za-z][A-Za-z' \-]{2,30}?\s+(?:district\s+local\s+gov(?:ernment|'?t)|"
        r"district|dlg|municipal\s+council|municipality|city\s+council|city|mc|lg|dc))\b[\s,:\-–]*(.+)$",
        text,
    )
    if match and resolve_lg(match.group(1)) and looks_like_facility(match.group(2)):
        return match.group(1), clean(match.group(2))
    return "", text


PLACE_HEADER = re.compile(r"(?i)district|local\s*gov|facility|school|health|location|hospital|lg\b|hc\b|centre|center")


def trailing_place_columns(rows: list[list[object]], mapping: dict[str, int], header: list[object] | None = None) -> dict[str, int]:
    """Map unlabeled columns that hold local government and facility names.

    Only a column whose header is blank or names a place is considered, and it
    counts as the local-government column only when most of its values name a
    known local government, and as the facility column only when most of its
    values read as a school or health facility. Item names, counts, amounts, and
    section labels such as "ENT Basic for HCIII" do not qualify.
    """
    if "lg" in mapping and "facility" in mapping:
        return mapping
    used = set(mapping.values())
    if header is not None:
        for index, value in enumerate(header):
            text = clean(value)
            if text and not PLACE_HEADER.search(text) and resolve_lg(text, fuzzy=False) is None \
                    and not (looks_like_facility(text) and facility_kind(text)) and not bare_place_name(text):
                # A place name typed into the header cell ("KYIKWAKYA") still heads a
                # place column; any other heading rules the column out.
                used.add(index)
    lg_scores: Counter[int] = Counter()
    facility_scores: Counter[int] = Counter()
    bare_scores: Counter[int] = Counter()
    filled: Counter[int] = Counter()
    for row in rows[:120]:
        for index, value in enumerate(row):
            text = clean(value)
            if index in used or not text:
                continue
            filled[index] += 1
            if len(text) > 80:
                continue
            if resolve_lg(text) is not None:
                lg_scores[index] += 1
            elif looks_like_facility(text):
                facility_scores[index] += 1
            header_text = clean(header[index]) if header is not None and index < len(header) else ""
            if bare_place_name(text) and (not header_text or bare_place_name(header_text)):
                # An unlabeled column of bare names beside the district column ("MASAKA"
                # is a health centre here even though a district shares the name).
                bare_scores[index] += 1
    updated = dict(mapping)
    if "lg" not in updated:
        best = [index for index, count in lg_scores.most_common() if count >= 3 and count >= 0.6 * filled[index]]
        if best:
            updated["lg"] = best[0]
    if "facility" not in updated:
        best = [
            index for index, count in facility_scores.most_common()
            if count >= 3 and count >= 0.6 * filled[index] and index != updated.get("lg")
        ]
        if best:
            updated["facility"] = best[0]
        elif "lg" in updated:
            near = [
                index for index, count in bare_scores.most_common()
                if count >= 3 and count >= 0.6 * filled[index] and index > updated["lg"] and index <= updated["lg"] + 2
            ]
            if near:
                updated["facility"] = near[0]
                updated["facility_header_kind"] = ""
    return updated


NAME_OF_LG = (
    r"(?:(?:name\s+of\s+(?:the\s+)?(?:lg|l\.g\.?|local\s+government|district)|district|local\s+government)\s*[:\-]\s*"
    r"|name\s+of\s+(?:the\s+)?(?:lg|l\.g\.?|local\s+government|district)\s*)"
)
NAME_OF_FACILITY = (
    r"(?:(?:name\s+of\s+(?:the\s+)?(?:health\s+)?(?:facility|school|health\s+cent(?:re|er)|seed\s+school|h/?c(?:\s*iii)?)|"
    r"facility\s+name|health\s+facility|facility|school\s+name|school|health\s+cent(?:re|er))\s*[:\-]\s*|"
    r"name\s+of\s+(?:the\s+)?(?:health\s+)?(?:facility|school|health\s+cent(?:re|er)|seed\s+school|h/?c(?:\s*iii)?)\s*)"
)


def banner_part(text: str) -> str:
    """A name written on a dotted template line, without the dots and placeholders."""
    text = clean(re.sub(r"[.…_]{2,}|(?<=\s)[.…_](?=\s)", " ", text)).strip(" :-–—.,")
    # The name ends where the next template heading starts ("... SCHOOL VERIFICATION
    # CHECKLIST NAME OF THE SCHOOL ...", "... NB: ...").
    text = re.split(
        r"(?i)\s+(?:school|health\s+cent(?:re|er))\s+verification\s+checklist|\s+checklist\b|\s+name\s+of\s+|\s+n\.?b\.?\s*:|\s+notes?\s*:",
        text,
    )[0].strip(" :-–—.,")
    if is_placeholder(text):
        return ""
    # "MAMBA SEED SECONDARY SCHOOL – NEBBI DLG" names the government after a dash.
    return text


def split_trailing_lg(facility: str) -> tuple[str, str]:
    match = re.search(r"^(.*?)[\s,]+[–\-]\s*([A-Za-z][A-Za-z .'\-]{2,40}(?:\bdlg|\bdistrict|\bmc|\bcity|\blg|\bdc)\.?)\s*$", facility, re.I)
    if match and resolve_lg(match.group(2)):
        return clean(match.group(1)), clean(match.group(2))
    # "KYEIHARA HC III MITOOMA": the last word or two name the local government.
    words = clean(facility).split()
    for size in (2, 1):
        if len(words) > size + 1:
            head, tail = " ".join(words[:-size]), " ".join(words[-size:])
            if resolve_lg(tail, fuzzy=False) and facility_kind(head) and looks_like_facility(head):
                return head, tail
    return facility, ""


def parse_banner(text: str) -> tuple[str, str] | None:
    raw = clean(text)
    if len(raw) < 6 or is_placeholder(raw):
        return None
    low = raw.casefold()
    asset_word = re.search(
        r"\b(machine|desks?|chairs?|stools?|autoclave|bowl|computer|monitor|cupboard|printer|tables?|"
        r"set|instrument|beds?|mask|kit|apparatus|shelves|screen|projector|trolley|stand|scale|meter)\b",
        low,
    )
    if asset_word and "name of" not in low and "ugift asset" not in low and not re.search(r"(?i)\b(district|dlg|d\.?c\.?|mc|city)\b", low):
        return None
    match = re.search(rf"(?i){NAME_OF_LG}(.*?)\s*{NAME_OF_FACILITY}(.*)$", raw)
    if match:
        facility, extra = split_trailing_lg(banner_part(match.group(2)))
        return banner_part(match.group(1)) or extra, facility
    match = re.search(rf"(?i){NAME_OF_FACILITY}(.*)$", raw)
    if match:
        facility, extra = split_trailing_lg(banner_part(match.group(1)))
        return extra, facility
    match = re.search(rf"(?i){NAME_OF_LG}(.*)$", raw)
    if match:
        return banner_part(match.group(1)), ""
    match = re.search(r"(?i)^(.+?)\s+ugift\s+asset\s+verification.*?,\s*(.+)$", raw)
    if match:
        return clean(match.group(2)), clean(match.group(1))
    match = re.search(r"(?i)^(.+?)\s+ugift\s+asset\s+verification\s+tool\s*kit\s+(.+)$", raw)
    if match:
        return clean(match.group(2)), clean(match.group(1))
    if "," in raw and re.search(r"(?i)hc\s*(?:ii+|iv)?|health|school|\bsss\b|seed", raw):
        # "BITSYA HCIII, BUHWEJU DC" names the facility, then its government. An item
        # such as "Instrument set, ENT Basic for HCIII" has no government after its comma.
        parts = [clean(part) for part in raw.split(",") if clean(part)]
        if len(parts) >= 2 and resolve_lg(parts[-1]) is not None and looks_like_facility(parts[0]):
            return parts[-1], parts[0]
    match = re.search(
        r"(?i)^(.+?\b(?:HC\s*III|HCIII|HEALTH\s+CENT(?:RE|ER).*|SEED(?:\s+SECONDARY)?\s+SCHOOL|SSS))\s+(.+\bDISTRICT)\s*$",
        raw,
    )
    if match:
        return clean(match.group(2)), clean(match.group(1))
    if re.fullmatch(
        r"(?i)[A-Z0-9][A-Z0-9 .'/()-]{2,80}(?:SEED(?:\s+SECONDARY)?\s+SCHOOL|HEALTH\s+CENT(?:RE|ER)(?:\s+[IVX0-9]+)?|HC\s*III|HCIII|SSS)",
        raw,
    ):
        return "", raw
    # A cell holding nothing but a facility name ("ADEKNINO HC II", "Kigumba S.S.S").
    if len(raw) <= 60 and looks_like_facility(raw) and not re.search(r"(?i)\b(machine|set|kit|bin|bed|chair|desk|stool|table)\b", raw):
        return "", raw
    return None


def first_text(row: list[object]) -> str:
    for value in row:
        text = clean(value)
        if text:
            return text
    return ""


def row_text(row: list[object]) -> str:
    """Every filled cell of a row in one string, so a banner split over cells
    ("NAME OF LG:" | "DOKOLO" | "NAME OF HEALTH FACILITY:" | "ADOK HC III") reads whole."""
    return clean(" ".join(clean(value) for value in row if clean(value)))


def bare_integer(value: object) -> int | None:
    """A cell holding only a whole number. Excel turns a typed "(3)" into -3, so the
    sign is dropped."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return abs(int(value))
    text = clean(value).replace(",", "")
    if re.fullmatch(r"[-(]?\d{1,4}\)?", text):
        return abs(int(re.sub(r"[()\s]", "", text)))
    # "2+3 = 5" or "2 + 3": the count written as a sum of the per-room counts.
    if re.fullmatch(r"[\d\s+=()\-]+", text) and re.search(r"\d", text):
        numbers = [int(number) for number in re.findall(r"\d{1,4}", text)]
        if "=" in text:
            return numbers[-1]
        return sum(numbers) if "+" in text else numbers[0]
    return None


def placeholder_only(asset: Asset) -> bool:
    """A template line with nothing but the item name: every other field blank, a
    dash, or dots. It records no asset."""
    fields = [
        asset.department, asset.asset_number, asset.description, asset.tag, asset.status, asset.remarks,
        asset.life, asset.purchase, asset.service, asset.recoverable, asset.cost, asset.acc_dep, asset.nbv, asset.ytd,
    ]
    for value in fields:
        if value in (None, ""):
            continue
        if isinstance(value, str) and (is_placeholder(value) or re.fullmatch(r"[\-–—.…_/ ]+", value)):
            continue
        return False
    return True


FRAGMENT = re.compile(r"(?i)^(?:[\(\[]?\d+[\)\]]?\s+\w+|one|two|three|four|five|six|[a-z(]).{0,40}$")
NOT_RECEIVED = re.compile(r"(?i)^(?:didn'?t\s+receive|not\s+received|none\s+received|nil)\b")
# A department or room written where the item name goes, on a continuation row:
# "MCH (2)", "OPD(1) STORE (1)", "General ward".
DEPARTMENT_WORD = re.compile(
    r"(?i)^(?:(?:mch|opd|ipd|store|stores|lab|laboratory|maternity|general|ward|theatre|theater|kitchen|office|"
    r"admin|administration|dispensary|pharmacy|records|reception|staff\s*room|library|classroom|dormitory|"
    r"emergency|paediatric|pediatric|antenatal|postnatal|delivery|injection|dressing|treatment|inpatient|outpatient)"
    r"\s*(?:ward|room|block|unit)?\s*[\(\[]?\s*\d{0,3}\s*[\)\]]?\s*[,;/&+]?\s*){1,4}$"
)


def cell(row: list[object], mapping: dict[str, int], name: str) -> object:
    index = mapping.get(name)
    if index is None or index >= len(row):
        return None
    return row[index]


def looks_like_subheader(asset: Asset) -> bool:
    return norm(asset.description) == "description" and not any([
        asset.item, asset.tag, asset.status, asset.remarks, asset.cost, asset.asset_number,
    ])


def looks_like_label(asset: Asset, row: list[object] | None = None) -> bool:
    if asset.tag or asset.status or asset.cost not in (None, "") or asset.asset_number:
        return False
    text = asset.item or asset.description
    if not text and row is not None:
        filled = [clean(value) for value in row if clean(value)]
        text = filled[0] if len(filled) == 1 else ""
    if not text or asset.department:
        return False
    banner = parse_banner(text)
    return banner is not None and bool(banner[0] or banner[1])


def facility_type_for(facility: str, department: str, sheet: str, filename: str) -> str:
    """'School' or 'Health centre', the only two facility types the register uses."""
    kind = facility_kind(facility)
    if kind:
        return kind
    if re.search(r"(?i)blood\s*bank|hospital|clinic|referral", facility):
        # A regional blood bank or hospital is a health facility in the two-kind scheme.
        return "Health centre"
    kind = facility_kind(department)
    if kind:
        return kind
    if re.search(r"(?i)\beducation(al)?\b|\bedn\b|academics?|classroom|laborator|library|dormitor", department):
        return "School"
    if re.search(r"(?i)\bopd\b|\bmch\b|maternity|\bward\b|\bipd\b|theatre|dispens", department):
        return "Health centre"
    broader = f"{sheet} {filename}"
    if re.search(r"health|hospital|\bhc\b|hciii", broader, re.I) and not re.search(r"school|\bsss?\b|seed", broader, re.I):
        return "Health centre"
    if re.search(r"school|education|\bsss?\b|seed", broader, re.I) and not re.search(r"health|hospital|\bhc\b|hciii", broader, re.I):
        return "School"
    return ""


def join_text(left: str, right: str, divider: str = " ") -> str:
    left, right = clean(left), clean(right)
    if not right or right.casefold() in left.casefold():
        return left
    if not left:
        return right
    if left.endswith(("/", "-")):
        return f"{left}{right}"
    return f"{left}{divider}{right}"


def same_place(left: str, right: str) -> bool:
    """Two spellings of one facility name: equal bases, a few edits apart with the
    same opening, or the same letters in another order."""
    a, b = facility_base(left), facility_base(right)
    if not a or not b:
        return False
    if a == b or sorted(a) == sorted(b):
        return True
    if a[:1] != b[:1]:
        return False
    distance = _edit_distance(a, b)
    prefix = len(os.path.commonprefix([a, b]))
    return distance <= 2 or (distance == 3 and prefix >= 4 and min(len(a), len(b)) >= 6)


def bare_place_name(text: str) -> bool:
    """A value of a labelled facility column without a type word ("BUTAWATA" under
    a HEALTH CENTRE header); status, department and header words do not qualify."""
    text = clean(text)
    if not text or is_placeholder(text) or len(text.split()) > 4 or re.search(r"\d", text):
        return False
    if DEPARTMENT_WORD.match(text) or re.match(
        r"(?i)^(?:education(?:al)?|health|hospital|school|facility|location|name|n/?a|none|nil|yes|no|functional|"
        r"faulty|good|new|not|non|in\s+use|verified|broken|damaged|missing|total)(?![A-Za-z])",
        text,
    ):
        return False
    return re.fullmatch(r"[A-Za-z][A-Za-z .'\-/&]{2,45}", text) is not None


def unit_row(asset: Asset) -> bool:
    """A row without an item name that still records one unit: a description together
    with a department, tag, status, remark, cost or date. A wrapped line of the row
    above carries only a fragment of one cell."""
    if not asset.description:
        return False
    if re.match(r"^[a-z]", asset.description) and len(asset.description) < 25 and not asset.tag and not asset.status:
        return False
    others = sum(1 for value in (asset.department, asset.tag, asset.status, asset.remarks) if value)
    others += sum(1 for value in (asset.cost, asset.purchase, asset.service) if value not in (None, ""))
    return others >= 1


def merge_continuation(previous: Asset, current: Asset) -> None:
    previous.department = join_text(previous.department, current.department, "; ")
    previous.description = join_text(previous.description, current.description)
    previous.tag = join_text(previous.tag, current.tag)
    previous.status = join_text(previous.status, current.status)
    previous.remarks = join_text(previous.remarks, current.remarks)
    for name in ("life", "purchase", "service", "recoverable", "cost", "acc_dep", "nbv", "ytd", "asset_number"):
        if getattr(previous, name) in (None, "") and getattr(current, name) not in (None, ""):
            setattr(previous, name, getattr(current, name))
    if not previous.lg and current.lg:
        previous.lg = current.lg
    if not previous.facility and current.facility:
        previous.facility = current.facility


def take_asset(row: list[object], mapping: dict[str, int], source: str, location: str) -> Asset:
    quantity = number_or_text(cell(row, mapping, "explicit_qty"))
    if isinstance(quantity, str):
        # "2 pcs", "120 desks": the number in a quantity cell written with a unit.
        match = re.match(r"\s*(\d{1,3})\b(?!\s*(?:months?|yrs?|years?|inch|%|/))", quantity)
        quantity = int(match.group(1)) if match else None
    explicit = quantity if isinstance(quantity, int) and 0 < quantity <= MAX_QUANTITY else None
    cost = number_or_text(cell(row, mapping, "cost"))
    unit_cost = number_or_text(cell(row, mapping, "unit_cost")) if "unit_cost" in mapping else None
    asset = Asset(
        item=clean(cell(row, mapping, "item")),
        department=clean(cell(row, mapping, "department")),
        asset_number=clean(cell(row, mapping, "asset_number")),
        description=clean(cell(row, mapping, "description")),
        life=number_or_text(cell(row, mapping, "life")),
        tag=clean(cell(row, mapping, "tag")),
        purchase=excel_date(cell(row, mapping, "purchase")),
        service=excel_date(cell(row, mapping, "service")),
        recoverable=number_or_text(cell(row, mapping, "recoverable")),
        cost=cost if cost not in (None, "") else unit_cost,
        acc_dep=number_or_text(cell(row, mapping, "acc_dep")),
        nbv=number_or_text(cell(row, mapping, "nbv")),
        ytd=number_or_text(cell(row, mapping, "ytd")),
        status=clean(cell(row, mapping, "status")),
        remarks=clean(cell(row, mapping, "remarks")),
        lg=clean(cell(row, mapping, "lg")),
        facility=clean(cell(row, mapping, "facility")),
        explicit_qty=explicit,
        source_file=source,
        source_location=location,
    )
    if norm(asset.item) in {"equipment item", "equipment", "description", "asset description", "item", "facility fitting installation"}:
        # A repeated header line, not an asset.
        asset.item = ""
        asset.extras["header_repeat"] = True
    elif asset.item and (is_placeholder(asset.item) or TOTAL_ITEM.match(asset.item)):
        # A dash, "N/A" or a "TOTAL ..." line names no asset.
        asset.item = ""
        asset.extras["placeholder_item"] = True
    # A merged Word cell read once per spanned column ("KYESS KYESS KYESS").
    asset.tag = re.sub(r"(?i)^(\S+)(?:\s+\1)+$", r"\1", asset.tag)
    # Model, serial and manufacturer columns are kept as stated facts in the description.
    for role, label in (("model", "Model"), ("serial", "Serial number"), ("manufacturer", "Made by")):
        value = clean(cell(row, mapping, role)) if role in mapping else ""
        if value and not is_placeholder(value):
            asset.description = join_text(asset.description, f"{label} {value}", "; ")
    if (cost in (None, "") and unit_cost not in (None, "")) or ("unit_cost" in mapping and "cost" not in mapping):
        # A unit price is already one item's share: it is not divided by the count.
        asset.extras["unit_cost"] = True
    if "explicit_qty" in mapping:
        asset.extras["qty_column"] = True
    return asset


def sheet_is_duplicate_draft(name: str, mapping: dict[str, int], workbook_has_places: bool) -> bool:
    """Skip OCR fragments and single-facility drafts when a consolidated sheet exists."""
    if not workbook_has_places:
        return False
    if re.fullmatch(r"table\s*[1-3]", name.strip(), flags=re.I):
        return True
    unnamed = re.fullmatch(r"sheet\s*[1-3]", name.strip(), flags=re.I)
    return bool(unnamed and "lg" not in mapping and "facility" not in mapping)


def parse_template_sheet(
    sheet_name: str, rows: list[list[object]], source: str, filename: str, fallback_lg: str,
) -> list[Asset]:
    mapped_rows: list[tuple[int, dict[str, int]]] = []
    for index, row in enumerate(rows):
        mapping = classify_header(row)
        if mapping:
            mapped_rows.append((index, trailing_place_columns(rows[index + 1:index + 40], mapping, row)))
    if not mapped_rows:
        return []
    assets: list[Asset] = []
    carry_lg = fallback_lg
    carry_facility = ""
    carry_column = ""
    mapping = None
    header_indexes = {index for index, _ in mapped_rows}
    active = {index: item for index, item in mapped_rows}
    block_last: Asset | None = None
    pending: Asset | None = None
    for row_number, row in enumerate(rows, 1):
        if (row_number - 1) in header_indexes:
            mapping = active[row_number - 1]
            block_last = None
            pending = None
            # The facility carries across the furniture, ICT and buildings sub-tables of
            # one block until a banner, label row or place column names another.
            # A banner written on the rows just above the header names the block. A
            # data row of the previous block is not a banner.
            for above in rows[max(0, row_number - 9):row_number - 1]:
                filled = [clean(value) for value in above if clean(value)]
                text = row_text(above)
                if len(filled) > 4 and not re.search(r"(?i)name\s+of", text):
                    continue
                banner = parse_banner(text)
                if banner:
                    carry_lg = banner[0] or carry_lg
                    carry_facility = banner[1] or carry_facility
            continue
        if mapping is None:
            banner = parse_banner(row_text(row))
            if banner:
                carry_lg = banner[0] or carry_lg
                carry_facility = banner[1] or carry_facility
            continue
        if not any(clean(value) for value in row):
            continue
        asset = take_asset(row, mapping, source, f"{sheet_name} row {row_number}")
        if asset.extras.get("placeholder_item"):
            # A row whose item cell is a dash or "N/A": a ruled-out template line.
            continue
        if looks_like_subheader(asset) or asset.extras.get("header_repeat"):
            # The "Description" sub-header row may carry the block's place columns.
            if asset.lg and resolve_lg(asset.lg):
                carry_lg = asset.lg
            if asset.facility and (looks_like_facility(asset.facility) or ("facility" in mapping and bare_place_name(asset.facility))):
                if not (carry_facility and same_place(asset.facility, carry_facility)):
                    carry_facility = asset.facility
            continue
        continuation = not asset.item and block_last is not None and (asset.description or asset.remarks or asset.tag or asset.status)
        if not continuation and looks_like_label(asset, row):
            banner = parse_banner(row_text(row)) or parse_banner(asset.item or first_text(row))
            if banner:
                carry_lg = banner[0] or carry_lg
                carry_facility = banner[1] or carry_facility
                block_last = None
            continue
        if asset.lg and (resolve_lg(asset.lg) or lg_from_text(asset.lg)):
            new_lg, old_lg = resolve_lg(asset.lg), resolve_lg(carry_lg) if carry_lg else None
            if new_lg is not None and old_lg is not None and new_lg.key != old_lg.key:
                # A new district block: the facility of the block above does not carry.
                carry_facility, carry_column = "", ""
            carry_lg = asset.lg
        else:
            asset.lg = carry_lg
        header_kind = mapping.get("facility_header_kind") or ("" if "facility" not in mapping else facility_type_for("", "", sheet_name, filename))
        if asset.facility and carry_facility and same_place(asset.facility, carry_facility):
            # "ALAORMIT HC III" under the banner "AKOROMIT HC III": one facility, the
            # banner's spelling.
            asset.facility = carry_facility
        if asset.facility and (looks_like_facility(asset.facility) or ("facility" in mapping and bare_place_name(asset.facility))):
            carry_facility = asset.facility
            carry_column = header_kind if "facility" in mapping and not looks_like_facility(asset.facility) else ""
        elif asset.facility and "facility" in mapping and not is_placeholder(asset.facility):
            # The column's own value stands, whatever the tests say of it.
            carry_facility, carry_column = asset.facility, header_kind
        else:
            asset.facility = carry_facility
        if carry_column and asset.facility == carry_facility:
            asset.extras["column_facility"] = carry_column
        count = bare_integer(asset.item) if asset.item else None
        department_word = bool(asset.item) and DEPARTMENT_WORD.match(asset.item) is not None
        # "Classroom Block" with its own cost and status is a building, not the
        # department of the row above.
        department_row = department_word and not (
            asset.status or asset.cost not in (None, "") or asset.purchase not in (None, "") or (asset.tag and not blank_tag(asset.tag))
        )
        if pending is not None and (count is not None or (department_word or not asset.item) and not placeholder_only(asset)):
            department_row = department_word
            # Count-on-next-row layout: the item name sat alone on the previous row and
            # this row starts with the count (or the department), then the item's data.
            if department_row:
                asset.department = join_text(asset.department, asset.item, "; ")
            asset.item = ""
            if count is not None and 0 < count <= MAX_QUANTITY:
                pending.explicit_qty = count
            merge_continuation(pending, asset)
            pending.facility_type = facility_type_for(pending.facility, pending.department, sheet_name, filename)
            assets.append(pending)
            block_last = pending
            pending = None
            continue
        if pending is not None:
            if not placeholder_only(pending):
                # A filled asset held back for a count row that never came.
                pending.facility_type = facility_type_for(pending.facility, pending.department, sheet_name, filename)
                assets.append(pending)
                block_last = pending
            # The lone item name was a section label, a template line the facility
            # never received, or a fragment of the previous item's text.
            elif block_last is not None and FRAGMENT.match(pending.item) and squash(pending.item) not in TOOLKIT_INDEX:
                block_last.description = join_text(block_last.description, pending.item)
            pending = None
        if count is not None and (asset.status or asset.cost not in (None, "") or asset.purchase not in (None, "")):
            # A line number in the item column of a row that carries its own status,
            # cost or date: the row is an asset named by its description, not a count.
            if not asset.description:
                continue
            asset.item, asset.description = asset.description, ""
            count = None
        if (count is not None or department_row) and block_last is not None:
            if department_row:
                asset.department = join_text(asset.department, asset.item, "; ")
            asset.item = ""
            if count is not None and block_last.explicit_qty is None and 0 < count <= MAX_QUANTITY:
                block_last.explicit_qty = count
            merge_continuation(block_last, asset)
            continue
        if not asset.item and block_last is not None and unit_row(asset):
            # One unit per row under the item's name: "Drip Stand" on the first row,
            # then a row per stand with its own description, tag or status.
            asset.item = block_last.extras.get("unit_item") or identity_item(block_last.item)
            asset.extras["unit_row"] = True
            asset.extras["unit_item"] = asset.item
            if not block_last.extras.get("unit_row"):
                block_last.extras["has_unit_rows"] = True
            if not asset.department:
                asset.department = block_last.department
            asset.facility_type = facility_type_for(asset.facility, asset.department, sheet_name, filename)
            assets.append(asset)
            block_last = asset
            continue
        if not asset.item and block_last is not None and (asset.description or asset.tag or asset.status or asset.remarks or asset.department):
            merge_continuation(block_last, asset)
            continue
        if not asset.item:
            continue
        if NOT_RECEIVED.match(asset.item) and placeholder_only(asset):
            continue
        if placeholder_only(asset):
            if block_last is not None and squash(f"{block_last.item} {asset.item}") in TOOLKIT_INDEX:
                # "Instrument set, ENT" / "Basic for HCIII": one item name wrapped over
                # two rows; a count row may still follow it.
                block_last.item = clean(f"{block_last.item} {asset.item}")
                if assets and assets[-1] is block_last:
                    assets.pop()
                pending = block_last
                block_last = None
                continue
            pending = asset
            continue
        asset.facility_type = facility_type_for(asset.facility, asset.department, sheet_name, filename)
        assets.append(asset)
        block_last = asset
    return assets


def parse_qty_register(sheet_name: str, rows: list[list[object]], source: str, fallback_lg: str) -> list[Asset]:
    header_at = None
    mapping: dict[str, int] = {}
    for index, row in enumerate(rows[:15]):
        texts = [norm(value) for value in row]
        if "qty" in texts and "description" in texts and "condition" in texts:
            header_at = index
            for col, text in enumerate(texts):
                if text == "description":
                    mapping["item"] = col
                elif text in {"qty", "quantity"}:
                    mapping["explicit_qty"] = col
                elif "cost" in text:
                    mapping["cost"] = col
                elif "purchase" in text or text.startswith("date"):
                    mapping["purchase"] = col
                elif text == "user":
                    mapping["remarks"] = col
                elif text == "location":
                    mapping["facility"] = col
                elif text in {"department", "section"}:
                    mapping["department"] = col
                elif text == "condition":
                    mapping["status"] = col
                elif "asset number" in text or text == "no":
                    mapping["asset_number"] = col
            break
    if header_at is None:
        return []
    assets: list[Asset] = []
    facility = ""
    for row_number, row in enumerate(rows[header_at + 1:], header_at + 2):
        asset = take_asset(row, mapping, source, f"{sheet_name} row {row_number}")
        if asset.remarks and not re.search(r"\d", asset.remarks):
            asset.remarks = f"User: {asset.remarks}"
        filled = [clean(value) for value in row if clean(value)]
        if len(filled) == 1 and parse_banner(filled[0]):
            facility = filled[0]
            continue
        if not asset.item:
            continue
        asset.facility = asset.facility or facility
        asset.lg = asset.lg or fallback_lg
        asset.facility_type = facility_type_for(asset.facility, asset.department, sheet_name, source)
        assets.append(asset)
    return assets


def parse_pallisa(sheet_name: str, rows: list[list[object]], source: str) -> list[Asset]:
    header_at = None
    texts: list[str] = []
    for index, row in enumerate(rows[:12]):
        texts = [norm(value) for value in row]
        if any("asset description" in text or text == "description" for text in texts) and "condition" in texts:
            header_at = index
            break
    if header_at is None:
        return []
    columns = {text: index for index, text in enumerate(texts) if text}
    def col(*names: str) -> int | None:
        for name in names:
            if name in columns:
                return columns[name]
        for text, index in columns.items():
            if any(name in text for name in names):
                return index
        return None
    item_col = col("asset description", "description")
    tag_col = col("tag number", "tag")
    dept_col = col("section", "department-location", "location- department", "location-department")
    facility_col = col("location")
    cost_col = col("cost ugx", "cost")
    date_col = col("date of purchase", "date of acquisition")
    status_col = col("condition")
    user_col = col("user name")
    title_col = col("user title")
    assets: list[Asset] = []
    for row_number, row in enumerate(rows[header_at + 1:], header_at + 2):
        def value(index: int | None) -> object:
            if index is None or index >= len(row):
                return None
            return row[index]
        item = clean(value(item_col))
        if not item or norm(item) in {"description", "asset description"}:
            continue
        facility = clean(value(facility_col))
        department = clean(value(dept_col))
        if facility and not re.search(r"school|health|hc|hospital", facility, re.I) and re.search(r"school|health|hc", department, re.I):
            facility, department = department, facility
        user = clean(value(user_col))
        title = clean(value(title_col))
        remarks = ", ".join(part for part in (user, title) if part)
        asset = Asset(
            item=item,
            department=department,
            tag=clean(value(tag_col)),
            purchase=excel_date(value(date_col)),
            cost=number_or_text(value(cost_col)),
            status=clean(value(status_col)),
            remarks=remarks,
            lg="Pallisa",
            facility=facility,
            facility_type=facility_type_for(facility, department, sheet_name, "pallisa"),
            source_file=source,
            source_location=f"{sheet_name} row {row_number}",
        )
        if not asset.facility_type:
            asset.facility_type = "Seed school"
        assets.append(asset)
    return assets


def context_from_line(text: str) -> tuple[str, str] | None:
    """A line above or between asset rows that names the place: a banner, a known
    local government, or a facility name standing alone."""
    raw = clean(text)
    if not raw or len(raw) > 160 or raw.endswith(":") or is_placeholder(raw):
        return None
    if re.match(r"(?i)class\s+\d", raw) or re.match(r"(?i)cost cent|source\b|vote\b", raw):
        return None
    # "Local Government: Rubirizi District | Facility: Munyonyi HC III"
    pieces = [piece for piece in re.split(r"\s*\|\s*", raw) if piece]
    if len(pieces) > 1:
        lg_text = facility_text = ""
        for piece in pieces:
            banner = parse_banner(piece)
            if banner:
                lg_text = lg_text or banner[0]
                facility_text = facility_text or banner[1]
        if lg_text or facility_text:
            return lg_text, facility_text
    banner = parse_banner(raw)
    if banner and (banner[0] or banner[1]):
        return banner
    if len(raw) <= 60 and resolve_lg(raw) is not None:
        return raw, ""
    if len(raw) <= 60 and lg_from_text(raw) is not None and not facility_kind(raw):
        return raw, ""
    if len(raw) <= 80 and looks_like_facility(raw):
        return "", raw
    return None


def flexible_mapping(row: list[object]) -> dict[str, int] | None:
    texts = [norm(value) for value in row]
    if not any(texts):
        return None
    roles: dict[str, int] = {}
    descriptions: list[int] = []
    for index, text in enumerate(texts):
        if not text or text in {"sn", "s n", "no", "item no", "pg no"}:
            continue
        if "imei" in text or "serial" in text or "engrave" in text or text.startswith("tag"):
            roles.setdefault("tag", index)
        elif "remark" in text or text.endswith("notes") or "notes" in text:
            roles.setdefault("remarks", index)
        elif "status" in text or text == "condition" or "functionality" in text or text.startswith("functional"):
            roles.setdefault("status", index)
        elif "unit cost" in text or "unit price" in text or "cost per unit" in text:
            roles.setdefault("unit_cost", index)
        elif "total cost" in text or "total amount" in text:
            roles.setdefault("cost", index)
        elif ("cost" in text and "center" not in text and "centre" not in text and "control" not in text):
            roles.setdefault("cost", index)
        elif re.search(r"\b(supplied|delivered|call\s*off|ordered|deficit|register)\b", text):
            # What was supplied is not what was found.
            continue
        elif "qty" in text or "qtty" in text or "qnty" in text or "quantity" in text or text in {"found", "verified", "counted", "physical count", "number found"}:
            roles.setdefault("explicit_qty", index)
        elif text in {"department", "section"} or text.startswith("department"):
            roles.setdefault("department", index)
        elif "asset category" in text and "sub" not in text:
            roles.setdefault("department", index)
        elif "physical location" in text or text in {"facility", "school", "schools", "location"}:
            roles.setdefault("facility", index)
        elif "location" in text and "department" in text:
            roles.setdefault("department", index)
        elif text in {"district", "local government"} or "local gov" in text:
            roles.setdefault("lg", index)
        elif any(token in text for token in ("date of purchase", "date purchased", "date of acquisition", "delivery date")) or text.startswith("date of pur"):
            roles.setdefault("purchase", index)
        elif "placed" in text or "issue date" in text:
            roles.setdefault("service", index)
        elif text in {"equipment", "equipment item", "equipment name"} or text.startswith("equipment /") or text.startswith("equipment item"):
            roles.setdefault("item", index)
        elif "description" in text or text in {"asset", "equipment asset", "item", "items", "item name", "name of item", "asset name"}:
            descriptions.append(index)
        elif "model" in text:
            roles.setdefault("description", index)
    if "item" not in roles and descriptions:
        roles["item"] = descriptions.pop(0)
    if descriptions and "description" not in roles:
        roles["description"] = descriptions[0]
    # The prompt's test: a description column together with condition, quantity,
    # tag, or cost. A device list with serial numbers but no item column is not one.
    useful = {"tag", "status", "cost", "explicit_qty"}
    if "item" in roles and len(useful & set(roles)) >= 1 and len(roles) >= 3:
        return roles
    return None


def sheet_label(sheet_name: str, filename: str) -> str:
    blob = f"{sheet_name} {filename}".casefold()
    if "tela" in blob or "phone" in blob:
        return "Phone"
    if "tab" in blob:
        return "Tablet"
    if "inspection" in blob or "device" in blob:
        return "Inspection device"
    return ""


def looks_like_budget_codes(assets: list[Asset]) -> bool:
    if len(assets) < 8:
        return False
    codes = sum(1 for asset in assets if re.fullmatch(r"\d{5,8}", asset.item))
    return codes / len(assets) > 0.5


def parse_flexible_sheet(
    sheet_name: str, rows: list[list[object]], source: str, filename: str, fallback_lg: str,
) -> list[Asset]:
    header_at = None
    mapping: dict[str, int] = {}
    for index, row in enumerate(rows[:15]):
        found = flexible_mapping(row)
        if found:
            header_at = index
            mapping = found
            break
    if header_at is None:
        return []
    carry_lg = fallback_lg
    carry_facility = ""
    for row in rows[:header_at]:
        label_text = norm(row[0]) if row else ""
        value = clean(row[1]) if len(row) > 1 else ""
        if value and label_text in {"school name", "name of school", "facility", "health facility", "name of health facility"}:
            carry_facility = value
            continue
        if value and label_text in {"district", "name of lg", "local government"}:
            carry_lg = value
            continue
        context = context_from_line(first_text(row))
        if not context:
            continue
        carry_lg = context[0] or carry_lg
        carry_facility = context[1] or carry_facility
    assets: list[Asset] = []
    for row_number, row in enumerate(rows[header_at + 1:], header_at + 2):
        if not any(clean(value) for value in row):
            continue
        context = context_from_line(row_text(row))
        probe = take_asset(row, mapping, source, "")
        if context and not any([probe.tag, probe.status, probe.cost, probe.explicit_qty, probe.description]):
            carry_lg = context[0] or carry_lg
            carry_facility = context[1] or carry_facility
            continue
        asset = probe
        asset.source_location = f"{sheet_name} row {row_number}"
        if not asset.item or norm(asset.item) in {"description", "asset description", "equipment", "total"}:
            continue
        if placeholder_only(asset) and not asset.explicit_qty:
            continue
        if asset.department and not asset.facility and re.search(r"school|health|hospital|blood bank|\bhc\b", asset.department, re.I):
            asset.facility = asset.department
            asset.department = ""
        asset.lg = asset.lg or carry_lg
        asset.facility = asset.facility or carry_facility
        asset.facility_type = facility_type_for(asset.facility, asset.department, sheet_name, filename)
        assets.append(asset)
    if looks_like_budget_codes(assets):
        return []
    if not any(asset.tag or asset.status or asset.cost not in (None, "") or asset.explicit_qty or asset.description for asset in assets):
        return []
    return assets


def reconciliation_sources() -> dict[str, list[tuple[str, str, str]]]:
    """Source file -> [(local government, facility, type), ...] from the
    reconciliation. Used only after the row, the banner and the folder path have
    all failed to name the place, and only when one facility of the row's kind
    is tied to the file."""
    found: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    if not RECONCILIATION.exists():
        return {}
    with RECONCILIATION.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            name = row.get("field_name") or row.get("name") or ""
            kind = row.get("type") or ""
            if not name or kind not in {"School", "Health centre"}:
                continue
            for column in ("source", "additional_sources"):
                for source in (row.get(column) or "").split(";"):
                    source = source.strip()
                    place = (row.get("lg") or "", name, kind)
                    if source and place not in found[source]:
                        found[source].append(place)
    return dict(found)


RECONCILED: dict[str, list[tuple[str, str, str]]] = {}
PLACE_NOTES: Counter = Counter()


def only_of_kind(places, kind: str):
    """The single place of this kind in a list of (lg, facility) or (lg, facility, kind)."""
    matches = []
    for place in places:
        facility = place[1]
        place_kind = place[2] if len(place) > 2 else facility_kind(facility)
        if kind and place_kind == kind and facility:
            matches.append(place)
    return matches[0] if len(matches) == 1 else None


def resolve_places(asset: Asset, contexts: list[tuple[str, str]], relative: str) -> None:
    """Settle the local government and facility of one row.

    Order: the row's own values, then the banner already carried into the row,
    then the folder path (every folder a byte-identical copy was filed under, matched
    to the table's facility kind), then the reconciliation file when it ties this
    source to one facility of that kind. A banner that names another facility than the
    folder of the same kind is a template line left unedited: the filing decides, and
    the count is reported on Read Me. Names are written in one spelling.
    """
    kind = asset.facility_type if asset.facility_type in {"School", "Health centre"} else ""
    banner_lg, banner_facility = strip_lg_prefix(asset.facility) if asset.facility else ("", "")
    if banner_facility and not banner_lg:
        # "KYEIHARA HC III MITOOMA": the district written after the facility.
        banner_facility, banner_lg = split_trailing_lg(banner_facility)
    column_kind = asset.extras.get("column_facility") or ""
    if not (banner_facility and not is_placeholder(banner_facility)
            and (looks_like_facility(banner_facility) or (column_kind and bare_place_name(banner_facility)))):
        banner_facility = ""
    if not kind and column_kind and banner_facility:
        kind = column_kind
    if not kind:
        kind = facility_kind(banner_facility) or facility_kind(asset.facility) if asset.facility else ""
    lg = resolve_lg(asset.lg) if asset.lg else None
    if lg is None and asset.lg:
        lg = lg_from_text(asset.lg)
    if lg is None and banner_lg:
        lg = resolve_lg(banner_lg)
    if lg is None and banner_facility:
        lg = lg_from_text(banner_facility)

    folder = only_of_kind(contexts, kind) if kind else None
    if folder is None and len(contexts) == 1:
        folder = contexts[0]
        if kind and facility_kind(folder[1]) and facility_kind(folder[1]) != kind:
            folder = None
    folder_lg = (resolve_lg(folder[0]) or lg_from_text(folder[0])) if folder and folder[0] else None
    if folder and folder_lg is None and folder[1]:
        folder_lg = lg_of_known_facility(folder[1], kind)
    folder_facility = folder[1] if folder else ""

    facility_text = banner_facility
    if folder_facility:
        if not facility_text:
            facility_text = folder_facility
        elif facility_key(facility_text, kind) != facility_key(folder_facility, kind):
            probe_lg = lg or folder_lg
            banner_display = canonical_facility(facility_text, probe_lg, kind)[0]
            folder_display = canonical_facility(folder_facility, folder_lg, kind)[0]
            if facility_key(banner_display, kind) != facility_key(folder_display, kind):
                PLACE_NOTES[f"{relative}: banner '{facility_text}' replaced by filed facility '{folder_display}'"] += 1
                facility_text = folder_facility
                lg = folder_lg or lg
    if lg is None:
        lg = folder_lg
    if lg is None and contexts:
        governments = {resolve_lg(context[0]) for context in contexts}
        governments.discard(None)
        if len(governments) == 1:
            lg = governments.pop()

    if not facility_text or lg is None:
        reconciled = only_of_kind(RECONCILED.get(relative, []), kind) if kind else None
        if reconciled is None and len(RECONCILED.get(relative, [])) == 1:
            reconciled = RECONCILED[relative][0]
        if reconciled:
            if lg is None:
                lg = resolve_lg(reconciled[0])
            if not facility_text:
                facility_text = reconciled[1]
                kind = kind or reconciled[2]
    if not kind and facility_text:
        kind = facility_kind(facility_text)
    if lg is None and facility_text and not is_placeholder(facility_text):
        # The reconciliation names the one local government of this facility.
        lg = lg_of_known_facility(facility_text, kind)
    if facility_text and not is_placeholder(facility_text) and lg is not None and not folder_facility:
        # A consolidation sheet that carries the previous block's district onto a
        # facility the reconciliation places in another local government.
        bucket = known_facilities().get(lg.key, {})
        if facility_key(facility_text, kind) not in bucket:
            # The district and its municipality share a name (Sheema, Sheema MC): a
            # facility the sibling vote lists, even misspelt, belongs there before a
            # namesake in a distant district.
            owner = None
            for sibling in known_local_governments().values():
                if sibling.key != lg.key and norm(sibling.base) == norm(lg.base):
                    display, _found = canonical_facility(facility_text, sibling, kind)
                    if facility_key(display, kind) in known_facilities().get(sibling.key, {}):
                        owner = sibling
                        break
            owner = owner or lg_of_known_facility(facility_text, kind) or fuzzy_owner(facility_text, kind)
            if owner is not None and owner.key != lg.key:
                PLACE_NOTES[f"{relative}: '{facility_text}' moved from {lg.display} to {owner.display}"] += 1
                lg = owner
    elif facility_text and lg is not None and folder_lg is not None and folder_lg.key != lg.key:
        # The banner wrote the district ("Apac") for a facility filed, and reconciled,
        # under the municipality ("Apac MC"): the filing names the vote.
        key = facility_key(facility_text, kind)
        if key in known_facilities().get(folder_lg.key, {}) and key not in known_facilities().get(lg.key, {}):
            PLACE_NOTES[f"{relative}: '{facility_text}' moved from {lg.display} to {folder_lg.display} (filed there)"] += 1
            lg = folder_lg
    asset.extras["facility_raw"] = facility_text
    if facility_text and not is_placeholder(facility_text):
        display, found_kind = canonical_facility(facility_text, lg, kind)
        asset.facility = display
        asset.facility_type = found_kind or kind
    else:
        asset.facility = ""
        asset.facility_type = kind
    asset.lg = lg.display if lg else ""
    asset.extras["lg_key"] = lg.key if lg else ""
    asset.extras["facility_key"] = facility_key(asset.facility, asset.facility_type) if asset.facility else ""


def mark_serial_asset_numbers(assets: list[Asset]) -> None:
    """An Asset Number column that counts 1, 2, 3 down the sheet is a line number,
    never a quantity."""
    numbers = [bare_integer(asset.asset_number) for asset in assets if asset.asset_number]
    if len(numbers) < 5:
        return
    integers = [number for number in numbers if number is not None]
    if len(integers) < 0.8 * len(numbers):
        return
    steps = sum(1 for previous, current in zip(integers, integers[1:]) if current - previous == 1)
    if steps >= 0.6 * (len(integers) - 1):
        for asset in assets:
            asset.extras["serial_asset_number"] = True


FACILITY_LEVEL_FOLDERS = ("_team-documents", "_district-documents")


def team_level(relative: str) -> bool:
    """A team, district or multi-team document names no facility in its path."""
    parts = relative.split("/")
    return relative.startswith("_multi-team/") or any(part in FACILITY_LEVEL_FOLDERS for part in parts)


_KNOWN_ANYWHERE: set[str] | None = None


def known_anywhere(key: str) -> bool:
    """A facility the reconciliation lists under any local government (a source
    that writes Jinja for a Jinja City facility still names a UgIFT facility)."""
    global _KNOWN_ANYWHERE
    if _KNOWN_ANYWHERE is None:
        _KNOWN_ANYWHERE = {facility for bucket in known_facilities().values() for facility in bucket}
    return bool(key) and key in _KNOWN_ANYWHERE


def in_scope(asset: Asset) -> bool:
    """A row of a team- or district-level register belongs to the register only when
    it names a UgIFT facility: one the reconciliation lists for that local
    government, or a seed school, Health Centre III or regional blood bank by name."""
    if not asset.facility:
        return False
    key = asset.extras.get("facility_key") or facility_key(asset.facility, asset.facility_type)
    bucket = known_facilities().get(asset.extras.get("lg_key") or "", {})
    if key in bucket or known_anywhere(key):
        return True
    if asset.extras.get("column_facility") and asset.facility_type in {"School", "Health centre"}:
        # Named under a HEALTH CENTRE or SCHOOL column of a programme verification sheet.
        return True
    raw = asset.extras.get("facility_raw") or asset.facility
    if "/_district-documents/" in asset.source_file:
        # A district-wide register lists every health centre of the district; only a
        # seed school named as such, or a reconciled facility, is the programme's.
        return bool(re.search(r"(?i)\bseed\b", raw))
    return bool(re.search(r"(?i)\bseed\b|health\s*cent(?:re|er)\s*(?:iii|111|3)\b|\bhc\s*(?:iii|111|3)\b|regional blood bank", raw))


def sheet_places(name: str) -> tuple[str, str]:
    """A sheet named "KABALE- LG" or "Kigarama Seed SS" is a banner for its rows."""
    text = clean(name)
    if re.fullmatch(r"(?i)sheet\s*\d*|table\s*\d*|master|data|register|list|summary", text):
        return "", ""
    lg = resolve_lg(text)
    if lg:
        return lg.display, ""
    banner = parse_banner(text)
    if banner:
        return banner
    if looks_like_facility(text):
        return "", text
    return "", ""


STEM_NOISE = re.compile(
    r"(?i)\b(ugift|ugft|asset|assets|verification|recording|tool|kit|toolkit|final|edited|copy|of|and|form|report|"
    r"template|field|excel|updated|register|data|team|\d+)\b|[\(\)\[\]_\-.,]+"
)


def stem_places(path: Path) -> tuple[str, str]:
    """A file named "UGIFT BULAGA_HCIII.docx" or "KIGANDO SEED.docx" names its facility."""
    text = clean(STEM_NOISE.sub(" ", path.stem))
    if not text:
        return "", ""
    banner = parse_banner(text)
    if banner and (banner[0] or banner[1]):
        return banner
    lg = lg_from_text(text)
    if looks_like_facility(text) and not re.search(r"(?i)\b(and|&)\b", text):
        return (lg.display if lg else ""), text
    return (lg.display if lg else ""), ""


def read_workbook(path: Path) -> list[Asset]:
    relative = path.relative_to(GROUPED).as_posix()
    contexts = CONTEXTS.get(path) or [context for context in [path_context(path)] if context[0]]
    weak_context: tuple[str, str] | None = None
    if not contexts:
        stem_lg, stem_facility = stem_places(path)
        if stem_facility and not re.search(
            r"(?i)\b(schools|centres|centers|hcs|dtb|team|list|register|facilities|hospitals|health|centre|center|"
            r"one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve)\b",
            stem_facility,
        ):
            # "UGIFT BULAGA_HCIII.docx" names its one facility; used only when no
            # banner or column in the file names one.
            weak_context = (stem_lg, stem_facility)
    fallback_lg = contexts[0][0] if contexts and len({context[0] for context in contexts}) == 1 else ""
    assets: list[Asset] = []
    offsets: dict[str, int] = {}
    if path.suffix.lower() == ".docx":
        if not docx_has_asset_table(path):
            return []
        sheets = docx_tables(path)
        # Banner paragraphs are prepended to each table; row numbers in Source
        # location still count the table's own rows.
        for name, rows, _kind in sheets:
            offsets[name] = sum(1 for row in rows if len(row) == 1 and not isinstance(row, tuple)) if rows else 0
        sheets = [(name, rows, kind) for name, rows, kind in sheets]
    else:
        sheets = [(name, rows, "") for name, rows in sheet_rows(path)]
    template_maps = []
    for _name, rows, _kind in sheets:
        for row in rows[:40]:
            mapping = classify_header(row)
            if mapping:
                template_maps.append(mapping)
                break
    workbook_has_places = any("lg" in mapping and "facility" in mapping for mapping in template_maps)
    last_docx_places: dict[str, tuple[str, str]] = {}
    last_docx_header: dict[str, list[object]] = {}
    for name, rows, section_kind in sheets:
        if SUPPLY_SHEET.search(name):
            # A delivery note, requisition voucher or interview sheet lists what was
            # issued, not what was found.
            continue
        header_row = next((row for row in rows[:40] if classify_header(row)), None)
        mapping = classify_header(header_row) if header_row is not None else None
        if path.suffix.lower() == ".docx":
            previous_header = last_docx_header.get(section_kind)
            if header_row is None and rows and previous_header is not None and len(rows[-1]) == len(previous_header):
                # A table split across pages continues the previous table's columns.
                lead = 0
                while lead < len(rows) and len(rows[lead]) == 1:
                    lead += 1
                rows = rows[:lead] + [list(previous_header)] + rows[lead:]
                offsets[name] = offsets.get(name, 0) + 1
                mapping = classify_header(previous_header)
            elif header_row is not None:
                last_docx_header[section_kind] = header_row
        if mapping and sheet_is_duplicate_draft(name, mapping, workbook_has_places):
            continue
        sheet_lg, sheet_facility = sheet_places(name) if path.suffix.lower() != ".docx" else ("", "")
        parsed: list[Asset] = []
        if mapping:
            parsed = parse_template_sheet(name, rows, relative, path.name, sheet_lg or fallback_lg)
        if not parsed:
            parsed = parse_qty_register(name, rows, relative, sheet_lg or fallback_lg)
        if not parsed:
            parsed = parse_flexible_sheet(name, rows, relative, path.name, sheet_lg or fallback_lg)
        offset = offsets.get(name, 0)
        mark_serial_asset_numbers(parsed)
        if path.suffix.lower() == ".docx" and parsed:
            # A Word table split across pages continues the facility named before
            # the previous table of the same checklist section.
            previous = last_docx_places.get(section_kind)
            if previous and not any(asset.facility for asset in parsed):
                for asset in parsed:
                    asset.lg, asset.facility = asset.lg or previous[0], previous[1]
            places = {(asset.lg, asset.facility) for asset in parsed if asset.facility}
            if len(places) == 1:
                last_docx_places[section_kind] = places.pop()
        for asset in parsed:
            if section_kind:
                asset.facility_type = section_kind
            if not asset.facility and sheet_facility:
                asset.facility = sheet_facility
            if offset:
                asset.source_location = re.sub(
                    r"row (\d+)$", lambda found: f"row {int(found.group(1)) - offset}", asset.source_location
                )
        assets.extend(parsed)
    for asset in assets:
        if not asset.facility_type:
            asset.facility_type = facility_type_for(asset.facility, asset.department, "", path.name)
        if weak_context and not (asset.facility and looks_like_facility(asset.facility)):
            asset.facility = weak_context[1]
            asset.lg = asset.lg or weak_context[0]
    # A workbook whose banners name two or more facilities of one kind is a team's
    # consolidation filed under one facility folder: its banners decide, not the folder.
    named: dict[str, set[str]] = defaultdict(set)
    for asset in assets:
        if asset.facility and not is_placeholder(asset.facility) and looks_like_facility(asset.facility):
            kind = asset.facility_type if asset.facility_type in {"School", "Health centre"} else facility_kind(asset.facility)
            named[kind or ""].add(facility_key(strip_lg_prefix(asset.facility)[1] or asset.facility, kind or ""))
    for asset in assets:
        kind = asset.facility_type if asset.facility_type in {"School", "Health centre"} else facility_kind(asset.facility)
        many = bool(asset.facility) and len(named.get(kind or "", ())) >= 2
        resolve_places(asset, [] if many else contexts, relative)
        if not asset.facility_type and asset.facility:
            # The folder named the facility only now (a blood bank, a hospital).
            asset.facility_type = facility_type_for(asset.facility, asset.department, "", path.name)
    if team_level(relative) and not any(context[1] for context in contexts) and weak_context is None:
        # A district-wide or team-wide register: rows that name no UgIFT facility
        # (district offices, sub-counties, primary schools) are outside this register.
        kept = [asset for asset in assets if in_scope(asset)]
        if len(kept) < len(assets):
            SCOPE_NOTES[relative] = (len(assets) - len(kept), len(assets))
        assets = kept
    return assets


SCOPE_NOTES: dict[str, tuple[int, int]] = {}
SUPPLY_SHEET = re.compile(r"(?i)voucher|delivery\s*note|requisition|invoice|receipt|interview|distribution\s*list|dispatch|\bdeployed\b")


def bare_count(value: object) -> int | None:
    """A cell whose whole value is a count, not a year, serial or amount."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        count = abs(int(value))
    else:
        text = clean(value).replace(",", "")
        if not re.fullmatch(r"[-(]?\d{1,4}\)?", text):
            return None
        count = abs(int(re.sub(r"[()]", "", text)))
    if count in range(1990, 2036) or not 1 < count <= MAX_QUANTITY:
        return None
    return count


def trailing_quantity(item: str) -> tuple[str, int] | None:
    """'Examination Couch 2' and 'Desks 125' state the count after the item name.

    A model suffix such as LaserJet 1320 or Laptop 840 is not a count.
    Larger counts are accepted only for bulk items such as desks and chairs.
    """
    match = re.fullmatch(r"(.+?)\s+(\d{1,4})", clean(item))
    if not match:
        return None
    count = int(match.group(2))
    name = clean(match.group(1)).rstrip(" ,.;")
    if not name or count in range(1990, 2036) or not 1 < count <= MAX_QUANTITY:
        return None
    last = name.split()[-1]
    if MODEL_TAIL.match(last) or MODEL_FAMILY.search(name):
        return None
    if last.casefold().strip(".,") in PORT_WORDS and count in {4, 8, 12, 16, 24, 48}:
        # "Network switch 24", "Patch panel 24": the port count of one device.
        return None
    if count > 40 and not BULK_ITEM.search(name):
        return None
    return name, count


PORT_WORDS = {"switch", "switches", "panel", "panels", "hub", "hubs", "router", "routers"}


def phrase_count(text: object) -> tuple[int, str] | None:
    """'2 microscopes, white' states a count in front of the description. A size
    or specification ('15 inch', '3 stance', '24 port') is not a count."""
    raw = clean(text)
    match = LEADING_COUNT.match(raw)
    if not match:
        return None
    count = int(match.group(1))
    if count in range(1990, 2036) or not 1 < count <= MAX_QUANTITY:
        return None
    return count, clean(match.group(2))


def embedded_quantity(text: str) -> int | None:
    """A count written inside a description, such as 'qty 100' or '3 classrooms'."""
    raw = clean(text)
    if not raw or bare_count(raw):
        return None
    patterns = (
        rf"(?i)\b(?:qty|quantity|no\.?\s*of\s+(?:assets|items|pieces|units)|number\s+of\s+(?:assets|items|pieces|units))\s*[:\-]?\s*({COUNT_WORD})\b",
        rf"(?i)\b({COUNT_WORD})\s*(?:pcs|pieces|units|desks?|chairs?|bench(?:es)?|stools?|tables?|beds?|shel(?:f|ves))\b",
        # "They are 13 metallic grey steel", "They received 20 beds": a count in prose.
        rf"(?i)\b(?:they\s+are|there\s+are|received|recieved|delivered|has|have|got)\s+({COUNT_WORD})\s+[a-z]",
    )
    for pattern in patterns:
        match = re.search(pattern, raw)
        if match and (count := as_count(match.group(1))) and 1 < count <= MAX_QUANTITY:
            return count
    return None


def per_item_amount(value: object, total: int) -> object:
    """Split a line total across the rows created from that line."""
    if isinstance(value, bool) or not isinstance(value, (int, float)) or total <= 1:
        return value
    share = value / total
    if abs(share - round(share)) < 1e-6:
        return int(round(share))
    return round(share, 2)


def furniture_quantity(item: str) -> tuple[str, int] | None:
    match = re.fullmatch(
        r"(?i)(desks?|chairs?|stools?|tables?|shel(?:f|ves)|bench(?:es)?|beds?|cupboards?)\s+(\d{2,4})",
        clean(item),
    )
    if not match:
        return None
    count = int(match.group(2))
    if 1 < count <= MAX_QUANTITY:
        return match.group(1), count
    return None


def parenthetical_quantity(item: str) -> tuple[str, int] | None:
    match = re.fullmatch(r"(.+?)\s*[\(\[]\s*(\d{1,4})\s*[\)\]]", clean(item))
    if not match:
        return None
    count = int(match.group(2))
    if 1 < count <= MAX_QUANTITY:
        return clean(match.group(1)), count
    return None


def condition_groups(status: str, remarks: str) -> list[tuple[int, str | None]] | None:
    text = clean(f"{status}. {remarks}")
    if not text:
        return None
    def valid(number: int) -> bool:
        return 0 < number <= MAX_QUANTITY and number not in range(1990, 2036)

    match = re.search(r"(?i)(\d{1,4})\s+verified as good\s+then\s+(\d{1,4})\s+damaged", text)
    if match and valid(int(match.group(1))) and valid(int(match.group(2))):
        return [(int(match.group(1)), "Verified as good"), (int(match.group(2)), "Damaged")]
    match = re.search(r"(?i)(\d{1,4})\s+verified and in good use\s*\.?\s*(\d{1,4})\s+broken", text)
    if match and valid(int(match.group(1))) and valid(int(match.group(2))):
        return [(int(match.group(1)), "Verified and in good use"), (int(match.group(2)), "Broken")]
    match = re.search(r"(?i)all\s+(\d{1,4})\s+verified as good.*?(\d{1,4})\s+broken", text)
    if match and valid(int(match.group(1))) and valid(int(match.group(2))):
        return [(int(match.group(1)), "Verified as good"), (int(match.group(2)), "Broken")]
    functional = re.search(rf"(?i)\b({COUNT_WORD})\s+function(?:al|ing)\b", text)
    store = re.search(rf"(?i)\b({COUNT_WORD})\s+(?:other\s+)?in\s+(?:the\s+)?store\b", text)
    if functional or store:
        groups = []
        if functional and (count := as_count(functional.group(1))):
            groups.append((count, "Functional"))
        if store and (count := as_count(store.group(1))):
            groups.append((count, "In the store"))
        if sum(count for count, _ in groups) > 1:
            return [(count, label) for count, label in groups if 0 < count <= MAX_QUANTITY]
    # "2 in use and 2 kept in store", "16 in use & 13 not in use": each part counted.
    parts = re.findall(
        rf"(?i)\b({COUNT_WORD})\s+(?:are\s+|were\s+|is\s+|was\s+)?"
        r"((?:not\s+)?in\s+(?:active\s+)?use|functional|functioning|working|kept\s+in\s+(?:the\s+)?store|in\s+(?:the\s+)?store|boxed|"
        r"broken|damaged|faulty|spoilt|not\s+functional|non[- ]functional|not\s+working)\b",
        text,
    )
    groups = [(as_count(number), phrase) for number, phrase in parts]
    groups = [(count, phrase) for count, phrase in groups if count and valid(count)]
    if len(groups) >= 2 and sum(count for count, _ in groups) > 1:
        return [(count, phrase[:1].upper() + phrase[1:].lower()) for count, phrase in groups]
    verified = [int(item) for item in re.findall(r"(?i)(?<!not )(\d{1,4})\s+verified\b", text)]
    received = [int(item) for item in re.findall(r"(?i)(\d{1,4})\s+(?:were\s+|was\s+)?(?:received|recieved|supplied)\b", text)]
    received += [int(item) for item in re.findall(r"(?i)\b(?:received|recieved|supplied|counted)\s+(\d{1,4})\b", text)]
    received += [int(item) for item in re.findall(r"(?i)\ball the\s+(\d{1,4})\b", text)]
    verified = [item for item in verified if valid(item)]
    received = [item for item in received if valid(item)]
    if verified and received:
        return [(verified[0], None)]
    if len(verified) == 1:
        return [(verified[0], None)]
    if len(set(received)) == 1 and received:
        return [(received[0], None)]
    match = re.search(
        r"(?i)\b(\d{1,4})\s+(?:desks?|chairs?|stools?|tables?|shelves|benches|items?|units?)\b(?:\s+\w+){0,3}\s+(?:supplied|received|verified)",
        text,
    )
    if match and 1 < int(match.group(1)) <= MAX_QUANTITY:
        return [(int(match.group(1)), None)]
    return None


MEASURE_WORD = (
    r"(?:seater|seaters|stances?|stanza|ports?|inch(?:es)?|in|phases?|drawers?|months?|mths?|years?|yrs?|"
    r"lit(?:re|er)s?|ltrs?|l|ml|kgs?|g|gm|kva|kw|hp|volts?|v|watts?|w|gb|tb|mm|cm|m|x|way|tiers?|doors?|steps?|pins?|"
    r"ohms?|amps?|a|pieces?\s+set|piece\s+set|in\s+1|blocks?|classrooms?|rooms?|labs?|bed\s+capacity|capacity)"
)
LEADING_COUNT = re.compile(rf"(?i)^0?([1-9]\d{{0,3}})\s+(?!{MEASURE_WORD}\b)([A-Za-z].+)$")


def stated_count(asset: Asset) -> tuple[str, list[tuple[int, str | None]]] | None:
    """The count a source line states, in the prompt's order, or None."""
    if asset.explicit_qty and asset.explicit_qty > 1:
        return asset.item, [(asset.explicit_qty, None)]
    parenthetical = parenthetical_quantity(asset.item)
    if parenthetical:
        name, count = parenthetical
        return name, [(count, None)]
    trailing = trailing_quantity(asset.item)
    if trailing:
        name, count = trailing
        return name, [(count, None)]
    from_asset_number = blank_tag(asset.tag) and not asset.extras.get("serial_asset_number") and not asset.extras.get("qty_column")
    if from_asset_number and asset.extras.get("alone") and (count := bare_count(asset.asset_number)):
        return asset.item, [(count, None)]
    if from_asset_number and (found := phrase_count(asset.asset_number)):
        return asset.item, [(found[0], None)]
    if count := bare_count(asset.description):
        return asset.item, [(count, None)]
    if found := phrase_count(asset.description):
        return asset.item, [(found[0], None)]
    if count := embedded_quantity(asset.description):
        return asset.item, [(count, None)]
    if from_asset_number and (count := embedded_quantity(asset.asset_number)):
        return asset.item, [(count, None)]
    furniture = furniture_quantity(asset.item)
    if furniture:
        name, count = furniture
        return name, [(count, None)]
    leading = LEADING_COUNT.match(asset.item)
    if leading and 1 < int(leading.group(1)) <= MAX_QUANTITY:
        return clean(leading.group(2)), [(int(leading.group(1)), None)]
    groups = condition_groups(asset.status, asset.remarks)
    if groups and sum(count for count, _ in groups) > 1:
        return asset.item, groups
    return None


def decide_groups(asset: Asset, alone: bool, repeats: int = 1, per_unit_block: bool = False) -> tuple[str, list[tuple[int, str | None]]]:
    """Choose how many asset rows a source line represents.

    A stated quantity is taken from a quantity column, a number in brackets, a
    trailing count on the item name, a bare number in Asset Number (only line for
    the item, tag blank), a count in the description, or a count in the status or
    remarks. The same line typed once per unit is one asset per row: when the
    rows repeating this line number at least the stated count, or the block lists
    its items once per unit, the count is a group total and is not applied.
    """
    asset.extras["alone"] = alone
    if asset.extras.get("has_unit_rows"):
        # The units of this line are listed on the rows below it, one each.
        return identity_item(asset.item), [(1, None)]
    found = stated_count(asset)
    if not found:
        return asset.item, [(1, None)]
    name, groups = found
    total = sum(count for count, _ in groups)
    if repeats >= total or (per_unit_block and not asset.explicit_qty):
        return name, [(1, None)]
    return name, groups


def presence_key(asset: Asset) -> tuple[str, str, str, str]:
    """One return's rows for one item at one facility."""
    return (
        asset.source_file,
        asset.extras.get("lg_key") or norm(asset.lg),
        asset.extras.get("facility_key") or norm(asset.facility),
        norm(identity_item(asset.item)),
    )


def exact_key(asset: Asset) -> tuple[str, ...]:
    """The same line typed again: same item text, description, department and tag
    (or no tag). The same count in two departments is two stated counts."""
    return presence_key(asset)[:3] + (
        norm(asset.item), norm(asset.description), norm(asset.department), norm(asset.tag) if not blank_tag(asset.tag) else "",
    )


def explode(assets: list[Asset]) -> list[Asset]:
    presence = Counter(presence_key(asset) for asset in assets)
    exact = Counter(exact_key(asset) for asset in assets)
    # A block that lists its items once per unit (most rows repeat an item text)
    # states no group counts on the item names.
    block_rows: Counter = Counter()
    block_items: dict[tuple, set] = defaultdict(set)
    for asset in assets:
        block = presence_key(asset)[:3]
        block_rows[block] += 1
        block_items[block].add(norm(asset.item))
    per_unit = {
        block: rows >= 20 and len(block_items[block]) <= 0.5 * rows
        for block, rows in block_rows.items()
    }
    exploded: list[Asset] = []
    for asset in assets:
        alone = presence[presence_key(asset)] == 1
        item, groups = decide_groups(asset, alone, exact[exact_key(asset)], per_unit[presence_key(asset)[:3]])
        total = sum(count for count, _ in groups)
        running = 0
        for count, status in groups:
            for _ in range(count):
                running += 1
                copy = Asset(**{key: getattr(asset, key) for key in (
                    "item", "department", "asset_number", "description", "life", "tag",
                    "purchase", "service", "recoverable", "cost", "acc_dep", "nbv", "ytd",
                    "status", "remarks", "lg", "facility", "facility_type", "explicit_qty",
                    "source_file", "source_location",
                )})
                copy.item = canonical_item(item)
                if status:
                    copy.status = status
                if total > 1:
                    if blank_tag(asset.tag) and bare_count(asset.asset_number) == total:
                        copy.asset_number = ""
                    elif blank_tag(asset.tag) and (found := phrase_count(asset.asset_number)) and found[0] == total:
                        copy.asset_number = ""
                        if not copy.description or blank_tag(copy.description):
                            copy.description = found[1]
                    if bare_count(asset.description) == total:
                        copy.description = ""
                    elif (found := phrase_count(asset.description)) and found[0] == total:
                        copy.description = found[1]
                    for field_name in ("recoverable", "cost", "acc_dep", "nbv", "ytd"):
                        if field_name == "cost" and asset.extras.get("unit_cost"):
                            continue
                        setattr(copy, field_name, per_item_amount(getattr(copy, field_name), total))
                    if copy.description and (found := LEADING_COUNT.match(copy.description)) and int(found.group(1)) == total:
                        copy.description = clean(found.group(2))
                    copy.extras["unit"] = f"item {running} of {total}"
                exploded.append(copy)
    return exploded


def blank_tag(value: str) -> bool:
    """A tag cell that names no tag: blank, N/A, none, nil, not engraved and the
    like. Anything holding a digit is a real number (UG/NILE/01, KAN/A/12)."""
    text = norm(value)
    if not text or is_placeholder(value):
        return True
    # "Not engraved (2)" or "N/A 2" is still no tag.
    text = re.sub(r"\s*\(?\d{1,3}\)?$", "", text).strip() if re.match(r"(not|n a|na|none|nil|no tag)", text) else text
    if re.search(r"\d", text):
        return False
    if re.search(
        r"\b(?:all|some|they|none|no|not|yet|were|are|was)\b.*\b(?:engrav|tagg|label|number)|\bengrav\w*\b.*\b(?:not|no)\s+number|"
        r"^labell?ed$|inad\w*\s+tagg|unable\s+to\s+see|no+t\s*engrav|^not\s+yet$|^no\s+engrave|^engraved\s+(?:but\s+)?not\b|^(?:blank|nothing)$",
        text,
    ):
        # "All not engraved yet", "some are engraved but not numbered", "labelled".
        return True
    return re.fullmatch(
        r"(not\s*(yet\s*)?engrav\w*|not\s*tagged|no\s*tags?|not\s*labell?ed|not\s*(available|applicable|indicated)|"
        r"engraved|nil+|none|n\s*a|na|yes|no|not|tag\s*number.*|to\s+be\s+engraved|pending)(\s.*)?",
        text,
    ) is not None


TOOLKIT_ITEMS = ROOT / "reference" / "toolkit_items.json"


def _toolkit_index() -> dict[str, str]:
    index: dict[str, str] = {}
    if TOOLKIT_ITEMS.exists():
        import json

        for names in json.load(TOOLKIT_ITEMS.open(encoding="utf-8")).values():
            for name in names:
                index.setdefault(squash(name), name)
                index.setdefault(squash(re.sub(r"[^A-Za-z0-9 ]", "", name)), name)
    return index


TOOLKIT_INDEX = _toolkit_index()


def canonical_item(item: str) -> str:
    """The toolkit's spelling of an equipment name where the source spells the same
    name without spaces or with letters split across runs ("B.P.Machine,Digital",
    "Autoclav e 20 Liters.,Du o Operated")."""
    text = clean(item)
    if not text:
        return text
    key = squash(text)
    found = TOOLKIT_INDEX.get(key) or TOOLKIT_INDEX.get(squash(re.sub(r"[^A-Za-z0-9 ]", "", text)))
    return found or text


def identity_item(item: str) -> str:
    parenthetical = parenthetical_quantity(item)
    if parenthetical:
        return canonical_item(parenthetical[0])
    trailing = trailing_quantity(item)
    if trailing:
        return canonical_item(trailing[0])
    return canonical_item(item)


def line_signature(asset: Asset) -> tuple[str, ...]:
    """Identity of one recorded line, ignoring which workbook it came from."""
    tag = "" if blank_tag(asset.tag) else norm(asset.tag)
    item = norm(identity_item(asset.item))
    if tag:
        return ("tag", tag, item)
    description = "" if bare_count(asset.description) else norm(asset.description)
    return ("line", item, description, norm(asset.status), norm(asset.department))


def represented_count(asset: Asset) -> int:
    _, groups = decide_groups(asset, alone=True)
    return sum(count for count, _ in groups)


def union_facility_submissions(assets: list[Asset]) -> tuple[list[Asset], list[str]]:
    """Keep every distinct item when a facility was submitted more than once.

    The same line in two returns is kept once, at the larger recorded count.
    An item that appears in only one return is kept.
    """
    grouped: dict[tuple[str, str], dict[str, list[Asset]]] = defaultdict(lambda: defaultdict(list))
    for asset in assets:
        key = asset.extras.get("facility_key") or facility_key(asset.facility, asset.facility_type)
        if not key:
            continue
        grouped[(asset.extras.get("lg_key") or norm(asset.lg), key)][asset.source_file].append(asset)
    drop: set[int] = set()
    notes: list[str] = []
    for sources in grouped.values():
        if len(sources) < 2:
            continue
        buckets: dict[tuple[str, ...], list[Asset]] = defaultdict(list)
        for rows in sources.values():
            for asset in rows:
                buckets[line_signature(asset)].append(asset)
        folded = 0
        for rows in buckets.values():
            by_file: dict[str, list[Asset]] = defaultdict(list)
            for asset in rows:
                by_file[asset.source_file].append(asset)
            if len(by_file) < 2:
                continue
            winner = max(
                by_file,
                key=lambda name: (
                    sum(represented_count(asset) for asset in by_file[name]),
                    len(by_file[name]),
                ),
            )
            for name, file_rows in by_file.items():
                if name == winner:
                    continue
                folded += len(file_rows)
                for asset in file_rows:
                    drop.add(id(asset))
        if folded:
            sample = next(iter(next(iter(sources.values()))))
            notes.append(
                f"{sample.facility} ({sample.lg}): {len(sources)} returns combined; "
                f"{folded} repeated lines folded in; items found in only one return were kept."
            )
    kept = [asset for asset in assets if id(asset) not in drop]
    return kept, notes


def check_examples() -> None:
    def groups(item="", status="", remarks="", qty=None, alone=True, asset_number="", description="", tag="", repeats=None):
        asset = Asset(
            item=item, status=status, remarks=remarks, explicit_qty=qty,
            asset_number=asset_number, description=description, tag=tag,
        )
        if repeats is None:
            repeats = 1 if alone else 2
        return decide_groups(asset, alone, repeats)

    assert groups(status="2 functional and one in the store")[1] == [(2, "Functional"), (1, "In the store")]
    assert groups(remarks="116 verified as good then 4 damaged")[1] == [(116, "Verified as good"), (4, "Damaged")]
    assert groups(remarks="Received 2 and all still functioning.")[1] == [(2, None)]
    assert groups(remarks="39 desks were supplied")[1] == [(39, None)]
    assert groups(item="B.P. Machine, Digital(2)", alone=True) == ("B.P. Machine, Digital", [(2, None)])
    # The same bracketed line in two departments of one return still states two each.
    assert groups(item="B.P. Machine, Digital(2)", alone=False, repeats=1) == ("B.P. Machine, Digital", [(2, None)])
    assert groups(item="Office Chairs (20)", alone=False, repeats=21) == ("Office Chairs", [(1, None)])
    # The same bracketed line typed once per unit (193 rows of 192) is one asset per row.
    assert groups(item="Laboratory stools (192)", alone=False, repeats=193) == ("Laboratory stools", [(1, None)])
    assert groups(item="Desks", qty=60, alone=False) == ("Desks", [(60, None)])
    assert groups(item="3 SEATER SCHOOL DESK")[1] == [(1, None)]
    assert groups(item="24 PORT SWITCH")[1] == [(1, None)]
    assert groups(item="Mattresses", description="60 mattresses, foam", tag="N/A")[1] == [(60, None)]
    assert groups(item="Laptop Dell 15")[1] == [(1, None)]
    assert groups(item="Bench 50") == ("Bench", [(50, None)])
    assert groups(remarks="600 verified as good then 4 damaged")[1] == [(1, None)]
    assert groups(item="Bed, Adult", asset_number="14 beds and mattress", tag="BUKMB/HC/01")[1] == [(1, None)]
    assert bare_integer("(3)") == 3 and bare_integer(-9) == 9 and bare_integer("12 stances") is None
    assert blank_tag("Not yet engraved") and blank_tag("No tag") and blank_tag("Engraved") and blank_tag("Tag Number ( engrave no.)")
    assert not blank_tag("UG/NILE/01") and not blank_tag("579-BUN-ACTF-01") and not blank_tag("KAN/A/12")
    assert looks_like_facility("ADAGMON HC 111") and not looks_like_facility("2 at school") and not looks_like_facility("Educational")
    assert parse_banner("Instrument set, ENT Basic for HCIII") is None
    assert parse_banner("Facility Name: Abalang HC III") == ("", "Abalang HC III")
    assert parse_banner("ADEKNINO HC II") == ("", "ADEKNINO HC II")
    assert groups(item="Desks 125") == ("Desks", [(125, None)])
    assert groups(item="Desks 107", alone=False, repeats=108)[1] == [(1, None)]
    assert groups(remarks="40 verified and in good use.45 broken")[1][0][0] == 40
    assert groups(remarks="Counted 21 and verified them")[1] == [(21, None)]
    assert groups(status="Not received")[1] == [(1, None)]
    assert groups(remarks="120 receive , 118 verified and in use")[1] == [(118, None)]
    assert groups(item="CPU", status="working well", alone=False)[1] == [(1, None)]
    assert groups(item="Office chairs", asset_number="286", tag="N/A") == ("Office chairs", [(286, None)])
    assert groups(item="Examination Couch 2") == ("Examination Couch", [(2, None)])
    assert groups(item="Autoclave 20 Liters., Duo Operated 2")[0].endswith("Operated")
    assert groups(item="Autoclave 20 Liters., Duo Operated 2")[1] == [(2, None)]
    assert groups(item="Desks", description="120", tag="N/A")[1] == [(120, None)]
    assert groups(item="CPU", asset_number="15", tag="SSUGU-SEED PC/02")[1] == [(1, None)]
    assert groups(item="Desks", asset_number="120", alone=False)[1] == [(1, None)]
    assert groups(item="HP LASERJET 1300")[1] == [(1, None)]
    assert groups(item="HP LAPTOP 840")[1] == [(1, None)]
    assert groups(item="LASERJET 1320")[1] == [(1, None)]
    assert groups(item="School desks 120") == ("School desks", [(120, None)])
    assert groups(item="Counting chamber", asset_number="2 microscopes, white", tag="N/A")[1] == [(2, None)]
    assert groups(item="Monitor", description="15 inch screen")[1] == [(1, None)]
    banner = parse_banner("NAME OF LG: DOKOLO DISTRICT LOCAL GOVERNMENT NAME OF HEALTH FACILITY: ABALANG HEALTH CENTER III")
    assert banner and "ABALANG" in banner[1].upper()
    assert parse_banner("BITSYA HCIII, BUHWEJU DC")[0].upper().startswith("BUHWEJU")
    assert parse_banner("School Desks") is None


def display(value: object) -> object:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def write_workbook(assets: list[Asset], sources: list[str], notes: list[str]) -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Asset Register"
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    sheet.append(HEADERS)
    for cell_item in sheet[1]:
        cell_item.font = header_font
        cell_item.fill = header_fill
        cell_item.alignment = Alignment(wrap_text=True, vertical="center")
    for asset in assets:
        sheet.append([
            asset.item, asset.department, asset.asset_number, asset.description, display(asset.life),
            asset.tag, display(asset.purchase), display(asset.service), asset.recoverable, asset.cost,
            asset.acc_dep, asset.nbv, asset.ytd, asset.status, asset.remarks,
            asset.lg, asset.facility, asset.facility_type, asset.extras.get("unit", ""),
            asset.source_file, asset.source_location,
        ])
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 2)}"
    widths = {
        1: 36, 2: 18, 3: 18, 4: 42, 5: 14, 6: 28, 7: 18, 8: 20, 9: 16, 10: 14,
        11: 14, 12: 16, 13: 14, 14: 28, 15: 42, 16: 28, 17: 36, 18: 16, 19: 12,
        20: 55, 21: 28,
    }
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30
    notes_sheet = workbook.create_sheet("Read Me")
    notes_sheet["A1"] = "UgIFT shared asset register"
    notes_sheet["A1"].font = Font(bold=True, size=14, color="1F4E79")
    lines = [
        "One workbook. Health centres and seed schools are rows in Asset Register, not separate files.",
        "Columns follow the health-centre and seed-school templates in new-templates-to-follow. Facility holds the facility name in one spelling; Facility type is School or Health centre.",
        "Each physical item is one row, the same rule as the 22 September register. Where a source line stated a quantity, that line was repeated once per item. Unit shows item 1 of 100, and the units column in the GOU template is 1.",
        "A quantity was taken from a quantity column (2 to 500), a number in brackets, a trailing count on the item name such as Examination Couch 2 (above 40 only for desks, chairs, stools, tables, shelves, benches, beds, cupboards, couches and cylinders), "
        "a bare number in Asset Number when that was the only line for the item and the tag was blank, a bare number or leading count in the description or Asset Number, or a count written in the status or remarks. "
        "Model numbers (LaserJet 1320, Laptop 840), measures (15 inch, 20 litres, 3 seater, 2 stance, 24 port), calendar years, numbers above 500, an Asset Number beside a real engraved tag, and an Asset Number column that counts down the sheet were not read as quantities.",
        "The same line typed once per unit (Laboratory stools (192) on 193 rows, School Desks with tags 001 to 122) is one asset per row: a bracket or trailing number on such rows is the group total and was not applied. A block that lists its items once per unit states no group counts on its item names.",
        "Where a grouped line carried one cost, recoverable cost, accumulated depreciation, net book value or year-to-date depreciation, that figure was treated as the line total and divided by the quantity, so each asset row holds its share. A unit price was not divided.",
        "Equipment names spelt without spaces or with letters split across Word runs were written in the toolkit's spelling (B.P. Machine, Digital; Autoclave 20 Liters.,Duo Operated).",
        "Spreadsheets and Word verification tables under raw-data-grouped were read, except programme-documents. The data-management chat in that folder is the exception. A file is an asset source when it has an asset table: an Equipment/Item header, or a description column with condition, quantity, tag or cost.",
        "Photographs, narrative reports, reconciliation lists and distribution lists are not asset lines. In one folder the spreadsheet was kept over a Word file of the same stem, and byte-for-byte duplicates were read once.",
        "Draft sheets (Table 1, Sheet 1) were left out where the same workbook already had a consolidated sheet with local government and facility columns.",
        "Facility and local government come from the row, then from a banner on the sheet (Name of LG, Name of health facility, Name of school, or a facility name standing alone, also in Word paragraphs and details tables), then from the folder path team-NN/<Local government>/<Facility>/. "
        "A combined health-centre-and-school return filed under both facility folders gives its health-centre tables to the health centre and its school tables to the school. Where none of these names the place, the reconciliation's link from the file to one facility of that kind, or the facility named in the file name, was used.",
        "Local governments are written in the spelling of facility-reconciliation.csv (one typing slip tolerated: Muyuge is Mayuge); facility names in the reconciliation spelling where the name is listed, otherwise ending in Seed Secondary School or Health Centre III. "
        "A district-wide or team-wide register contributes only the rows that name a UgIFT health centre or seed school; district offices, sub-counties, primary schools and water schemes are outside this register.",
        "Where a team submitted more than one workbook for the same facility, the lists were combined. Two returns are the same facility when the normalised local government and facility name match. A repeated line (same tag and item, or same item, description, status and department) was kept once, from the return with the larger stated count. An item present in only one return was kept. "
        "A re-saved copy of a return whose lines are all in another copy was read once.",
        "",
        "Sources:",
        *sources,
        "",
        "Facilities combined from more than one return, files left out, and other notes:",
        *(notes or ["None."]),
    ]
    for index, line in enumerate(lines, 3):
        notes_sheet.cell(index, 1, line)
    notes_sheet.column_dimensions["A"].width = 140
    workbook.save(OUTPUT)


def row_signature(asset: Asset) -> tuple[str, ...]:
    # The facility is part of the line: the same lines filed for another facility are
    # another return, not a copy.
    return (
        norm(identity_item(asset.item)), norm(asset.description), norm(asset.tag), norm(asset.status), norm(asset.department),
        asset.extras.get("facility_key") or norm(asset.facility),
    )


def drop_near_duplicates(parsed: dict[str, list[Asset]]) -> list[str]:
    """Drop a workbook whose lines are all (98 percent or more) already in another
    workbook: a re-saved copy of the same return that differs only in its label
    cells, so byte comparison did not catch it. The copy filed under a facility
    folder is the one kept. Returns notes."""
    notes: list[str] = []
    signatures = {relative: Counter(row_signature(asset) for asset in assets) for relative, assets in parsed.items()}
    # Candidates share at least one distinctive line, so compare only within groups.
    by_line: dict[tuple[str, ...], set[str]] = defaultdict(set)
    for relative, counter in signatures.items():
        for signature in counter:
            if signature[1] or signature[2] or signature[3]:
                by_line[signature].add(relative)
    neighbours: dict[str, set[str]] = defaultdict(set)
    for members in by_line.values():
        if 1 < len(members) <= 12:
            for member in members:
                neighbours[member] |= members - {member}

    def rank(relative: str) -> tuple:
        under_facility = "/_team-documents/" not in relative and "/_district-documents/" not in relative and not relative.startswith("_multi-team/")
        return (under_facility, len(parsed[relative]), -len(relative))

    def folder_key(relative: str) -> tuple[str, str] | None:
        parts = relative.split("/")
        if len(parts) >= 4 and parts[0].startswith("team-") and not parts[1].startswith("_") and not parts[2].startswith("_"):
            place = parts[2].replace("-", " ")
            return norm(parts[1]), facility_key(place, facility_kind(place))
        return None

    dropped: set[str] = set()
    for smaller in sorted(parsed, key=rank):
        if smaller in dropped or len(parsed[smaller]) < 20:
            continue
        for larger in sorted(neighbours.get(smaller, ()), key=rank, reverse=True):
            if larger in dropped or rank(larger) <= rank(smaller):
                continue
            if folder_key(smaller) and folder_key(larger) and folder_key(smaller) != folder_key(larger):
                # The same lines filed under two facilities are two returns; both are kept.
                continue
            shared = sum(min(count, signatures[larger].get(signature, 0)) for signature, count in signatures[smaller].items())
            if shared >= 0.98 * len(parsed[smaller]):
                dropped.add(smaller)
                notes.append(f"{smaller}: {len(parsed[smaller]):,} lines, {shared:,} already in {larger}; read once.")
                break
    for relative in dropped:
        del parsed[relative]
    return notes


def main() -> None:
    check_examples()
    RECONCILED.update(reconciliation_sources())
    files = candidate_files()
    parsed: dict[str, list[Asset]] = {}
    skipped: list[str] = []
    for path in files:
        relative = path.relative_to(GROUPED).as_posix()
        try:
            assets = read_workbook(path)
        except Exception as error:
            skipped.append(f"{relative} ({error})")
            continue
        if not assets:
            skipped.append(relative)
            continue
        parsed[relative] = assets
        print(f"{len(assets):6,}  {relative}", flush=True)
    duplicate_notes = drop_near_duplicates(parsed)
    for note in duplicate_notes:
        print("near duplicate:", note, flush=True)
    collected: list[Asset] = [asset for assets in parsed.values() for asset in assets]
    used: list[str] = [f"{relative} ({len(assets):,} source rows)" for relative, assets in parsed.items()]
    collected, overlap_notes = union_facility_submissions(collected)
    overlap_notes = list(overlap_notes) + [f"Near-duplicate workbook left out: {note}" for note in duplicate_notes]
    overlap_notes += [f"Left out: {relative}. {reason}" for relative, reason in EXCLUDED_SOURCES.items()]
    overlap_notes += [f"Left out folder: {folder}/. {reason}" for folder, reason in EXCLUDED_FOLDERS.items()]
    if PLACE_NOTES:
        overlap_notes.append(
            f"Banner facility replaced by the filed facility of the same kind, or a facility moved to the local government "
            f"the reconciliation lists it under, on {sum(PLACE_NOTES.values()):,} rows in {len(PLACE_NOTES)} tables: "
            + "; ".join(sorted(PLACE_NOTES)[:12])
        )
    for relative, (dropped, total) in sorted(SCOPE_NOTES.items(), key=lambda item: -item[1][0]):
        overlap_notes.append(
            f"Out of scope: {relative}: {dropped:,} of {total:,} lines name no UgIFT health centre or seed school "
            "(district offices, sub-counties, primary schools, water schemes) and were left out."
        )
    exploded = explode(collected)
    multi = sum(1 for asset in exploded if asset.extras.get("unit"))
    print(f"source rows kept {len(collected):,}; output rows {len(exploded):,}; rows from a quantity split {multi:,}")
    largest = Counter(
        (asset.source_location, asset.item, asset.extras.get("unit", "").split(" of ")[-1])
        for asset in exploded if asset.extras.get("unit")
    )
    print("largest splits:")
    seen = set()
    for (location, item, total), _count in largest.most_common(12):
        key = (location, item, total)
        if key in seen:
            continue
        seen.add(key)
        print(f"  {total:>4}  {item}  {location}")
    write_workbook(exploded, used, overlap_notes)
    print(f"wrote {OUTPUT}")
    print(f"skipped {len(skipped)} shared workbooks with no template register")


if __name__ == "__main__":
    main()
