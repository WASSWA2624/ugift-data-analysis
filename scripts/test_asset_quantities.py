"""Regression checks for one source group -> one row per physical asset."""

import unittest
import math
import tempfile
import csv
from unittest.mock import patch
from pathlib import Path

import merge_shared_asset_registers as registers
from openpyxl import load_workbook

from merge_shared_asset_registers import (
    Asset, QUANTITY_AUDIT, decide_groups, explode, take_asset,
    union_facility_submissions, check_examples,
    drop_near_duplicates,
)


def asset(**values):
    return Asset(source_file="return.xlsx", source_location="Sheet row 2",
                 lg="Test district", facility="Test HC III", **values)


class PhysicalAssetQuantityTests(unittest.TestCase):
    def test_bp_machine_quantity_in_each_source_location(self):
        examples = [
            asset(item="BP machines", explicit_qty=120),
            asset(item="BP machines (120)"),
            asset(item="BP machines 120"),
            asset(item="120 BP machines"),
            asset(item="BP machines", description="120"),
            asset(item="BP machines", description="BP machines (120)"),
            asset(item="BP machines", description="120 BP machines, digital"),
            asset(item="BP machines", remarks="120 units"),
            asset(item="BP machines", department="Quantity: 120"),
            asset(item="BP machines", extras={"source_cells": ["one hundred and twenty units"]}),
        ]
        for source in examples:
            with self.subTest(source=source):
                rows = explode([source])
                self.assertEqual(len(rows), 120)
                self.assertEqual(rows[0].extras["unit"], "item 1 of 120")
                self.assertEqual(rows[-1].extras["unit"], "item 120 of 120")

    def test_unmapped_column_and_unlimited_positive_counts(self):
        source = take_asset(["BP machine", "Quantity: 2026", "Good"],
                            {"item": 0, "status": 2}, "source.xlsx", "Sheet row 2")
        self.assertEqual(len(explode([source])), 2026)
        for value in (123456, "123,456", "123456.00", "one hundred and twenty thousand"):
            source = take_asset(["BP machine", value], {"item": 0, "explicit_qty": 1}, "s.xlsx", "r2")
            expected = 120000 if isinstance(value, str) and value.startswith("one") else 123456
            self.assertEqual(decide_groups(source, True)[1], [(expected, None)])
        source = take_asset(["Assorted Tyres", "60Pcs"], {"item": 0, "explicit_qty": 1}, "s.xlsx", "r2")
        self.assertEqual(len(explode([source])), 60)
        self.assertEqual(len(explode([asset(item="BP machine", remarks="120pcs")])), 120)

    def test_repeated_grouped_lines_retain_all_units(self):
        sources = [asset(item="BP machines (120)") for _ in range(121)]
        self.assertEqual(len(explode(sources)), 14520)
        self.assertEqual(len(explode([asset(item="Solid flush doors", explicit_qty=count) for count in (14, 7, 6)])), 27)
        self.assertEqual(len(explode([asset(item="Chair", status="Good") for _ in range(4)])), 4)

    def test_repeated_submissions_choose_larger_return_with_multiplicity(self):
        first = [asset(item="BP machines", explicit_qty=2) for _ in range(3)]
        second = [asset(item="BP machines", explicit_qty=4)]
        second[0].source_file = "resubmission.xlsx"
        kept, _ = union_facility_submissions(first + second)
        self.assertEqual(len(kept), 3)
        self.assertEqual(len(explode(kept)), 6)

    def test_workbooks_with_different_quantities_are_not_resaved_copies(self):
        parsed = {}
        for filename, count in (("original-long-name.xlsx", 120), ("copy.xlsx", 2)):
            path = f"team-01/Test/Test-HC-III/{filename}"
            parsed[path] = [asset(item=f"BP machine type{index}", status="Good", explicit_qty=count) for index in range(20)]
        self.assertEqual(drop_near_duplicates(parsed), [])
        self.assertEqual(len(parsed), 2)

    def test_proven_unit_records_are_not_expanded_twice(self):
        sources = [asset(item="BP machines (3)", tag=f"BP/{index}") for index in range(3)]
        self.assertEqual(len(explode(sources)), 3)
        layout = [asset(item="BP machines (120)", extras={"has_unit_rows": True}),
                  asset(item="BP machines", description="120", extras={"unit_row": True})]
        self.assertEqual(len(explode(layout)), 2)
        # Two separately tagged grouped lines are still groups, not proof of 120 units.
        grouped = [asset(item="BP machines (3)", tag=f"BP/{index}") for index in range(2)]
        self.assertEqual(len(explode(grouped)), 6)

    def test_condition_subsets_sum_without_repeating_total(self):
        source = asset(item="BP machine 120", explicit_qty=120,
                       status="116 verified as good then 4 damaged", remarks="120 units")
        rows = explode([source])
        self.assertEqual(len(rows), 120)
        self.assertEqual(sum(row.status == "Damaged" for row in rows), 4)
        self.assertEqual(len(explode([asset(item="BP machine", status="one hundred and sixteen good and four damaged")])), 120)
        self.assertEqual(len(explode([asset(item="BP machine", status="40 verified and in good use.45 broken")])), 85)
        self.assertEqual(len(explode([asset(item="BP machine", status="18 stolen, 10 in use")])), 28)
        self.assertEqual(len(explode([asset(item="BP machine", status="2 functional", remarks="All 2 are in use")])), 2)
        self.assertEqual(len(explode([asset(item="BP machine", status="116 functional and 4 damaged", remarks="120 units")])), 120)
        self.assertEqual(len(explode([asset(item="BP machine", status="116 verified as good", remarks="4 damaged")])), 120)
        for description in (
            "all 120 verified as good but 4 broken",
            "all 120 verified as good of which 4 broken",
            "120 verified as good including 4 broken",
        ):
            rows = explode([asset(item="BP machine", description=description)])
            self.assertEqual(len(rows), 120, description)
            self.assertEqual(sum(row.status == "Broken" for row in rows), 4, description)

    def test_line_totals_unit_prices_and_provenance(self):
        source = asset(item="BP machine", explicit_qty=120, description="120 BP machines, digital",
                       cost=120000, recoverable=60000, acc_dep=12000, nbv=108000, ytd=2400,
                       extras={"filled_from": {"donor.xlsx"}})
        rows = explode([source])
        self.assertEqual(sum(row.cost for row in rows), source.cost)
        self.assertEqual((rows[0].cost, rows[0].recoverable, rows[0].acc_dep, rows[0].nbv, rows[0].ytd),
                         (1000, 500, 100, 900, 20))
        self.assertEqual(rows[0].description, "BP machines, digital")
        self.assertEqual(rows[-1].source_location, "Sheet row 2")
        self.assertEqual(rows[-1].extras["filled_from"], {"donor.xlsx"})
        source.extras["unit_cost"] = True
        self.assertEqual(explode([source])[0].cost, 120000)
        for total, count in ((100, 3), (1000000, 601)):
            rows = explode([asset(item="BP machine", explicit_qty=count, cost=total)])
            self.assertAlmostEqual(math.fsum(row.cost for row in rows), total)

    def test_streamed_workbook_keeps_layout_and_all_unit_rows(self):
        old_output = registers.OUTPUT
        try:
            with tempfile.TemporaryDirectory() as folder:
                registers.OUTPUT = Path(folder) / "register.xlsx"
                registers.write_workbook(explode([asset(item="BP machine", explicit_qty=120)]), ["return.xlsx"], ["Quantity checked."])
                workbook = load_workbook(registers.OUTPUT)
                sheet = workbook["Asset Register"]
                self.assertEqual(sheet.max_row, 121)
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:U121")
                self.assertTrue(sheet["A1"].font.bold)
                self.assertEqual(sheet["S121"].value, "item 120 of 120")
                self.assertTrue(workbook["Read Me"]["A3"].alignment.wrap_text)
                workbook.close()
        finally:
            registers.OUTPUT = old_output

    def test_models_measures_dates_money_and_identifiers_are_not_counts(self):
        for item in ("HP LASERJET 1320", "Laptop 840", "24 PORT SWITCH", "3 SEATER SCHOOL DESK", "HP LaserJet (1320)", "Network switch (24)"):
            self.assertEqual(len(explode([asset(item=item)])), 1, item)
        for description in ("Serial number 123456", "Engine No. 123456", "Model 840", "Acquired June 2018", "Cost UGX 120000", "Power 120", "Operating voltage 220", "RAM 16", "Core i5 8400"):
            self.assertEqual(len(explode([asset(item="BP machine", description=description)])), 1, description)
            self.assertEqual(len(explode([asset(item="BP machine", description=description, explicit_qty=120)])), 120, description)
        for description in (
            "Dell Latitude 5440 Laptop Computer; Serial number 6YY2TV3",
            "Lenovo ThinkPad X1 Carbon Gen 11 Laptop i5-1335U,16GB; Serial number ABC",
            "Model THERMOFISHER SCIENTIFC TSX60086V; Serial number 1146352501240400",
        ):
            self.assertEqual(len(explode([asset(item="Computer", description=description, explicit_qty=1)])), 1)
        self.assertEqual(len(explode([asset(item="MS Office Pro 2021", explicit_qty=5)])), 5)
        self.assertEqual(len(explode([asset(item="Bags & Covers", explicit_qty=745)])), 745)
        source = take_asset(["BP machine", 2026, 120000, "2026-09-27", "BP/120"],
                            {"item": 0, "life": 1, "cost": 2, "purchase": 3, "tag": 4}, "s.xlsx", "r2")
        self.assertEqual(len(explode([source])), 1)
        self.assertEqual(len(explode([asset(item="BP machine", asset_number="120", tag="BP/1")])), 1)

    def test_negative_and_fractional_counts_are_not_partial_integers(self):
        for value in ("-120 units", "-2 functional and -3 damaged", "1.5 functional and 2.5 damaged", "Quantity: -120", "Quantity: 1.5", "⁹", "m²"):
            self.assertEqual(len(explode([asset(item="BP machine", description=value, status=value)])), 1, value)

    def test_sheet_capacity_guard_and_conflict_audit(self):
        with tempfile.TemporaryDirectory() as folder:
            audit_path = Path(folder) / "quantity-audit.csv"
            with self.assertRaisesRegex(ValueError, "exceeding Excel"):
                explode([asset(item="BP machine", explicit_qty=1048576)], audit_path=audit_path)
            with audit_path.open(encoding="utf-8-sig", newline="") as handle:
                self.assertEqual(next(csv.DictReader(handle))["output_rows"], "1048576")
        source = asset(item="BP machine", explicit_qty=120, description="130 units")
        self.assertEqual(len(explode([source])), 130)
        self.assertTrue(QUANTITY_AUDIT[0]["conflicting_counts"])

    def test_parsed_checkpoint_reuses_only_matching_source_metadata(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "return.xlsx"
            source.write_bytes(b"fixture")
            cache = root / "parsed.pkl"
            with patch.object(registers, "GROUPED", root), patch.object(registers, "read_central_sources", return_value={}), patch.object(registers, "read_workbook", return_value=[asset(item="BP machine", explicit_qty=120)]) as reader:
                first, audit = registers.parse_sources([source], cache)
                second, cached_audit = registers.parse_sources([source], cache)
                self.assertEqual(reader.call_count, 1)
                self.assertEqual(second["return.xlsx"][0].explicit_qty, 120)
                self.assertEqual(cached_audit, audit)
                source.write_bytes(b"changed fixture")
                registers.parse_sources([source], cache)
                self.assertEqual(reader.call_count, 2)

    def test_source_parse_errors_are_recorded_separately_from_empty_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with patch.object(registers, "GROUPED", root), patch.object(registers, "read_central_sources", return_value={}), patch.object(registers, "read_workbook", side_effect=[ValueError("damaged file"), []]):
                parsed, audit = registers.parse_sources([root / "broken.xlsx", root / "empty.xlsx"])
                self.assertEqual(parsed, {})
                self.assertEqual([row["status"] for row in audit], ["error", "no_asset_rows"])
                self.assertIn("damaged file", audit[0]["error"])

    def test_existing_examples(self):
        check_examples()


if __name__ == "__main__":
    unittest.main()
