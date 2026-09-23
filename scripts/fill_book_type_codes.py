"""Repair BOOK_TYPE_CODE and FACILITY_NAME in the UgIFT asset register.

The resolver follows a strict evidence order:

1. reconciled facility ID;
2. reconciled facility folder/source path;
3. the source-line LG and facility fields;
4. the exact worksheet row named by ``Source locator`` (including section
   headers and carried-down district/facility columns);
5. a facility/LG explicitly present in the source description;
6. the owning district or MDA when the source is an entity-wide register.

Unsupported values remain blank. The script rebuilds from the committed
workbook when requested, inserts FACILITY_NAME beside BOOK_TYPE_CODE, and
keeps the native Excel table metadata synchronized with the worksheet.
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import subprocess
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path, PurePosixPath
from typing import Iterable, Iterator

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_PATH = ROOT / "outputs" / "asset-register-2026-09-22" / "UgIFT Asset Register.xlsx"
RECONCILIATION_PATH = ROOT / "raw-data-grouped" / "facility-reconciliation.csv"
LOCATION_MASTER_PATH = ROOT / "Location(3)2.xlsx"
GIT_WORKBOOK_PATH = "outputs/asset-register-2026-09-22/UgIFT Asset Register.xlsx"

SOURCE_LOCATOR_RE = re.compile(
    r"^Worksheet\s+(.*?),\s*rows?\s+(\d+)(?:\s*-\s*\d+)?\s*$", re.I
)
LG_PREFIX_RE = re.compile(
    r"^(.*?(?:DISTRICT LOCAL GOV'?T|MUNICIPAL COUNCIL|CITY COUNCIL|DISTRICT LG|\sDLG|\sMC|\sCITY))(?=-)",
    re.I,
)

MDA_NAMES = {
    "MAAIF": "Ministry of Agriculture, Animal Industry and Fisheries",
    "MOES": "Ministry of Education and Sports",
    "MOFPED": "Ministry of Finance, Planning and Economic Development",
    "MOGLSD": "Ministry of Gender, Labour and Social Development",
    "MOH": "Ministry of Health",
    "MOLG": "Ministry of Local Government",
    "MOWE": "Ministry of Water and Environment",
    "MOWT": "Ministry of Works and Transport",
    "NEMA": "National Environment Management Authority",
    "OAG": "Office of the Auditor General",
    "OPM": "Office of the Prime Minister",
    "PPDA": "Public Procurement and Disposal of Public Assets Authority",
}

MDA_TEXT_MARKERS = (
    ("ministry of agriculture animal industry and fisheries", "MAAIF"),
    ("ministry of education and sports", "MOES"),
    ("ministry of finance planning and economic development", "MOFPED"),
    ("ministry of gender labour and social development", "MOGLSD"),
    ("ministry of health", "MOH"),
    ("ministry of local government", "MOLG"),
    ("ministry of water and environment", "MOWE"),
    ("ministry of works and transport", "MOWT"),
    ("office of the auditor general", "OAG"),
    ("office of the prime minister", "OPM"),
    ("national environment management authority", "NEMA"),
    ("public procurement and disposal of public assets authority", "PPDA"),
)

LG_ALIASES = {
    "amolata": "amolatar",
    "buyend": "buyende",
    "fortportal city": "fort portal city",
    "kabalore": "kabarole",
    "kakumuro": "kakumiro",
    "kyakwanzi": "kyankwanzi",
    "kyankwazi": "kyankwanzi",
    "lira disrict": "lira",
    "liras city": "lira city",
    "nebbi municipality": "nebbi mc",
    "ssembabule": "sembabule",
}

INVALID_LG_KEYS = {
    "central",
    "district",
    "eastern",
    "not allocated",
    "northern",
    "unspecified",
    "western",
}

FACILITY_MARKERS = re.compile(
    r"\b(?:HC\s*(?:II|III|IV|2|3|4|111|LLL)|HEALTH\s*CENT(?:RE|ER)|"
    r"HOSPITAL|BLOOD\s*BANK|SEED\s*(?:SECONDARY\s*)?SCHOOL|SECONDARY\s*SCHOOL)\b",
    re.I,
)

RAW_FACILITY_RE = re.compile(
    r"\b([A-Z][A-Z0-9 .&'()-]{1,55}?\s+(?:HC\s*(?:II|III|IV|2|3|4|111|LLL)|"
    r"HEALTH\s*CENT(?:RE|ER)(?:\s*(?:II|III|IV|2|3|4))?|BLOOD\s*BANK|"
    r"SEED\s*(?:SECONDARY\s*)?SCHOOL|SECONDARY\s*SCHOOL|HOSPITAL))\b",
    re.I,
)


def clean_text(value: object) -> str:
    """Return readable text without spreadsheet escape and ellipsis noise."""
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\\", "")
    text = re.sub(r"[…]+", " ", text)
    return re.sub(r"\s+", " ", text).strip(" \t\r\n-–—|,;:")


def normalized_key(value: object) -> str:
    text = clean_text(value).casefold().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def normalized_path(value: str) -> str:
    return "/".join(
        normalized_key(part)
        for part in value.replace("\\", "/").split("/")
        if part
    )


def lg_key(value: object) -> str:
    text = normalized_key(value)
    text = re.sub(r"\b(?:district local government|district local govt|local government)\b", " ", text)
    text = re.sub(r"\bdistrict\b", " ", text)
    text = re.sub(r"\bdlg\b", " ", text)
    text = re.sub(r"\bmunicipal(?:ity| council)?\b", " mc ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return LG_ALIASES.get(text, text)


def facility_key(value: object) -> str:
    text = normalized_key(value)
    text = re.sub(r"\bhealth\s*cent(?:re|er)\b", " hc ", text)
    text = re.sub(r"\b(?:h\s*c)\s*(?:111|lll|3)\b", " hc iii ", text)
    text = re.sub(r"\b(?:h\s*c)\s*(?:11|2)\b", " hc ii ", text)
    text = re.sub(r"\b(?:h\s*c)\s*(?:1v|4)\b", " hc iv ", text)
    text = re.sub(r"\bsecondary\s+school\b", " ss ", text)
    text = re.sub(r"\bincharge\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def mda_from_text(value: object) -> str | None:
    key = normalized_key(value)
    for marker, code in MDA_TEXT_MARKERS:
        if marker in key:
            return code
    return None


def format_raw_facility(value: object) -> str | None:
    """Extract a readable facility name from a tag, user title, or heading."""
    text = clean_text(value)
    for segment in re.split(r"[/|]", text):
        match = RAW_FACILITY_RE.search(clean_text(segment))
        if not match:
            continue
        facility = re.sub(r"\bINCHARGE\b.*$", "", match.group(1), flags=re.I).strip()
        facility = facility.title()
        facility = re.sub(r"\bHc\s*(?:111|Lll|3|Iii)\b", "HC III", facility, flags=re.I)
        facility = re.sub(r"\bHc\s*(?:11|2|Ii)\b", "HC II", facility, flags=re.I)
        facility = re.sub(r"\bHc\s*(?:1v|4|Iv)\b", "HC IV", facility, flags=re.I)
        return facility
    return None


@dataclass(frozen=True)
class Facility:
    name: str
    lg: str


@dataclass(frozen=True)
class Resolution:
    lg: str | None
    facility: str | None
    lg_rule: str
    facility_rule: str


class MasterData:
    """Canonical LG/facility names and their supported aliases."""

    def __init__(self) -> None:
        self.facility_by_id: dict[str, Facility] = {}
        self.folder_to_facility: dict[str, Facility] = {}
        self.lg_display: dict[str, str] = {}
        self.facility_aliases_by_lg: dict[str, dict[str, Facility]] = defaultdict(dict)
        self.global_facility_aliases: dict[str, list[Facility]] = defaultdict(list)
        self._load_location_master()
        self._load_reconciliation()
        self._lg_search = sorted(self.lg_display, key=len, reverse=True)

    def _remember_lg(self, value: object, *, prefer: bool = False) -> str:
        display = clean_text(value)
        key = lg_key(display)
        if key and (prefer or key not in self.lg_display):
            self.lg_display[key] = display
        return key

    def _load_location_master(self) -> None:
        workbook = load_workbook(LOCATION_MASTER_PATH, read_only=True, data_only=True)
        sheet = workbook.active
        for (value,) in sheet.iter_rows(min_row=2, values_only=True):
            match = LG_PREFIX_RE.search(clean_text(value))
            if match:
                self._remember_lg(match.group(1), prefer=True)
        workbook.close()

    def _load_reconciliation(self) -> None:
        with RECONCILIATION_PATH.open("r", encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                facility_id = clean_text(row.get("id"))
                name = clean_text(row.get("field_name") or row.get("name"))
                lg = clean_text(row.get("lg"))
                if not facility_id or not name or not lg:
                    continue

                lgk = self._remember_lg(lg)
                canonical_lg = self.lg_display.get(lgk, lg)
                facility = Facility(name, canonical_lg)
                self.facility_by_id.setdefault(facility_id, facility)

                aliases = {
                    facility_key(name),
                    facility_key(row.get("name")),
                    facility_key(row.get("field_name")),
                }
                aliases.update(
                    alias.replace(" seed ss", " ss")
                    for alias in tuple(aliases)
                    if " seed ss" in alias
                )
                aliases.update(
                    alias.replace(" regional blood bank", " blood bank")
                    for alias in tuple(aliases)
                    if " regional blood bank" in alias
                )
                folder = clean_text(row.get("folder")).replace("\\", "/").strip("/")
                if folder:
                    folder_key = normalized_path(folder)
                    self.folder_to_facility.setdefault(folder_key, facility)
                    aliases.add(facility_key(PurePosixPath(folder).name))

                for alias in {item for item in aliases if item}:
                    self.facility_aliases_by_lg[lgk].setdefault(alias, facility)
                    if facility not in self.global_facility_aliases[alias]:
                        self.global_facility_aliases[alias].append(facility)

    def canonical_lg(self, value: object) -> str | None:
        raw = clean_text(value)
        if not raw:
            return None
        upper = re.sub(r"[^A-Z0-9]+", "", raw.upper())
        for code in MDA_NAMES:
            if upper == code:
                return code

        key = lg_key(raw)
        if not key or key in INVALID_LG_KEYS:
            return None
        if key in self.lg_display:
            return self.lg_display[key]

        scores = sorted(
            ((SequenceMatcher(None, key, candidate).ratio(), candidate) for candidate in self.lg_display),
            reverse=True,
        )
        if scores and scores[0][0] >= 0.86 and (
            len(scores) == 1 or scores[0][0] - scores[1][0] >= 0.06
        ):
            return self.lg_display[scores[0][1]]
        return None

    def find_lg_in_text(self, value: object) -> str | None:
        text = f" {normalized_key(value)} "
        for key in self._lg_search:
            if len(key) >= 4 and f" {key} " in text:
                return self.lg_display[key]
        return None

    def facility_from_path(self, value: object) -> Facility | None:
        path = normalized_path(clean_text(value))
        prefix = "raw data grouped/"
        if path.startswith(prefix):
            path = path[len(prefix) :]
        parts = path.split("/")
        for length in range(len(parts) - 1, 1, -1):
            candidate = "/".join(parts[:length])
            if candidate in self.folder_to_facility:
                return self.folder_to_facility[candidate]
        return None

    def canonical_facility(self, value: object, lg: object = None) -> Facility | None:
        key = facility_key(value)
        if not key:
            return None
        lgk = lg_key(lg)
        aliases = self.facility_aliases_by_lg.get(lgk, {}) if lgk else {}

        if key in aliases:
            return aliases[key]
        contained = [
            facility
            for alias, facility in aliases.items()
            if len(alias) >= 5 and f" {alias} " in f" {key} "
        ]
        if contained:
            return max(contained, key=lambda item: len(facility_key(item.name)))

        if aliases and len(key) >= 5:
            scores = sorted(
                ((SequenceMatcher(None, key, alias).ratio(), alias) for alias in aliases),
                reverse=True,
            )
            if scores and scores[0][0] >= 0.90 and (
                len(scores) == 1 or scores[0][0] - scores[1][0] >= 0.05
            ):
                return aliases[scores[0][1]]

        matches = self.global_facility_aliases.get(key, [])
        return matches[0] if len(matches) == 1 else None

    def find_facility_in_text(self, value: object, lg: object = None) -> Facility | None:
        key = facility_key(value)
        if not key:
            return None
        aliases = self.facility_aliases_by_lg.get(lg_key(lg), {}) if lg else {}
        if not aliases:
            aliases = {
                alias: facilities[0]
                for alias, facilities in self.global_facility_aliases.items()
                if len(facilities) == 1
            }
        matches = [
            (len(alias), facility)
            for alias, facility in aliases.items()
            if len(alias) >= 5 and f" {alias} " in f" {key} "
        ]
        if matches:
            return max(matches, key=lambda item: item[0])[1]
        return None

    def school_from_tag(self, value: object) -> Facility | None:
        match = re.search(r"\b([A-Z][A-Z'-]{3,})\s*/\s*SSS\b", clean_text(value), re.I)
        if not match:
            return None
        prefix = normalized_key(match.group(1))
        candidates = {
            facilities[0]
            for alias, facilities in self.global_facility_aliases.items()
            if len(facilities) == 1 and alias.startswith(prefix + " ") and alias.endswith(" ss")
        }
        return next(iter(candidates)) if len(candidates) == 1 else None


def usable_facility_text(value: object, master: MasterData, lg: object = None) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    key = facility_key(text)
    if key in {"unspecified", "not allocated", "not applicable", "n a", "district"}:
        return None
    compact = re.sub(r"[^A-Za-z0-9]", "", text)
    if " " not in text and len(compact) >= 10 and any(ch.isdigit() for ch in compact):
        return None
    if lg and lg_key(text) == lg_key(lg):
        return None
    canonical = master.canonical_facility(text, lg)
    if canonical:
        return canonical.name
    if len(text) > 100 or text.count("/") >= 3:
        return None
    return text


def entity_name(lg: str) -> str:
    if lg in MDA_NAMES:
        return MDA_NAMES[lg]
    base = re.sub(
        r"\s+(?:DISTRICT LOCAL GOV'?T|DISTRICT LOCAL GOVERNMENT|DISTRICT|DLG)$",
        "",
        clean_text(lg),
        flags=re.I,
    ).strip()
    upper = base.upper()
    if upper.endswith(" MC"):
        return f"{base[:-3].strip().title()} Municipal Council"
    if upper.endswith(" CITY"):
        return f"{base.title()} Council"
    return f"{base.title()} District Local Government"


def book_code(lg: str | None) -> str | None:
    if not lg:
        return None
    if lg in MDA_NAMES:
        return f"{lg} BK"
    text = clean_text(lg).upper()
    for old, new in (
        (" DISTRICT LOCAL GOV'T", ""),
        (" DISTRICT LOCAL GOVT", ""),
        (" DISTRICT LOCAL GOVERNMENT", ""),
        (" MUNICIPAL COUNCIL", " MC"),
        (" DISTRICT LG", ""),
        (" LOCAL GOVERNMENT", ""),
        (" DISTRICT", ""),
        (" DLG", ""),
    ):
        if text.endswith(old):
            text = (text[: -len(old)] + new).strip()
            break
    return f"{text} BK" if text else None


def parse_labeled_value(text: str, label: str) -> str | None:
    match = re.search(rf"{label}\s*:\s*([^|]+)", text, re.I)
    return clean_text(match.group(1)) if match else None


def header_kind(value: object) -> str | None:
    key = normalized_key(value)
    if key in {"district", "local government", "local govvernment", "lg", "entity"}:
        return "lg"
    if key in {
        "facility",
        "health facility",
        "health centre",
        "health center",
        "hospital",
        "location",
        "school",
    }:
        return "facility"
    return None


class SourceTracer:
    """Read an exact source worksheet row and its nearest location context."""

    def __init__(self, master: MasterData) -> None:
        self.master = master

    def trace(self, source_file: object, locator: object) -> Resolution | None:
        match = SOURCE_LOCATOR_RE.match(clean_text(locator))
        if not match:
            return None
        relative = clean_text(source_file).replace("/", "\\")
        path = ROOT / relative
        if not path.exists() and path.parent.exists():
            wanted = normalized_key(path.name)
            matches = [item for item in path.parent.iterdir() if normalized_key(item.name) == wanted]
            if len(matches) == 1:
                path = matches[0]
        if not path.exists() or path.suffix.casefold() not in {".xlsx", ".xlsm", ".xls"}:
            return None
        sheet_name, row_number = match.group(1), int(match.group(2))
        result = self._contexts(str(path.resolve()), sheet_name).get(row_number)
        if not result:
            return None
        lg, facility = result
        return Resolution(lg, facility, "source worksheet", "source worksheet")

    @lru_cache(maxsize=None)
    def _contexts(self, absolute_path: str, requested_sheet: str) -> dict[int, tuple[str | None, str | None]]:
        return self._scan_rows(self._read_rows(Path(absolute_path), requested_sheet))

    def _read_rows(self, path: Path, requested_sheet: str) -> Iterator[tuple[int, tuple[object, ...]]]:
        if path.suffix.casefold() == ".xls":
            book = xlrd.open_workbook(path, on_demand=True)
            sheet = self._xlrd_sheet(book, requested_sheet)
            try:
                for row_index in range(sheet.nrows):
                    yield row_index + 1, tuple(sheet.row_values(row_index))
            finally:
                book.release_resources()
            return

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = self._openpyxl_sheet(workbook, requested_sheet)
        try:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                yield row_number, tuple(row)
        finally:
            workbook.close()

    @staticmethod
    def _openpyxl_sheet(workbook, requested: str):
        if requested in workbook.sheetnames:
            return workbook[requested]
        wanted = requested.strip().casefold()
        for name in workbook.sheetnames:
            if name.strip().casefold() == wanted:
                return workbook[name]
        raise KeyError(f"Worksheet {requested!r} not found")

    @staticmethod
    def _xlrd_sheet(workbook, requested: str):
        try:
            return workbook.sheet_by_name(requested)
        except xlrd.biffh.XLRDError:
            wanted = requested.strip().casefold()
            for name in workbook.sheet_names():
                if name.strip().casefold() == wanted:
                    return workbook.sheet_by_name(name)
            raise

    def _scan_rows(
        self, rows: Iterable[tuple[int, tuple[object, ...]]]
    ) -> dict[int, tuple[str | None, str | None]]:
        contexts: dict[int, tuple[str | None, str | None]] = {}
        lg_columns: set[int] = set()
        facility_columns: set[int] = set()
        current_lg: str | None = None
        current_facility: str | None = None

        for row_number, row in rows:
            texts = [clean_text(value) for value in row]
            nonblank = [(index, text) for index, text in enumerate(texts) if text]
            joined = " | ".join(text for _, text in nonblank)

            for index, text in nonblank:
                kind = header_kind(text)
                if kind == "lg":
                    lg_columns.add(index)
                elif kind == "facility":
                    facility_columns.add(index)

            explicit_lg = parse_labeled_value(joined, r"(?:NAME OF LG|ENTITY)")
            explicit_facility = parse_labeled_value(joined, r"NAME OF (?:HEALTH )?FACILITY")
            if explicit_lg:
                current_lg = self.master.canonical_lg(explicit_lg) or current_lg
            if explicit_facility:
                canonical = self.master.canonical_facility(explicit_facility, current_lg)
                current_facility = canonical.name if canonical else clean_text(explicit_facility)

            for index in lg_columns:
                if index < len(row):
                    candidate = self.master.canonical_lg(row[index])
                    if candidate:
                        current_lg = candidate
            for index in facility_columns:
                if index < len(row):
                    candidate = usable_facility_text(row[index], self.master, current_lg)
                    if candidate and header_kind(candidate) is None:
                        current_facility = candidate

            row_lg = self.master.find_lg_in_text(joined)
            if row_lg and len(nonblank) <= 5:
                current_lg = row_lg
            effective_lg = row_lg or current_lg

            # A facility named in a tag or description applies to that source
            # row only. Carry it forward only when the sparse row is a section
            # title; otherwise a single mention in a district-wide register
            # would incorrectly relabel every following asset.
            row_facility = self.master.find_facility_in_text(joined, effective_lg)
            if not row_facility:
                row_facility = self.master.school_from_tag(joined)
            if not row_facility:
                raw_candidates = [
                    text
                    for _, text in nonblank
                    if "/" in text or "incharge" in text.casefold()
                ]
                if len(nonblank) <= 5:
                    raw_candidates.append(joined)
                raw_facility = next(
                    (format_raw_facility(text) for text in raw_candidates if format_raw_facility(text)),
                    None,
                )
            else:
                raw_facility = None
            if row_facility and len(nonblank) <= 5:
                current_facility = row_facility.name
                current_lg = row_facility.lg
            if raw_facility and len(nonblank) <= 5:
                current_facility = raw_facility
            effective_facility = (
                row_facility.name if row_facility else raw_facility or current_facility
            )
            if row_facility:
                effective_lg = row_facility.lg

            contexts[row_number] = (effective_lg, effective_facility)
        return contexts


def facility_from_filename(source_file: object, master: MasterData, lg: object) -> Facility | None:
    stem = Path(clean_text(source_file)).stem
    if not FACILITY_MARKERS.search(stem):
        return None
    return (
        master.find_facility_in_text(stem, lg)
        or master.find_facility_in_text(stem)
        or master.canonical_facility(stem, lg)
    )


def facility_from_description(description: object, master: MasterData, lg: object) -> Facility | None:
    text = clean_text(description)
    found = master.find_facility_in_text(text, lg)
    if found:
        return found
    parts = [clean_text(part) for part in re.split(r"\s+-\s+", text) if clean_text(part)]
    for candidate in reversed(parts[-3:]):
        if FACILITY_MARKERS.search(candidate):
            return master.canonical_facility(candidate, lg)
    return None


def lg_from_district_path(source_file: object, master: MasterData) -> str | None:
    parts = [part for part in clean_text(source_file).replace("\\", "/").split("/") if part]
    for index, part in enumerate(parts):
        if normalized_key(part) == "district documents" and index >= 1:
            return master.canonical_lg(parts[index - 1])
    return None


def source_maps(
    ws_source, master: MasterData
) -> tuple[dict[int, str], dict[int, str], Counter, dict[str, Counter]]:
    """Map register rows to source-backed LG and facility values."""
    headers = [cell.value for cell in next(ws_source.iter_rows(min_row=1, max_row=1))]
    index = {name: i for i, name in enumerate(headers)}
    tracer = SourceTracer(master)
    lg_by_row: dict[int, str] = {}
    facility_by_row: dict[int, str] = {}
    rules: Counter = Counter()
    unresolved = {"lg": Counter(), "facility": Counter()}

    for row in ws_source.iter_rows(min_row=2, values_only=True):
        first = row[index["Register first row"]]
        last = row[index["Register last row"]]
        if first is None or last is None:
            continue

        facility_id = clean_text(row[index["Facility ID"]])
        source_file = row[index["Source file"]]
        source_path_key = normalized_path(clean_text(source_file))
        source_lg = row[index["Local government / MDA"]]
        source_facility = row[index["Facility / location"]]
        source_asset = row[index["Asset in source"]]

        lg: str | None = None
        facility: str | None = None
        lg_rule = ""
        facility_rule = ""

        reconciled = master.facility_by_id.get(facility_id)
        if reconciled:
            lg, facility = reconciled.lg, reconciled.name
            lg_rule = facility_rule = "reconciled facility ID"

        if not reconciled:
            path_facility = master.facility_from_path(source_file)
            if path_facility:
                lg, facility = path_facility.lg, path_facility.name
                lg_rule = facility_rule = "reconciled source folder"

        if not lg:
            lg = master.canonical_lg(source_lg)
            if lg:
                lg_rule = "source-line LG"
        if not lg:
            lg = lg_from_district_path(source_file, master)
            if lg:
                lg_rule = "district source path"

        if not facility:
            canonical = master.canonical_facility(source_facility, lg)
            if canonical:
                facility = canonical.name
                lg = canonical.lg
                facility_rule = "canonicalized source-line facility"
                lg_rule = lg_rule or "source-line facility"
            else:
                facility = usable_facility_text(source_facility, master, lg)
                if facility:
                    facility_rule = "source-line facility"

        traceable_source = any(
            marker in source_path_key
            for marker in (
                "team documents",
                "district documents",
                "busoga and part of central",
                "pdf to excel",
                "fwdugiftassets",
                "mda status register",
            )
        )
        if traceable_source and (not facility or not lg):
            try:
                traced = tracer.trace(source_file, row[index["Source locator"]])
            except (OSError, KeyError, ValueError, xlrd.biffh.XLRDError):
                traced = None
            if traced:
                traced_facility = master.canonical_facility(traced.facility, traced.lg or lg)
                if not lg and traced.lg:
                    lg, lg_rule = traced.lg, traced.lg_rule
                if not facility and traced_facility:
                    facility = traced_facility.name
                    lg = traced_facility.lg
                    facility_rule = "canonicalized source worksheet"
                    lg_rule = lg_rule or "source worksheet facility"
                elif not facility and traced.facility:
                    facility = usable_facility_text(traced.facility, master, traced.lg or lg)
                    if facility:
                        facility_rule = traced.facility_rule

        if not facility:
            described = facility_from_description(source_asset, master, lg)
            if described:
                facility, lg = described.name, described.lg
                facility_rule = "source description"
                lg_rule = lg_rule or "source description facility"

        if not facility:
            named = facility_from_filename(source_file, master, lg)
            if named:
                facility, lg = named.name, named.lg
                facility_rule = "source filename"
                lg_rule = lg_rule or "source filename facility"

        if (
            not facility
            and lg
            and source_facility
            and lg_key(source_facility) == lg_key(lg)
        ):
            facility = entity_name(lg)
            facility_rule = "LG named as source location"

        if not facility and lg and "district documents" in source_path_key:
            facility = entity_name(lg)
            facility_rule = "district-register owner"
        if not facility and lg in MDA_NAMES:
            facility = MDA_NAMES[lg]
            facility_rule = "MDA-register owner"

        source_context = normalized_key(
            f"{source_file} {row[index['Source locator']]} "
            f"{row[index['Source remarks and qualifications']]}"
        )
        if not lg:
            explicit_mda = mda_from_text(f"{source_facility} {facility} {source_asset}")
            if explicit_mda:
                lg = explicit_mda
                lg_rule = "explicit MDA name"
                facility = MDA_NAMES[lg]
                facility_rule = "explicit MDA name"
        if not lg:
            context_rules = (
                ("wemis", "MOWE", "WEMIS source context"),
                ("motorcycles for moh ugift", "MOH", "MOH worksheet heading"),
                ("for auditors", "OAG", "auditor-register context"),
                ("bped", "MOFPED", "BPED register context"),
                ("original entity heading nema", "NEMA", "MDA worksheet heading"),
                ("original entity heading ppda", "PPDA", "MDA worksheet heading"),
                ("fwdugiftassets", "MOFPED", "central UgIFT register owner"),
            )
            for marker, code, rule in context_rules:
                if marker in source_context:
                    lg, lg_rule = code, rule
                    break
        if not facility and lg in MDA_NAMES:
            facility = MDA_NAMES[lg]
            facility_rule = lg_rule

        represented = int(last) - int(first) + 1
        rules[f"LG: {lg_rule}" if lg else "LG: unresolved"] += represented
        rules[f"Facility: {facility_rule}" if facility else "Facility: unresolved"] += represented
        if not lg:
            unresolved["lg"][clean_text(source_file)] += represented
        if not facility:
            unresolved["facility"][clean_text(source_file)] += represented

        for register_row in range(int(first), int(last) + 1):
            if lg:
                lg_by_row[register_row] = lg
            if facility:
                facility_by_row[register_row] = facility

    return lg_by_row, facility_by_row, rules, unresolved


def shift_column_dimensions_for_insert(ws, insert_at: int) -> None:
    captured = {}
    for column_index in range(insert_at, ws.max_column + 1):
        letter = get_column_letter(column_index)
        if letter in ws.column_dimensions:
            captured[column_index] = copy(ws.column_dimensions[letter])
    for column_index in captured:
        del ws.column_dimensions[get_column_letter(column_index)]
    for column_index, dimension in captured.items():
        destination = get_column_letter(column_index + 1)
        dimension.index = destination
        dimension.min = column_index + 1
        dimension.max = column_index + 1
        ws.column_dimensions[destination] = dimension


def ensure_facility_column(ws) -> tuple[int, bool]:
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    if len(headers) > 1 and headers[0] == "BOOK_TYPE_CODE" and headers[1] == "FACILITY_NAME":
        return 2, False
    if "FACILITY_NAME" in headers:
        return headers.index("FACILITY_NAME") + 1, False

    shift_column_dimensions_for_insert(ws, 2)
    ws.insert_cols(2)
    ws.cell(1, 2).value = "FACILITY_NAME"
    ws.cell(1, 2)._style = copy(ws.cell(1, 1)._style)
    ws.column_dimensions["B"].width = 32
    return 2, True


def synchronize_asset_table(ws) -> None:
    table = ws.tables["AssetRegister"]
    headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    if len(table.tableColumns) == ws.max_column - 1 and headers[1] == "FACILITY_NAME":
        table.tableColumns.insert(1, TableColumn(id=2, name="FACILITY_NAME"))
    if len(table.tableColumns) != ws.max_column:
        raise ValueError(
            f"AssetRegister has {len(table.tableColumns)} table columns but the worksheet has {ws.max_column} columns"
        )
    for column_number, (column, header) in enumerate(zip(table.tableColumns, headers), 1):
        column.id = column_number
        column.name = str(header)
    table.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref


def update_read_me(ws) -> None:
    replacements = {
        "Template": (
            "The Asset Register sheet preserves the supplied template fields and adds "
            "FACILITY_NAME immediately after BOOK_TYPE_CODE, for 65 columns in total. "
            "The illustrative laptop row was removed."
        ),
        "Missing information": (
            "BOOK_TYPE_CODE and FACILITY_NAME are resolved from reconciled facility IDs, "
            "facility folders, source-line values, exact source worksheet rows and section "
            "headers, and explicit facility/LG text in source descriptions. District- or "
            "MDA-wide registers use the owning entity when no more specific facility is "
            "recorded. Values without defensible source evidence remain blank."
        ),
    }
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=3):
        label = row[0].value
        if label in replacements:
            row[1].value = replacements[label]
        if label == "BOOK_TYPE_CODE":
            row[1].value = (
                "Derive from the reconciled or source-traced local government/MDA, remove "
                "DLG/district wording, retain MC or City where applicable, and append ' BK'."
            )
            row[2].value = "Sample Header of Asset Register..xlsx; source-line reconciliation"


def load_input(path: Path, from_git_head: bool):
    if not from_git_head:
        return load_workbook(path)
    workbook_bytes = subprocess.check_output(
        ["git", "show", f"HEAD:{GIT_WORKBOOK_PATH}"], cwd=ROOT
    )
    return load_workbook(io.BytesIO(workbook_bytes))


def save_atomically(workbook, path: Path) -> None:
    temporary_path = path.with_name(f".{path.name}.tmp.xlsx")
    if temporary_path.exists():
        temporary_path.unlink()
    workbook.save(temporary_path)
    workbook.close()
    temporary_path.replace(path)


def build(path: Path, from_git_head: bool, dry_run: bool = False) -> dict[str, object]:
    workbook = load_input(path, from_git_head)
    master = MasterData()
    lg_by_row, facility_by_row, rules, unresolved = source_maps(
        workbook["Source Lines"], master
    )

    register = workbook["Asset Register"]
    facility_column, _ = ensure_facility_column(register)
    headers = [cell.value for cell in next(register.iter_rows(min_row=1, max_row=1))]
    location_column = headers.index("LOCATION_SEGMENT1") + 1

    stats: Counter = Counter()
    examples: list[tuple[int, str | None, str | None]] = []
    for row_number in range(2, register.max_row + 1):
        lg = lg_by_row.get(row_number)
        if not lg:
            lg = master.canonical_lg(register.cell(row_number, location_column).value)
            if lg:
                rules["LG: register location"] += 1

        book = book_code(lg)
        facility = facility_by_row.get(row_number)
        register.cell(row_number, 1).value = book
        register.cell(row_number, facility_column).value = facility
        stats["populated_book" if book else "blank_book"] += 1
        stats["populated_facility" if facility else "blank_facility"] += 1
        if len(examples) < 12 and (not book or not facility):
            examples.append((row_number, book, facility))

    synchronize_asset_table(register)
    update_read_me(workbook["Read Me"])
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    if dry_run:
        workbook.close()
    else:
        save_atomically(workbook, path)

    return {
        "populated_book": stats["populated_book"],
        "blank_book": stats["blank_book"],
        "populated_facility": stats["populated_facility"],
        "blank_facility": stats["blank_facility"],
        "rules": rules,
        "unresolved": unresolved,
        "unresolved_examples": examples,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", nargs="?", type=Path, default=DEFAULT_PATH)
    parser.add_argument(
        "--from-git-head", action="store_true", help="Start from the clean committed workbook."
    )
    parser.add_argument("--dry-run", action="store_true", help="Resolve and report values without saving.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    stats = build(args.path, args.from_git_head, args.dry_run)
    print(
        f"Filled BOOK_TYPE_CODE on {stats['populated_book']:,} rows; "
        f"left blank {stats['blank_book']:,}."
    )
    print(
        f"Filled FACILITY_NAME on {stats['populated_facility']:,} rows; "
        f"left blank {stats['blank_facility']:,}."
    )
    print("Resolution rules:")
    for rule, count in stats["rules"].most_common():
        print(f"  {count:>7,}  {rule}")
    for field in ("lg", "facility"):
        print(f"Top unresolved {field} sources:")
        for source_file, count in stats["unresolved"][field].most_common(12):
            print(f"  {count:>7,}  {source_file}")
    if stats["unresolved_examples"]:
        print("First unresolved rows:", stats["unresolved_examples"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
