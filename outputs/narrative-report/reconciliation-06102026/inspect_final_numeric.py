from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree as ET
import json

ROOT=Path(r'D:/coding/ugift-data-analysis');OUT=ROOT/'outputs/narrative-report/reconciliation-06102026'
PATH=ROOT/'outputs/narrative-report/ugift-working-report-06102026-1344-reconciled.docx'
doc=Document(PATH)
lines=[f'P{i} [{p.style.name}] {p.text}' for i,p in enumerate(doc.paragraphs) if p.text]
for i,t in enumerate(doc.tables):
    lines.append(f'TABLE {i}')
    lines.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
(OUT/'final-numeric-text.txt').write_text('\n'.join(lines),encoding='utf-8')
ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
charts=[]
with ZipFile(PATH) as z:
    for name in sorted(x for x in z.namelist() if x.startswith('word/charts/') and x.endswith('.xml')):
        root=ET.fromstring(z.read(name));data=[]
        for ser in root.findall('.//c:ser',ns):
            labels=ser.findall('.//c:cat//c:strCache/c:pt/c:v',ns)
            values=ser.findall('.//c:val//c:numCache/c:pt/c:v',ns)
            data.append({'label':ser.findtext('c:tx/c:strRef/c:strCache/c:pt/c:v',namespaces=ns) or ser.findtext('c:tx/c:v',namespaces=ns),'categories':[x.text for x in labels],'values':[float(x.text) for x in values]})
        charts.append({'name':name,'series':data})
(OUT/'final-chart-caches.json').write_text(json.dumps(charts,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps(charts,indent=2,ensure_ascii=False))
