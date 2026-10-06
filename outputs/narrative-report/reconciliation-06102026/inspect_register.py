from pathlib import Path
import json
import openpyxl
from docx import Document

base=Path(r'D:/coding/ugift-data-analysis')
out=base/'outputs/narrative-report/reconciliation-06102026'
out.mkdir(exist_ok=True)
wb=openpyxl.load_workbook(base/'outputs/asset-register/REF_ALL_UGIFT_ASSET_REGISTER_MF_TEMPLATE.xlsx',read_only=True,data_only=True)
for ws in wb:
    print('SHEET',ws.title,ws.max_row,ws.max_column)
    for row in ws.iter_rows(min_row=1,max_row=min(4,ws.max_row),values_only=True): print(json.dumps(row,default=str))
for suffix in ['1344','0000']:
    doc=Document(base/f'outputs/narrative-report/ugift-working-report-06102026-{suffix}.docx')
    lines=[]
    for i,p in enumerate(doc.paragraphs):
        if p.text: lines.append(f'P{i} [{p.style.name}] {p.text}')
    for i,t in enumerate(doc.tables):
        lines.append(f'TABLE {i}')
        lines.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
    (out/f'register-audit-report-{suffix}.txt').write_text('\n'.join(lines),encoding='utf-8')
    print('DOC',suffix,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
