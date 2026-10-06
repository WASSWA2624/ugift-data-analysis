"""Read-only numerical audit of chart caches against embedded workbook values."""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import hashlib
import json
import posixpath
import re
import sys

from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(r'D:/coding/ugift-data-analysis')
OUT = ROOT / 'outputs/narrative-report/reconciliation-06102026'
PATH = ROOT / 'outputs/narrative-report/ugift-working-report-06102026-1344-reconciled.docx'
NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'pr': 'http://schemas.openxmlformats.org/package/2006/relationships',
      'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def cache_values(ref):
    cache = ref.find('c:strCache', NS)
    if cache is None:
        cache = ref.find('c:numCache', NS)
    if cache is None:
        return []
    values = []
    for point in sorted(cache.findall('c:pt', NS), key=lambda p: int(p.get('idx'))):
        values.append(point.findtext('c:v', namespaces=NS))
    return values


def sheet_values(workbook, formula):
    sheet_name, cell_range = formula.split('!', 1)
    sheet_name = sheet_name.strip("'").replace("''", "'")
    sheet = workbook[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    return [sheet.cell(row, col).value for row in range(min_row, max_row + 1)
            for col in range(min_col, max_col + 1)]


def equal_value(a, b):
    try:
        return abs(float(a) - float(b)) <= max(1e-12, abs(float(a)) * 1e-12)
    except (TypeError, ValueError):
        return str(a) == str(b)


def signature(categories, labels):
    return tuple(categories), tuple(labels)


recipes = json.loads((OUT / 'reconciliation-applied.json').read_text(encoding='utf-8'))['charts']
only_chart = sys.argv[1] if len(sys.argv) > 1 else None
recipes_by_signature = {}
for recipe in recipes:
    recipes_by_signature.setdefault(signature(recipe['categories'], recipe['labels']), []).append(recipe)
result = {'report': str(PATH), 'report_sha256': hashlib.sha256(PATH.read_bytes()).hexdigest(),
          'checked_chart_count': 0, 'cache_workbook_cells_checked': 0,
          'cache_workbook_discrepancies': [], 'recipe_discrepancies': [], 'charts': []}
with ZipFile(PATH) as archive:
    document = ET.fromstring(archive.read('word/document.xml'))
    rels = ET.fromstring(archive.read('word/_rels/document.xml.rels'))
    part_by_id = {r.get('Id'): posixpath.normpath('word/' + r.get('Target')) for r in rels}
    caption_by_part = {}
    body = document.find('w:body', NS)
    for index, element in enumerate(body):
        nodes = element.findall('.//c:chart', NS)
        if not nodes:
            continue
        nearby = []
        for following in list(body)[index + 1:index + 5]:
            text = ''.join(following.itertext()) if following.tag.endswith('}t') else ''.join(following.xpath('.//w:t/text()', namespaces=NS))
            if text.strip():
                nearby.append(text)
        for node in nodes:
            caption_by_part[part_by_id[node.get('{'+NS['r']+'}id')]] = nearby
    for part in sorted((n for n in archive.namelist() if re.fullmatch(r'word/charts/chart[^/]*\.xml', n)), key=lambda n: int(re.search(r'(\d+)\.xml$', n).group(1))):
        if only_chart and not part.endswith('/chart'+only_chart+'.xml'):
            continue
        chart = ET.fromstring(archive.read(part))
        rel_path = posixpath.dirname(part) + '/_rels/' + posixpath.basename(part) + '.rels'
        chart_rels = ET.fromstring(archive.read(rel_path))
        external = chart.find('c:externalData', NS)
        external_id = external.get('{'+NS['r']+'}id')
        target = next(r.get('Target') for r in chart_rels if r.get('Id') == external_id)
        embedded_part = posixpath.normpath(posixpath.dirname(part) + '/' + target)
        workbook = load_workbook(BytesIO(archive.read(embedded_part)), read_only=True, data_only=True)
        series = []
        for ser in chart.findall('.//c:ser', NS):
            for reference in ser.findall('.//c:strRef', NS) + ser.findall('.//c:numRef', NS):
                formula = reference.findtext('c:f', namespaces=NS)
                cached = cache_values(reference)
                saved = sheet_values(workbook, formula)
                result['cache_workbook_cells_checked'] += len(cached)
                if len(cached) != len(saved) or any(not equal_value(a, b) for a, b in zip(cached, saved)):
                    result['cache_workbook_discrepancies'].append({'part': part, 'formula': formula, 'cache': cached, 'workbook': saved})
            label = ser.findtext('c:tx/c:strRef/c:strCache/c:pt/c:v', namespaces=NS) or ser.findtext('c:tx/c:v', namespaces=NS)
            categories = [p.text for p in ser.findall('.//c:cat//c:strCache/c:pt/c:v', NS)]
            values = [float(p.text) for p in ser.findall('.//c:val//c:numCache/c:pt/c:v', NS)]
            series.append({'label': label, 'categories': categories, 'values': values, 'sum': sum(values)})
        workbook.close()
        categories = series[0]['categories']
        candidates = recipes_by_signature.get(signature(categories, [s['label'] for s in series]), [])
        matches = [recipe for recipe in candidates if len(recipe['values']) == len(series) and all(len(rv) == len(s['values']) and all(equal_value(a, b) for a, b in zip(rv, s['values'])) for rv, s in zip(recipe['values'], series))]
        if not candidates:
            result['recipe_discrepancies'].append({'part': part, 'reason': 'No matching recipe labels/categories'})
        elif not matches:
            result['recipe_discrepancies'].append({'part': part, 'reason': 'Values differ', 'recipes': [r['values'] for r in candidates], 'actual': [s['values'] for s in series]})
        result['charts'].append({'part': part, 'embedded_workbook': embedded_part, 'series': series, 'following_text': caption_by_part.get(part, [])})
        result['checked_chart_count'] += 1
(OUT / ('final-native-chart'+only_chart+'-validation.json' if only_chart else 'final-native-chart-validation.json')).write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k != 'charts'}, indent=2, ensure_ascii=False))
for c in result['charts']:
    print(c['part'], [(s['label'],s['values']) for s in c['series']], c['following_text'][:1])
