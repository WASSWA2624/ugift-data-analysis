from docx import Document

d = Document(r"D:\coding\ugift-data-analysis\outputs\narrative-report\UgIFT Asset Verification Report.docx")
for i, p in enumerate(d.paragraphs[28:90]):
    text = p.text.strip()
    if text:
        print(f"{i+28}|{text[:130]}")
print("--- figures ---")
for i, p in enumerate(d.paragraphs):
    if p.text.startswith("Figure 46"):
        print(i, p.text[:120])
